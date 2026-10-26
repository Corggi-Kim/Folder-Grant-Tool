# =====================================
# 1) 폴더 구조
#    project/
#      ├─ Folder Grant API Client.pyw
#      └─ assets/
#            ├─ logo.png
#            └─ fgt.ico
#
# 2) 실행 시 생성/사용 경로 (코드에서 자동 생성)
#    - C:\FGT\Log       : 실행 로그(access.log)
#    - C:\FGT\conf      : 설정(login.json, theme 등)
#    - C:\FGT\ef         : 엑셀파일 다운로드
#    - C:\FGT\debug   : 디버그 파일
#
# 3) EXE 압축 시 모듈 임포트 필요 (--collect-all 옵션)
#    - selenium
# =====================================

import sys, os, glob, datetime, re, shutil, json, time
import html as htmllib
from typing import Dict, List
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import ElementClickInterceptedException, TimeoutException, StaleElementReferenceException
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QPushButton, QFileDialog, QTableWidget,
    QTableWidgetItem, QAbstractItemView, QVBoxLayout, QWidget, QTextEdit, QHBoxLayout,
    QLabel, QHeaderView, QCheckBox, QMenu, QToolButton, QMessageBox, QStyleOptionButton,
    QStyle, QDialog, QLineEdit, QComboBox, QDialogButtonBox, QFormLayout, QStatusBar,
    QTabWidget, QScrollArea, QSizePolicy, QSpinBox, QProgressBar, QAbstractButton, QTabBar, QStackedWidget, QAction, QStyledItemDelegate, QStyleOptionViewItem, QStyleOptionTab
)
from PyQt5.QtGui import QFont, QGuiApplication, QKeySequence, QPainter, QTextOption, QTextCursor, QPixmap, QIcon, QColor, QPen, QPainterPath
from PyQt5.QtCore import Qt, QRect, pyqtSignal, QThread, QObject, pyqtSlot, QTimer, QPoint, QSize
from openpyxl import load_workbook

APP_NAME = "Folder Grant API Client"
APP_VERSION = "1.1.0"  #구조변경, 기능추가/수정, 오류/버그수정
APP_BUILD = "2026-10-22"
APP_VERSION_STR = f"v{APP_VERSION}"

THEMES = {'light': {'bg': '#f4f6fa',
           'panel': '#ffffff',
           'header': '#f5f7fb',
           'text': '#1f2937',
           'muted': '#69788f',
           'border': '#dce3ec',
           'grid': '#e9edf4',
           'button_border': '#b9c6d8',
           'accent': '#2563eb',
           'accent_hover': '#1d4ed8',
           'select': '#e0ebff',
           'hover': '#f0f5ff',
           'badge': '#edf1f7',
           'blue_bg': '#eaf2ff',
           'green': '#16a34a',
           'amber': '#ce8109',
           'amber_bg': '#fff6e6',
           'amber_border': '#f1d6a5',
           'red': '#e5484d',
           'disabled': '#9da8b8'},
 'dark': {'bg': '#151b26',
          'panel': '#1e2634',
          'header': '#202b3b',
          'text': '#e5eaf2',
          'muted': '#a0aec2',
          'border': '#334155',
          'grid': '#2e3a4a',
          'button_border': '#5a6b82',
          'accent': '#3b82f6',
          'accent_hover': '#5292ff',
          'select': '#284468',
          'hover': '#29384d',
          'badge': '#2b3749',
          'blue_bg': '#243b59',
          'green': '#22c55e',
          'amber': '#f5b644',
          'amber_bg': '#3b3022',
          'amber_border': '#73592e',
          'red': '#fb7185',
          'disabled': '#6c788c'}}

HELP_PAGES = (
    ("처음 사용하기", (
        ("01 · 연결 확인", "프로그램을 켜면 서버 연결과 BUS 로그인 후 요청을 불러옵니다. 로그인창이 나오면 BUS 계정과 암호를 입력하세요. 계정 저장은 직접 선택할 수 있습니다."),
        ("02 · 요청 살펴보기", "권한 요청 또는 신규 프로젝트 탭에서 목록을 확인하세요. 행을 선택하면 오른쪽 선택 요청 상세에 해당 요청의 정보와 진행 상태가 표시됩니다."),
        ("03 · 확인하고 실행", "처리할 항목의 체크박스를 선택하세요. 권한 요청은 요청 검증으로 먼저 확인한 뒤 선택 실행을 누르세요. 신규 프로젝트도 요청 검증으로 확인한 뒤 선택실행을 누르세요."),
        ("04 · 완료 확인", "작업 결과를 목록과 오른쪽 상세에서 확인하세요. BUS 자동 완료 옵션을 켜 두면 실제 작업이 모두 성공한 BUS 요청만 완료 처리됩니다. 수동 입력과 엑셀로 추가한 항목은 BUS 완료 대상이 아닙니다."),
    )),
    ("버튼 안내", (
        ("권한 요청 · 목록 준비", "수동 입력은 직접 요청을 추가할 때, 엑셀 불러오기는 파일에 정리된 요청을 가져올 때 사용합니다. BUS 요청 조회를 누르면 BUS의 권한 요청을 다시 불러옵니다."),
        ("신규 프로젝트 · 목록 준비", "수동 추가로 프로젝트를 직접 입력하거나 BUS 신규조회로 프로젝트 생성 요청을 불러오세요. 필요한 항목을 체크한 뒤 아래의 선택 실행을 누르면 됩니다."),
        ("요청 검증 / 선택 실행", "요청 검증은 API 서버에서 사용자·보안 그룹·요청 경로와 현재 ACL을 읽기 전용으로 확인합니다. 실제 권한이나 폴더를 변경하지 않으며 항목별 결과는 작업 로그에 표시합니다. 선택 실행은 체크한 항목을 서버에 보내 처리합니다."),
        ("BUS 완료 재확인", "권한 작업은 성공했는데 BUS 완료가 되지 않았을 때 사용하세요. 서버 결과를 다시 확인한 뒤 완료 처리를 시도합니다. 권한 작업 자체를 다시 실행하는 버튼은 아닙니다."),
        ("중지", "다음 항목의 실행을 멈추고 진행 중인 작업에 취소를 요청합니다. 이미 진행 중인 처리는 멈추기까지 시간이 걸릴 수 있으며, 적용된 권한이 자동으로 되돌아가지는 않습니다."),
    )),
    ("문제 해결", (
        ("부분 성공은 어떤 뜻인가요?", "일부 단계만 처리된 상태입니다. 예를 들어 그룹에 사용자를 추가했지만 폴더 권한 적용에는 실패했을 수 있습니다. 오른쪽 상세에서 실패한 단계와 이유를 확인하세요. 이 상태에서는 BUS를 완료 처리하지 않습니다."),
        ("문제를 고친 뒤 다시 실행하려면?", "폴더 누락 등 표시된 원인을 먼저 해결하세요. 실패하거나 부분 성공한 권한 요청을 체크하고 선택 실행을 누른 뒤 재시도 안내를 확인하세요. 이전에 성공한 단계는 유지하고 남은 단계를 다시 처리합니다."),
        ("시간 초과 또는 결과 확인이 안 될 때", "오래 걸리는 작업은 아직 처리 중일 수 있습니다. 권한 요청의 BUS 완료 재확인으로 서버 상태를 확인하세요. 실패가 확인되면 원인을 해결한 뒤 재시도하고, 결과가 확실하지 않다면 중복 실행 전에 관리자에게 확인하세요."),
        ("목록이 비어 있거나 조회에 실패했을 때", "상단의 서버 연결과 BUS 로그인 상태를 확인하고 해당 탭의 BUS 조회 버튼을 눌러 주세요. 연결 정보나 계정이 바뀌었다면 톱니바퀴 메뉴의 설정에서 수정할 수 있습니다."),
        ("목록에서 수정하거나 삭제할 때", "BUS에서 조회한 권한 요청은 목록에서 수정할 수 없습니다. 직접 입력하거나 엑셀로 불러온 권한 요청만 수정할 수 있습니다. 두 목록에서 마우스 오른쪽 버튼을 누르면 행 삭제, 선택 행 삭제, 전체 삭제를 사용할 수 있습니다. 삭제는 화면 목록에만 적용되며 BUS 요청이나 서버 이력은 삭제하지 않습니다."),
    )),
    ("설정과 기록", (
        ("연결과 로그인 설정", "오른쪽 위 톱니바퀴 메뉴에서 서버 주소, API Key와 BUS 로그인 정보를 관리하세요. API Key는 관리자에게 받은 값을 입력합니다. 저장한 API Key는 이 PC의 설정 파일에 평문으로 보관됩니다."),
        ("숨겨진 정보 확인과 화면 테마", "이름·부서·Role 열은 톱니바퀴 메뉴의 상세 열 표시에서 볼 수 있습니다. 화면의 밝기는 오른쪽 위 달 또는 해 모양 아이콘으로 바꿀 수 있습니다."),
        ("작업 로그와 상세 기록", "화면 아래 작업 로그에는 처리 결과를 짧게 보여줍니다. 오류 원인을 자세히 확인할 때는 C:\\FGT\\Log\\access.log 파일을 확인하세요. 같은 파일에 날짜와 시간을 포함해 계속 기록합니다. 로그 지우기는 화면의 로그만 비웁니다."),
        ("다운로드한 파일", "BUS에서 내려받은 파일은 C:\\FGT\\ef 폴더에서 확인할 수 있습니다."),
    )),
)


class HelpDialog(QDialog):
    """User-facing task guide, using the same palette as the main window."""
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("도움말")
        self.resize(800, 660)
        t = THEMES[getattr(parent, 'current_theme', 'light')]
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 22, 24, 18)
        layout.setSpacing(14)
        title = QLabel("Folder Grant 도움말", self)
        title.setStyleSheet("font-size:22px;font-weight:700;background:transparent;")
        layout.addWidget(title)
        subtitle = QLabel("요청 조회부터 완료 처리까지, 프로그램의 주요 기능과 사용 방법을 안내합니다.", self)
        subtitle.setWordWrap(True)
        subtitle.setStyleSheet(f"color:{t['muted']};background:transparent;")
        layout.addWidget(subtitle)
        self.tabs = QTabWidget(self)
        self.tabs.setDocumentMode(True)
        layout.addWidget(self.tabs, 1)
        for label, cards in HELP_PAGES:
            page = QTextEdit(self.tabs)
            page.setReadOnly(True)
            page.setFrameShape(QTextEdit.NoFrame)
            page.document().setDefaultStyleSheet(
                f"body {{color:{t['text']};font-size:13px;}} "
                "h3 {font-size:15px;margin-top:0px;margin-bottom:8px;} "
                "p {margin-top:0px;margin-bottom:0px;line-height:150%;}")
            parts = []
            for heading, body in cards:
                parts.append(f'<table width="100%" cellspacing="0" cellpadding="14" '
                             f'bgcolor="{t["panel"]}"><tr><td>'
                             f'<h3 style="color:{t["accent"]}">{htmllib.escape(heading)}</h3>'
                             f'<p>{htmllib.escape(body)}</p></td></tr></table><p style="font-size:8px">&nbsp;</p>')
            page.setHtml('<html><body>'+''.join(parts)+'</body></html>')
            self.tabs.addTab(page, label)
        buttons = QDialogButtonBox(self)
        close = buttons.addButton("닫기", QDialogButtonBox.RejectRole)
        close.setDefault(True)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)
        self.setStyleSheet(f"QDialog {{background:{t['bg']};color:{t['text']};}} "
                           f"QTextEdit {{background:{t['bg']};color:{t['text']};padding:10px 0px;}} "
                           "QTabBar::tab {padding:10px 14px;min-height:22px;} "
                           f"QTabBar::tab:selected {{color:{t['accent']};border-bottom:2px solid {t['accent']};}}")


LOG_DIR = r"C:\FGT\Log"
CONF_DIR = r"C:\FGT\conf"
DL_DIR = r"C:\FGT\ef"
DEBUG_DIR = r"C:\FGT\debug"

GROUP_OU_PATH = r"OU=Group Project Folder,OU=0.Management Object Group,OU=lskglobal,DC=lskglobal,DC=com"
TEMPLATE_ROOT = r"\\LSK_S010\Study folder\_Template"

CONF_FILE = os.path.join(CONF_DIR, "login.json")
LOGO_FILE = "logo.png"
ICON_FILE = "fgt.ico"



DOMAIN_EMAIL_SUFFIX = "lskglobal.com"

BUS_LOGIN_URL = "https://bus.lskglobal.com/L4/Common/Login.aspx"
BUS_PROGRESS_LIST_URL = "https://bus.lskglobal.com/L4/Common/Default.aspx?7VHoVKC6bjriQDXQa/t/dQ=="
BUS_END_LIST_URL = "https://bus.lskglobal.com/L4/Common/Default.aspx?oRrCZc631pq4qaUZnht8Cg=="
BUS_NEW_URL = "https://bus.lskglobal.com/L4/Common/Default.aspx?2QiSDYhMfx5ql2mmNbos4A=="

REQ_GRANT = "권한부여"
REQ_RELEASE = "권한해제"
END_FLAG_GRANT = "END_GRANT"
END_FLAG_RELEASE = "END_RELEASE"

RELEASE_HINT_HEADERS = {
    "해제요청일", "해제요청자사번", "해제요청자성명",
    "해제요청자부서/팀", "해제요청자직책"
}

SHARE_ROOT = r"\\LSK_S010\Study Folder\{proj_seg}\{lv2}\{lv3}"
CLOSED_ROOT = r"\\192.168.1.95\Study_Closed"
CLOSED_ARCHIVE_ROOT = r"\\192.168.1.95\Study_Archive"

LEVEL3_CHOICES = ["ARS","CO","DM","ER","MW","PM","PV","RA","SSU","STAT","STAT_IDMC","ETC"]

ROLE_MAP: Dict[str, List[str]] = {
    "Trial STAT/SP": ["3.Dataset", "4.Analysis", "5.SDTM", "6.Validation"],
    "Verification SP": ["3.Dataset", "6.Validation", "8.Verification"],
    "SDTM": ["3.Dataset", "5.SDTM", "6.Validation"],
    "Manager": ["3.Dataset", "4.Analysis", "5.SDTM", "6.Validation", "8.Verification"],
    "Randomization Statistician": ["Random"],
    "Blind Reviewer": ["Reviewer"],
    "Unblind Reviewer": ["Random", "Reviewer"],
}

STUDY_ROLES = {"Trial STAT/SP", "Verification SP", "SDTM", "Manager"}
ISOLATED_ROLES = {"Randomization Statistician", "Blind Reviewer", "Unblind Reviewer"}

ISOLATED_STAT_IDMC_ROLE_MAP: Dict[str, List[str]] = {
    "Trial STAT/SP": ["8.Verification"],
    "Verification SP": ["4.Analysis", "5.SDTM"],
    "SDTM": ["4.Analysis", "8.Verification"],
    "Manager": [],
}

LEGACY_STUDY_MAP = {
    "Trial STAT/SP": [3, 4, 5],
    "Verification SP": [3, 5, 8],
    "SDTM": [3, 5],
    "Manager": [3, 4, 5, 8],
}

NEW_THRESHOLD = 25069
FORCE_NEW_CODES = {
    "25-038", "25-032", "25-028", "25-006", "25-004",
    "24-084", "24-076", "24-072", "24-061", "24-037", "24-033", "24-026", "24-009",
    "21-040",
}
STAT_IDMC_NEW_THRESHOLD = 26012
STAT_IDMC_FORCE_NEW_CODES = {"25-074", "25-077"}

HEADER_ALIASES = {
    "user":   {"대상자사번"},
    "name": {"대상자성명"},
    "proj":   {"프로젝트코드"},
    "level1": {"폴더level1","폴더Level1"},
    "level2": {"폴더level2","폴더Level2"},
    "level3": {"폴더level3","폴더Level3"},
    "role":   {"statrole","STATROLE"},
    "dept":   {"대상자부서/팀"},
}

def build_root_from_proj(proj_raw: str) -> str:
    seg = proj_segment_for_folder(proj_raw)
    return r"\\LSK_S010\Study Folder\{}".format(seg)

def resource_path(rel_path: str) -> str:
    try:
        base = sys._MEIPASS
    except Exception:
        base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, rel_path)



def make_hidden_chrome_options(download_dir: str | None = None):
    options = webdriver.ChromeOptions()
    if download_dir:
        prefs = {
            "download.default_directory": download_dir,
            "download.prompt_for_download": False,
            "download.directory_upgrade": True,
            "safebrowsing.enabled": True,
        }
        options.add_experimental_option("prefs", prefs)

    options.add_argument("--headless=new")
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    options.add_argument("--window-size=1280,900")
    options.add_argument("--window-position=-32000,-32000")
    return options

def _norm_k(h: str) -> str:
    return (h or "").strip().lower().replace(" ", "").replace("/", "").replace("_","")
RELEASE_HINT_HEADERS_NORM = {_norm_k(h) for h in RELEASE_HINT_HEADERS}

def _find_cols(header_row, wanted_norm_set):
    idxs = {}
    norm = [_norm_k(h) for h in header_row]
    for i, h in enumerate(norm):
        if h in wanted_norm_set and h not in idxs:
            idxs[h] = i
    return idxs

def _is_release_row_by_values(row_values, header_row):
    try:
        idxs = _find_cols(header_row, RELEASE_HINT_HEADERS_NORM)
        for k, i in idxs.items():
            if i < len(row_values):
                v = row_values[i]
                if v is not None and str(v).strip() != "":
                    return True
    except Exception:
        pass
    return False

def _norm_header(h: str) -> str:
    return (h or "").strip().lower().replace(" ", "").replace("/", "").replace("_","").replace("-","")

def auto_map_columns(header_row):
    idx = {}
    norm_headers = [_norm_header(h if h is not None else "") for h in header_row]
    for key, aliases in HEADER_ALIASES.items():
        aliases_norm = {_norm_header(a) for a in aliases}
        for i, h in enumerate(norm_headers):
            if h in aliases_norm:
                idx[key] = i
                break
    return idx

def normalize_lv2(s: str) -> str:
    t = (s or "").strip().lower()
    return "Isolated" if "iso" in t else "Study"

def is_stat_idmc_lv3(lv3: str) -> bool:
    return (lv3 or "").strip().upper() == "STAT_IDMC"

def is_stat_lv3(lv3: str) -> bool:
    return (lv3 or "").strip().upper() in {"STAT", "STAT_IDMC"}

def is_lv3_etc(lv3: str) -> bool:
    return (lv3 or "").strip().lower() == "etc"

def insert_zero_middle_4digit(code4: str) -> str:
    return code4[:2] + "0" + code4[2:]

def split_proj_and_suffix(raw: str):
    s = (raw or "").strip()
    if "-" in s:
        base, suf = s.split("-", 1)
        return base.strip(), suf.strip()
    return s, ""

def proj_segment_for_folder(raw: str) -> str:
    base, suf = split_proj_and_suffix(raw)
    digits = re.sub(r"\D", "", base or "")
    if len(digits) == 4:
        digits = insert_zero_middle_4digit(digits)
    return digits + ("A" + suf if suf else "")

def group_digits_from_proj_for_groupname(raw: str) -> (str, str):
    base, suf = split_proj_and_suffix(raw)
    digits = re.sub(r"\D", "", base or "")
    if len(digits) == 4:
        digits = insert_zero_middle_4digit(digits)
    last5 = digits[-5:] if len(digits) >= 5 else digits.zfill(5)
    return last5, suf

def format_group_name(proj_raw: str, lv2: str) -> str:
    last5, suf = group_digits_from_proj_for_groupname(proj_raw)
    yy, xxx = last5[:2], last5[2:]
    name = f"LSK {yy}-{xxx}"
    if suf:
        name += f"-{suf}"
    if normalize_lv2(lv2) == "Isolated":
        name += " Isolated"
    return name

def _yyxxx_from_proj_for_groupname(proj_raw: str) -> str:
    last5, _ = group_digits_from_proj_for_groupname(proj_raw)
    return last5

def _lsk_code_for_compare(proj_raw: str) -> str:
    last5 = _yyxxx_from_proj_for_groupname(proj_raw)
    return f"{last5[:2]}-{last5[2:]}"

def is_new_template(proj_raw: str) -> bool:
    code_yyxxx = int(_yyxxx_from_proj_for_groupname(proj_raw))
    lsk_code = _lsk_code_for_compare(proj_raw)
    if lsk_code in FORCE_NEW_CODES:
        return True
    return code_yyxxx >= NEW_THRESHOLD

def is_stat_idmc_new_policy(proj_raw: str) -> bool:
    lsk_code = _lsk_code_for_compare(proj_raw)
    if lsk_code in STAT_IDMC_FORCE_NEW_CODES:
        return True
    try:
        code_yyxxx = int(_yyxxx_from_proj_for_groupname(proj_raw))
    except Exception:
        return False
    return code_yyxxx >= STAT_IDMC_NEW_THRESHOLD

def build_path_l3(proj_raw: str, lv2: str, lv3: str) -> str:
    seg = proj_segment_for_folder(proj_raw)
    return SHARE_ROOT.format(
        proj_seg=seg,
        lv2=normalize_lv2(lv2),
        lv3=(lv3 or "").strip().replace("\\","").replace("/","")
    )

def closed_segment_from_proj(proj_raw: str) -> str:
    s = (proj_raw or "").strip()
    m = re.fullmatch(r"(\d{5})(?:A(\d+))?$", s)
    if m:
        return s

    base, suf = split_proj_and_suffix(s)
    digits = re.sub(r"\D", "", base or "")
    if len(digits) == 4:
        digits = insert_zero_middle_4digit(digits)
    seg = digits[-5:].zfill(5)

    suf_digits = re.sub(r"\D", "", suf or "")
    return seg + (("A" + suf_digits) if suf_digits else "")

def build_closed_path_from_proj(proj_raw: str) -> str:
    return build_closed_candidate_paths_from_proj(proj_raw)[0]

def build_closed_candidate_paths_from_proj(proj_raw: str) -> List[str]:
    seg = closed_segment_from_proj(proj_raw)
    return [
        os.path.join(CLOSED_ROOT, seg),
        os.path.join(CLOSED_ARCHIVE_ROOT, seg),
    ]


def _extract_tables_from_html(s: str):
    tables = re.findall(r"<table[^>]*>(.*?)</table>", s, re.I | re.S)
    out = []
    for tbl in tables:
        rows = re.findall(r"<tr[^>]*>(.*?)</tr>", tbl, re.I | re.S)
        parsed = []
        for r in rows:
            cells = re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", r, re.I | re.S)
            vals = []
            for c in cells:
                t = re.sub(r"<[^>]+>", "", c)
                t = htmllib.unescape(t).strip()
                vals.append(t)
            if vals:
                parsed.append(vals)
        if parsed:
            out.append(parsed)
    return out

def _parse_html_best_table(path: str):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        html = f.read()
    candidates = _extract_tables_from_html(html)
    if not candidates:
        return [], []
    for tbl in candidates:
        header = tbl[0]
        data = tbl[1:]
        return header, data
    return [], []

# API contract and transport. No server package dependency.
import hashlib
import threading
from uuid import uuid4
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlsplit
from urllib.request import Request, build_opener, HTTPRedirectHandler

DEFAULT_API_URL = 'http://192.168.1.10:8000'
API_CONFIG_FILE = os.path.join(CONF_DIR, 'api_client.json')
FINAL_JOB_STATUSES = {'simulated', 'succeeded', 'partially_succeeded', 'failed', 'cancelled'}


def canonical_project_code(raw: str) -> str:
    value = str(raw).strip().upper()
    # BUS displays YY-NNN; the API accepts YYNNN and optional A suffixes.
    value = re.sub(r'^(\d{2})-(\d{3})(?=A\d+$|$)', r'\1\2', value)
    if not re.fullmatch(r'(?:\d{4}|\d{5}(?:A\d+|-\d+)?)', value):
        raise ValueError('프로젝트 코드 형식을 확인하세요: ' + value)
    return value


def make_access_payload(values: dict, request_id: str, source: str, requested_by: str) -> dict:
    closed = values.get('kind') == '종료'
    user = str(values.get('user', '')).strip()
    if not re.fullmatch(r'[A-Za-z0-9._-]+', user):
        raise ValueError('대상자 사번을 확인하세요.')
    operation = values.get('req', REQ_GRANT)
    if operation not in {REQ_GRANT, REQ_RELEASE}:
        raise ValueError('권한부여 또는 권한해제 요청이어야 합니다.')
    lv2 = normalize_lv2(str(values.get('lv2', '')).strip())
    lv3 = str(values.get('lv3', '')).strip().upper()
    if not closed and (lv2 not in {'Study', 'Isolated'} or lv3 not in LEVEL3_CHOICES):
        raise ValueError('진행 과제의 Level2·Level3를 확인하세요.')
    return dict(request_id=request_id, source=source,
                operation='revoke' if operation == REQ_RELEASE else 'grant',
                project_status='closed' if closed else 'progress', employee_id=user,
                project_code=canonical_project_code(values.get('proj', '')),
                level2=None if closed else lv2, level3=None if closed else lv3,
                role=None if closed else (str(values.get('role', '')).strip() or None),
                requested_by=requested_by)


def make_project_payload(project_code: str, project_name: str, request_id: str,
                         source: str, requested_by: str) -> dict:
    return dict(request_id=request_id, source=source,
                project_code=canonical_project_code(project_code),
                project_name=project_name.strip(), requested_by=requested_by)


def can_complete_bus(result: dict, source: str, stopped: bool = False) -> bool:
    return bool(not stopped and source == 'bus' and result.get('status') == 'succeeded'
                and result.get('executor_mode') == 'powershell')


def make_bus_request_id(values: dict, bus_identity: str = '') -> str:
    stable = {k: str(v or '').strip() for k, v in values.items()
              if k in {'kind', 'req', 'user', 'proj', 'lv2', 'lv3', 'role', 'project_name'}}
    if 'proj' in stable:
        stable['proj'] = canonical_project_code(stable['proj'])
    stable['bus_identity'] = bus_identity.strip()
    digest = hashlib.sha256(json.dumps(stable, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    return 'BUS-' + digest


def new_client_request_id() -> str:
    return 'CLIENT-' + uuid4().hex


def bus_row_identity(header, values) -> str:
    fields = {str(h).strip(): str(v or '').strip() for h,v in zip(header,values)
              if str(h).strip().lower() not in {'','순번','번호','no','no.','#','선택'}}
    for name,value in fields.items():
        if re.sub(r'[\s_-]+','',name).lower() in {'요청번호','신청번호','requestid','요청id','신청id'} and value:
            return name + ':' + value
    # Keep request dates and other original matching fields when no ID is exported.
    return json.dumps(fields,ensure_ascii=False,sort_keys=True)


class ApiClientError(RuntimeError):
    def __init__(self, message, status_code=None):
        super().__init__(message)
        self.status_code = status_code


class _NoApiRedirects(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None


class FolderGrantApiClient:
    def __init__(self, base_url: str, api_key: str, timeout_seconds: int = 30):
        self.base_url = base_url.strip().rstrip('/')
        parsed = urlsplit(self.base_url)
        if parsed.scheme not in {'http', 'https'} or not parsed.hostname or parsed.username or parsed.password or parsed.query or parsed.fragment:
            raise ValueError('API 주소는 http://서버:포트 또는 https://서버 형식이어야 합니다.')
        self.api_key = api_key.strip()
        self.timeout_seconds = timeout_seconds
        self.opener = build_opener(_NoApiRedirects())

    def request(self, method: str, path: str, payload: dict | None = None) -> dict:
        authenticated = path != '/health'
        if authenticated and not self.api_key:
            raise ApiClientError('API Key를 입력하세요.')
        headers = {'Accept': 'application/json'}
        if authenticated:
            headers['X-API-Key'] = self.api_key
        body = None
        if payload is not None:
            body = json.dumps(payload, ensure_ascii=False).encode('utf-8')
            headers['Content-Type'] = 'application/json; charset=utf-8'
        try:
            with self.opener.open(Request(self.base_url + path, data=body, headers=headers,
                                         method=method), timeout=self.timeout_seconds) as response:
                data = json.loads(response.read().decode('utf-8'))
                if not isinstance(data, dict):
                    raise ValueError('JSON object required')
                return data
        except HTTPError as exc:
            raw = exc.read().decode('utf-8', errors='replace')
            try:
                detail = json.loads(raw).get('detail', raw)
                if isinstance(detail, dict):
                    detail = detail.get('message') or detail.get('code') or detail
                elif isinstance(detail, list):
                    detail = '; '.join(str(x.get('msg', x)) if isinstance(x, dict) else str(x) for x in detail)
            except (ValueError, AttributeError):
                detail = raw or exc.reason
            message = f'HTTP {exc.code}: {detail}'
            if self.api_key:
                message = message.replace(self.api_key, '[API Key]')
            raise ApiClientError(message, exc.code) from exc
        except (URLError, TimeoutError, OSError) as exc:
            raise ApiClientError('API 연결 실패 또는 시간 초과. 서버 주소와 연결을 확인하세요.') from exc
        except (ValueError, UnicodeError) as exc:
            raise ApiClientError('API 응답이 올바른 JSON 객체가 아닙니다.') from exc

    def health(self):
        return self.request('GET', '/health')

    def ready(self):
        return self.request('GET', '/ready')


class ApiJobWorker(QObject):
    progress = pyqtSignal(dict)
    succeeded = pyqtSignal(dict)
    failed = pyqtSignal(str)
    finished = pyqtSignal()

    def __init__(self, client: FolderGrantApiClient, kind: str, payload: dict, dry_run: bool = False, retry_job_id: str = None, validate_only: bool = False):
        super().__init__()
        if kind not in {'access', 'project'}:
            raise ValueError('알 수 없는 API 작업 종류입니다.')
        self.client, self.kind, self.payload, self.dry_run = client, kind, dict(payload), dry_run
        self.poll_interval = 1.5
        self.max_wait_seconds = 4 * 60 * 60
        self._stop_event = threading.Event()
        self.job_id = None
        self.retry_job_id = retry_job_id
        self.validate_only = validate_only

    def stop(self) -> None:
        self._stop_event.set()

    def _verify_recovered_job(self, job: dict, preview: dict) -> None:
        """Only recover a BUS actor conflict when the server execution plan matches."""
        if self.payload['request_id'] not in {job.get('request_id'),job.get('original_request_id')}:
            raise ApiClientError('기존 Job의 요청번호가 일치하지 않습니다.',409)
        if self.kind == 'access':
            if any(job.get(k) != self.payload.get(k) for k in ('operation','project_status')):
                raise ApiClientError('기존 Job의 권한 작업이 일치하지 않습니다.',409)
        elif any(job.get(k) != self.payload.get(k) for k in ('project_code','project_name')):
            raise ApiClientError('기존 Job의 프로젝트 입력이 일치하지 않습니다.',409)
        expected = []
        for step in preview.get('steps',[]):
            details = dict(step.get('metadata',{}))
            if self.kind == 'access':
                details['permission'] = step.get('permission')
            expected.append((step.get('type'),step.get('target'),details))
        actual = [(step.get('type'),step.get('target'),step.get('details',{}))
                  for step in job.get('steps',[])]
        if not expected or expected != actual:
            raise ApiClientError('기존 Job과 현재 Preview의 실행 계획이 다릅니다. 요청번호 충돌을 확인하세요.',409)

    def execute(self) -> dict:
        """Also used by the project batch worker in its own Qt thread."""
        path = '/api/v1/' + self.kind + '-jobs'
        if self._stop_event.is_set():
            if self.validate_only:
                return dict(status='validated',validation_status='unknown',request_id=self.payload['request_id'],client_stopped=True,
                            checks=[dict(id='cancelled',label='실제 자원 검증',target='',status='unknown',message='요청 검증을 중지했습니다.')])
            return {'status': 'cancelled', 'client_stopped': True}
        if self.validate_only:
            try:
                validation = self.client.request('POST', path + '/validate', self.payload)
            except ApiClientError as exc:
                if exc.status_code == 404:
                    raise ApiClientError('요청 검증 API가 없습니다. API 서버를 업데이트하고 재시작하세요.', 404) from exc
                raise
            if (validation.get('request_id') != self.payload['request_id']
                    or validation.get('validation_status') not in {'passed','failed','warning','unknown'}
                    or not isinstance(validation.get('checks'), list) or not validation['checks']
                    or any(not isinstance(check, dict) or check.get('status') not in
                           {'passed','failed','warning','unknown','not_applicable'} for check in validation['checks'])):
                raise ApiClientError('요청 검증 응답이 올바르지 않습니다. 실제 검증 결과를 확인하지 못했습니다.')
            states = {check['status'] for check in validation['checks']}
            expected = 'failed' if 'failed' in states else 'unknown' if 'unknown' in states else 'warning' if 'warning' in states else 'passed'
            if validation['validation_status'] != expected:
                raise ApiClientError('요청 검증 요약과 항목별 결과가 일치하지 않습니다.')
            return {**validation, 'status':'validated', 'client_stopped':self._stop_event.is_set()}
        if self.dry_run:
            preview = self.client.request('POST', path + '/preview', self.payload)
            return {**preview, 'status': 'preview', 'client_stopped': self._stop_event.is_set()}
        try:
            if self.retry_job_id:
                job = self.client.request('POST',path+'/'+quote(self.retry_job_id,safe='')+'/retry',
                                          {'requested_by':self.payload['requested_by']})
            else:
                job = self.client.request('POST', path, self.payload)
        except ApiClientError as exc:
            if self.retry_job_id and exc.status_code==404:
                raise ApiClientError('재실행 API를 사용할 수 없습니다. API 서버를 최신 버전으로 업데이트하고 재시작하세요.',404) from exc
            bus_conflict = not self.retry_job_id and exc.status_code == 409 and self.payload.get('source') == 'bus'
            if exc.status_code is not None and exc.status_code < 500 and not bus_conflict:
                raise
            # An accepted POST can lose its response. Never generate a new ID here.
            try:
                job = self.client.request('GET', path + '/by-request/' + quote(self.payload['request_id'], safe=''))
                if self.retry_job_id and job.get('job_id')==self.retry_job_id:
                    raise ApiClientError('재실행 접수 여부를 확인하지 못했습니다. 같은 항목으로 다시 시도하세요.')
                if bus_conflict:
                    preview = self.client.request('POST',path+'/preview',self.payload)
                    self._verify_recovered_job(job,preview)
                    job = {**job,'replayed':True,'recovered_after_actor_conflict':True}
            except ApiClientError as recovery:
                raise ApiClientError(f'등록 결과를 확인하지 못했습니다. 같은 요청으로 재시도하세요. request_id={self.payload["request_id"]}: {recovery}') from exc
        self.job_id = job.get('job_id')
        if not self.job_id:
            raise ApiClientError('API 응답에 job_id가 없습니다.')
        job_path = path + '/' + quote(str(self.job_id), safe='')
        deadline = time.monotonic() + self.max_wait_seconds
        cancel_sent = False
        while True:
            self.progress.emit(dict(job))
            status = job.get('status')
            if status in FINAL_JOB_STATUSES:
                return {**job, 'client_stopped': self._stop_event.is_set()}
            if status not in {'queued', 'running', 'cancel_requested'}:
                raise ApiClientError('알 수 없는 API Job 상태: ' + str(status))
            if time.monotonic() >= deadline:
                raise ApiClientError(f'작업 조회 대기 시간이 초과되었습니다. job_id={self.job_id}')
            if self._stop_event.is_set() and not cancel_sent:
                try:
                    self.client.request('POST', job_path + '/cancel', {'requested_by': self.payload['requested_by']})
                except ApiClientError as exc:
                    if exc.status_code != 409:
                        raise
                    job = self.client.request('GET',job_path)
                    if job.get('status') not in FINAL_JOB_STATUSES:
                        raise
                    continue
                cancel_sent = True
            # Once cancelled, wait normally to avoid busy looping on a set Event.
            if cancel_sent:
                threading.Event().wait(self.poll_interval)
            else:
                self._stop_event.wait(self.poll_interval)
            job = self.client.request('GET', job_path)

    @pyqtSlot()
    def run(self) -> None:
        try:
            self.succeeded.emit(self.execute())
        except Exception as exc:
            message = str(exc)
            if self.client.api_key:
                message = message.replace(self.client.api_key, '[API Key]')
            self.failed.emit(message)
        finally:
            self.finished.emit()


def load_api_config() -> dict:
    try:
        with open(API_CONFIG_FILE,encoding='utf-8') as handle:
            data = json.load(handle)
        if isinstance(data,dict):
            return dict(base_url=data.get('base_url',DEFAULT_API_URL),api_key=data.get('api_key',''))
    except (OSError,ValueError):
        pass
    return dict(base_url=DEFAULT_API_URL,api_key='')


class ApiJobRefreshWorker(QObject):
    refreshed = pyqtSignal(int,dict)
    failed = pyqtSignal(int,str)
    finished = pyqtSignal()

    def __init__(self,client,requests):
        super().__init__()
        self.client = client
        self.requests = list(requests)
        self._stop_event = threading.Event()

    def stop(self):
        self._stop_event.set()

    @pyqtSlot()
    def run(self):
        try:
            for row,request_id in self.requests:
                if self._stop_event.is_set():
                    break
                try:
                    result = self.client.request('GET','/api/v1/access-jobs/by-request/'+quote(request_id,safe=''))
                    if request_id not in {result.get('request_id'),result.get('original_request_id')}:
                        raise ApiClientError('서버 Job의 요청번호가 현재 행과 일치하지 않습니다.')
                    if not self._stop_event.is_set():
                        self.refreshed.emit(row,result)
                except Exception as exc:
                    message = str(exc)
                    if self.client.api_key:
                        message = message.replace(self.client.api_key,'[API Key]')
                    self.failed.emit(row,message)
        finally:
            self.finished.emit()


class ApiConnectionTask(QObject):
    succeeded = pyqtSignal(dict)
    failed = pyqtSignal(str)
    finished = pyqtSignal()

    def __init__(self,client):
        super().__init__()
        self.client = client

    @pyqtSlot()
    def run(self):
        try:
            self.succeeded.emit(dict(health=self.client.health(),ready=self.client.ready()))
        except Exception as exc:
            self.failed.emit(str(exc))
        finally:
            self.finished.emit()


class ApiSettingsDialog(QDialog):
    def __init__(self,config,parent=None):
        super().__init__(parent)
        self.setWindowTitle('API 연결 설정')
        self.resize(500,220)
        self.connection_verified = False
        self.thread = None
        self.task = None
        layout = QFormLayout(self)
        self.url = QLineEdit(config.get('base_url',DEFAULT_API_URL))
        self.key = QLineEdit(config.get('api_key',''))
        self.key.setEchoMode(QLineEdit.Password)
        self.message = QLabel('API Key는 C:\\FGT\\conf\\api_client.json에 저장됩니다.')
        self.message.setWordWrap(True)
        layout.addRow('API 주소',self.url)
        layout.addRow('X-API-Key',self.key)
        self.check = QPushButton('연결 확인 (health / ready)')
        self.check.clicked.connect(self._check_connection)
        layout.addRow(self.check)
        layout.addRow(self.message)
        self.buttons = QDialogButtonBox(QDialogButtonBox.Save|QDialogButtonBox.Cancel)
        self.buttons.accepted.connect(self._save)
        self.buttons.rejected.connect(self.reject)
        layout.addRow(self.buttons)
        self.url.textChanged.connect(lambda text:setattr(self,'connection_verified',False))
        self.key.textChanged.connect(lambda text:setattr(self,'connection_verified',False))

    def values(self):
        return dict(base_url=self.url.text().strip().rstrip('/'),api_key=self.key.text().strip())

    def _validated_client(self):
        values = self.values()
        if not values['api_key']:
            raise ValueError('API Key를 입력하세요.')
        return FolderGrantApiClient(values['base_url'],values['api_key'])

    def _check_connection(self):
        if self.thread is not None:
            return
        try:
            client = self._validated_client()
        except ValueError as exc:
            self.message.setText(str(exc)); return
        self.thread = QThread(self)
        self.task = ApiConnectionTask(client)
        self.task.moveToThread(self.thread)
        self.thread.started.connect(self.task.run)
        self.task.succeeded.connect(self._connection_succeeded)
        self.task.failed.connect(self.message.setText)
        self.task.finished.connect(self.task.deleteLater)
        self.task.finished.connect(self.thread.quit,type=Qt.DirectConnection)
        self.thread.finished.connect(self._connection_finished)
        self.thread.finished.connect(self.thread.deleteLater)
        self.check.setEnabled(False)
        self.buttons.setEnabled(False)
        self.url.setEnabled(False); self.key.setEnabled(False)
        self.connection_verified = False
        self.message.setText('연결 확인 중…')
        self.thread.start()

    def _connection_succeeded(self,result):
        self.connection_verified = True
        self.message.setText('API 연결 및 인증 확인 성공')

    def _connection_finished(self):
        self.thread = None; self.task = None
        self.check.setEnabled(True); self.buttons.setEnabled(True)
        self.url.setEnabled(True); self.key.setEnabled(True)

    def _save(self):
        if self.thread is not None:
            return
        try:
            self._validated_client()
            os.makedirs(os.path.dirname(API_CONFIG_FILE),exist_ok=True)
            temporary = API_CONFIG_FILE+'.tmp'
            with open(temporary,'w',encoding='utf-8') as handle:
                json.dump(self.values(),handle,ensure_ascii=False,indent=2)
            if os.name != 'nt':
                os.chmod(temporary,0o600)
            os.replace(temporary,API_CONFIG_FILE)
        except (OSError,ValueError) as exc:
            self.message.setText(str(exc)); return
        self.accept()

    def reject(self):
        if self.thread is None:
            super().reject()

    def closeEvent(self,event):
        if self.thread is not None:
            event.ignore()
        else:
            event.accept()


class SettingsDialog(QDialog):
    def __init__(self, parent=None, saved=None, remembered=False, debug_on=False, fail_tol_default=5):
        try:
            self._fail_tol_default = int(fail_tol_default)
        except Exception:
            self._fail_tol_default = 5

        super().__init__(parent)
        self.setWindowTitle("로그인 계정")
        self.le_id = QLineEdit()
        self.le_pw = QLineEdit()
        self.le_pw.setEchoMode(QLineEdit.Password)
        if saved:
            self.le_id.setText(saved.get("id",""))
            self.le_pw.setText(saved.get("pw",""))

        self.chk_save = QCheckBox("저장")
        self.chk_save.setChecked(bool(remembered))

        self.chk_debug = QCheckBox("디버그 모드")
        self.chk_debug.setChecked(bool(debug_on))

        self.sb_tol = QSpinBox()
        self.sb_tol.setRange(0, 999)
        try:
            tol_saved = int((saved or {}).get("remove_fail_tol", self._fail_tol_default))
            self.sb_tol.setValue(tol_saved)
        except Exception:
            self.sb_tol.setValue(self._fail_tol_default)

        form = QFormLayout()
        self._saved_notify_refresh_min = (saved or {}).get("notify_refresh_min",10)

        # self.sb_notify = QSpinBox()
        # self.sb_notify.setRange(5, 60)
        # self.sb_notify.setValue(int((saved or {}).get("notify_refresh_min", 10)))
        # form.addRow("알림 리프레시(분)", self.sb_notify)

        form.addRow("아이디", self.le_id)
        form.addRow("비밀번호", self.le_pw)
        form.addRow("", self.chk_save)
        form.addRow("", self.chk_debug)
        form.addRow("폴더해제 실패 허용 개수", self.sb_tol)

        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.accepted.connect(self.accept)
        btns.rejected.connect(self.reject)

        lay = QVBoxLayout(self)
        lay.addLayout(form)

        bottom_bar = QHBoxLayout()
        bottom_bar.setContentsMargins(0, 6, 0, 0)

        self.version_lbl = QLabel(f"{APP_VERSION_STR}")
        self.version_lbl.setObjectName("versionLabel")
        self.version_lbl.setStyleSheet("color:#8a8a8a; font-size:11px; padding:2px 4px;")
        self.version_lbl.setToolTip(f"Build: {APP_BUILD}")

        bottom_bar.addWidget(self.version_lbl, 0, Qt.AlignVCenter)
        bottom_bar.addStretch()
        bottom_bar.addWidget(btns)

        lay.addLayout(bottom_bar)

    def result(self):
        return (
            self.le_id.text().strip(),
            self.le_pw.text().strip(),
            self.chk_save.isChecked(),
            self.chk_debug.isChecked(),
            int(self.sb_tol.value()),
            # int(self.sb_notify.value()),  # 알림 기능 비활성화: 설정값만 보존
            int(self._saved_notify_refresh_min),
        )

    def get_fail_tol(self) -> int:
        try:
            return int(self.sb_tol.value())
        except Exception:
            return 5

class BusSessionManager(QObject):
    stopped = pyqtSignal()
    readyChanged = pyqtSignal(bool, str)
    downloaded = pyqtSignal(str, str)
    busyChanged = pyqtSignal(bool)
    processed = pyqtSignal(object)
    newDownloaded = pyqtSignal(str, str)
    countsReady = pyqtSignal(dict)

    MAX_INIT_RETRY = 3
    INIT_RETRY_BASE_DELAY = 2.0

    def __init__(self, dl_dir: str):
        super().__init__()
        self.dl_dir = dl_dir
        self.driver = None
        self._ready = False
        self._busy = False
        self._initializing = False
        self._authenticated_credentials = None
        self._cancel = False
        self.debug_enabled = False
        self.debug_dir = DEBUG_DIR
        self._user = ""
        self._pw = ""

    def set_debug(self, enabled: bool, debug_dir: str | None = None):
        self.debug_enabled = bool(enabled)
        if debug_dir:
            self.debug_dir = debug_dir
        try:
            if self.debug_enabled:
                os.makedirs(self.debug_dir, exist_ok=True)
        except Exception:
            pass

    def _stabilize_grid(self, tries=2, nap=0.8):
        self._wait_overlay_gone(10)
        for _ in range(tries):
            self._click_search()
            self._wait_overlay_gone(10)
            self._go_iframe()
            time.sleep(nap)

    def set_creds(self, user: str, pw: str):
        self._user, self._pw = (user or "").strip(), (pw or "").strip()

    def is_ready(self) -> bool:
        return (self._ready and self.driver is not None
                and self._authenticated_credentials == (self._user,self._pw))

    def is_busy(self) -> bool:
        return self._busy

    def _reset_filters(self):
        d = self.driver
        try:
            self._go_iframe()
            d.execute_script("""
              for (const id of ['approverYn','processYn','releaseYn']) {
                const el = document.getElementById(id);
                if (el) { el.value=''; el.dispatchEvent(new Event('change',{bubbles:true})); }
              }
            """)
        except Exception:
            pass

    def _nav_open_and_iframe(self, url: str, iframe_css: str = "iframe"):
        d = self.driver
        d.switch_to.default_content()
        d.get(url)
        WebDriverWait(d, 10).until(EC.presence_of_element_located((By.CSS_SELECTOR, iframe_css)))
        frame = d.find_element(By.CSS_SELECTOR, iframe_css)
        d.switch_to.frame(frame)

    def _set_select_value_and_fire(self, select_id: str, value: str):
        d = self.driver
        d.execute_script("""
            var sid = arguments[0], val = arguments[1];
            var s = document.getElementById(sid);
            if (s) {
                s.value = val;
                try { s.dispatchEvent(new Event('change', {bubbles:true})); } catch(e) {}
                if (typeof AllBtnYn === 'function') { try { AllBtnYn(); } catch(e) {} }
            }
        """, select_id, value)

    def _click_search_manual(self):
        d = self.driver
        try:
            WebDriverWait(d, 2).until(EC.element_to_be_clickable((By.ID, "btnSearch"))).click()
            return
        except Exception:
            pass

        d.switch_to.default_content()
        WebDriverWait(d, 5).until(EC.element_to_be_clickable((By.ID, "btnSearch"))).click()

        WebDriverWait(d, 5).until(EC.presence_of_element_located((By.CSS_SELECTOR, "iframe")))
        d.switch_to.frame(d.find_element(By.CSS_SELECTOR, "iframe"))

    def _goto_new_site(self):
        try:
            d = self.driver
            d.switch_to.default_content()
            d.get(BUS_NEW_URL)
            WebDriverWait(d, 20).until(EC.presence_of_all_elements_located((By.TAG_NAME, "iframe")))
            self._go_iframe()
            self._wait_overlay_gone(10)
            time.sleep(0.8)
            return True
        except Exception:
            return False

    def _goto_progress_site(self):
        try:
            d = self.driver
            d.switch_to.default_content()
            d.get(BUS_PROGRESS_LIST_URL)
            WebDriverWait(d, 20).until(EC.presence_of_all_elements_located((By.TAG_NAME, "iframe")))
            self._go_iframe()
            self._wait_overlay_gone(10)
            time.sleep(1.0)
            return True
        except Exception:
            return False

    def _goto_end_site(self):
        try:
            d = self.driver
            d.switch_to.default_content()
            d.get(BUS_END_LIST_URL)
            WebDriverWait(d, 20).until(EC.presence_of_all_elements_located((By.TAG_NAME, "iframe")))
            self._go_iframe()
            self._wait_overlay_gone(10)
            time.sleep(1.0)
            return True
        except Exception:
            return False

    @pyqtSlot()
    def download_new_list(self):
        if not self.is_ready():
            self.newDownloaded.emit("", "세션 준비 안됨"); return
        if self._busy:
            self.newDownloaded.emit("", "다른 작업 실행중"); return

        self._busy = True
        self.busyChanged.emit(True)
        self._cancel = False
        d = self.driver

        try:
            self._goto_new_site()
            d = self.driver
            self._go_iframe()
            try:
                d.execute_script("""
                  const p=document.getElementById('processYn');
                  if(p){ p.value='N'; p.dispatchEvent(new Event('change',{bubbles:true})); }
                """)
                time.sleep(0.2)
                self._go_iframe()
                WebDriverWait(d, 5).until(EC.element_to_be_clickable((By.ID, "btnSearch"))).click()
            except Exception:
                pass

            before = set(glob.glob(os.path.join(self.dl_dir, "*.xls"))) | set(glob.glob(os.path.join(self.dl_dir, "*.xlsx")))

            WebDriverWait(d, 10).until(EC.element_to_be_clickable((By.ID, "btnExcel"))).click()
            try:
                WebDriverWait(d, 8).until(EC.visibility_of_element_located((By.ID, "txtExcelDownReason")))
                try:
                    el = d.find_element(By.ID, "txtExcelDownReason")
                    el.clear(); el.send_keys("신규확인-미완료(N) 목록")
                except Exception:
                    pass
                ok_btn = WebDriverWait(d, 6).until(EC.element_to_be_clickable((By.CSS_SELECTOR, "button.swal2-confirm")))
                try: ok_btn.click()
                except ElementClickInterceptedException:
                    d.execute_script("arguments[0].click();", ok_btn)
            except TimeoutException:
                pass

            end = time.time() + 120
            latest = ""
            while time.time() < end:
                self._check_cancel()
                if not list(glob.glob(os.path.join(self.dl_dir, "*.crdownload"))):
                    now = set(glob.glob(os.path.join(self.dl_dir, "*.xls"))) | set(glob.glob(os.path.join(self.dl_dir, "*.xlsx")))
                    new_files = list(now - before)
                    if new_files:
                        latest = max(new_files, key=os.path.getmtime)
                        break
                time.sleep(0.3)

            if not latest:
                self.newDownloaded.emit("", "다운로드 실패"); return

            out_path = os.path.join(self.dl_dir, "신규_미완료.xls")
            try:
                if os.path.exists(out_path): os.remove(out_path)
            except Exception:
                pass
            try:
                shutil.move(latest, out_path)
            except Exception:
                try:
                    shutil.copyfile(latest, out_path)
                    try: os.remove(latest)
                    except Exception: pass
                except Exception:
                    self.newDownloaded.emit("", "파일 저장 실패"); return

            self.newDownloaded.emit(out_path, "")
        except SystemExit:
            self.newDownloaded.emit("", "사용자 취소")
        except Exception as e:
            self.newDownloaded.emit("", f"오류: {e}")
        finally:
            self._busy = False
            self.busyChanged.emit(False)

    def _read_total_entries(self) -> int:
        d = self.driver
        try:
            info_el = WebDriverWait(d, 5).until(
                EC.presence_of_element_located((By.CSS_SELECTOR, ".dataTables_info"))
            )
            txt = (info_el.text or "").strip().lower()
            m = re.search(r"(\d+)", txt)
            if m:
                return int(m.group(1))
        except Exception:
            pass

        try:
            self._stabilize_grid()
            WebDriverWait(d, 5).until(EC.presence_of_all_elements_located((By.CSS_SELECTOR, "tbody tr")))
            return len(d.find_elements(By.CSS_SELECTOR, "tbody tr"))
        except Exception:
            return 0

    def collect_request_counts(self):
        if not self.is_ready():
            try:
                self.countsReady.emit(getattr(self, "_last_counts", {}))
            except Exception:
                pass
            return
        if getattr(self, "_busy", False):

            try:
                self.countsReady.emit(getattr(self, "_last_counts", {}))
            except Exception:
                pass
            return

        d = self.driver
        counts = {
            "신규미완료": 0,
            "진행-부여": 0,
            "진행-제거": 0,
            "종료-부여": 0,
            "종료-제거": 0,
        }

        try:
            try:
                self._nav_open_and_iframe(BUS_NEW_URL)

                try:
                    pre = d.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                except Exception:
                    pre = ""

                self._set_select_value_and_fire("processYn", "N")
                self._click_search_manual()

                try:
                    WebDriverWait(d, 7).until(
                        lambda x: (x.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                                   if x.find_elements(By.CSS_SELECTOR, ".dataTables_info") else "") != (pre or "")
                    )
                except Exception:
                    time.sleep(0.6)

                txt = d.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip().lower()
                m = re.search(r"(\d+)", txt)
                counts["신규미완료"] = int(m.group(1)) if m else 0
            except Exception:
                pass

            try:
                self._nav_open_and_iframe(BUS_PROGRESS_LIST_URL)

                try:
                    pre = d.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                except Exception:
                    pre = ""

                self._set_select_value_and_fire("processYn", "N")
                self._click_search_manual()

                try:
                    WebDriverWait(d, 7).until(
                        lambda x: (x.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                                   if x.find_elements(By.CSS_SELECTOR, ".dataTables_info") else "") != (pre or "")
                    )
                except Exception:
                    time.sleep(0.6)

                txt = d.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip().lower()
                m = re.search(r"(\d+)", txt)
                counts["진행-부여"] = int(m.group(1)) if m else 0
            except Exception:
                pass

            try:
                self._nav_open_and_iframe(BUS_PROGRESS_LIST_URL)

                try:
                    pre = d.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                except Exception:
                    pre = ""

                self._set_select_value_and_fire("releaseYn", "N")
                self._click_search_manual()

                try:
                    WebDriverWait(d, 7).until(
                        lambda x: (x.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                                   if x.find_elements(By.CSS_SELECTOR, ".dataTables_info") else "") != (pre or "")
                    )
                except Exception:
                    time.sleep(0.6)

                txt = d.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip().lower()
                m = re.search(r"(\d+)", txt)
                counts["진행-제거"] = int(m.group(1)) if m else 0
            except Exception:
                pass

            try:
                self._nav_open_and_iframe(BUS_END_LIST_URL)

                try:
                    pre = d.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                except Exception:
                    pre = ""

                self._set_select_value_and_fire("processYn", "N")
                self._click_search_manual()

                try:
                    WebDriverWait(d, 7).until(
                        lambda x: (x.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                                   if x.find_elements(By.CSS_SELECTOR, ".dataTables_info") else "") != (pre or "")
                    )
                except Exception:
                    time.sleep(0.6)

                txt = d.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip().lower()
                m = re.search(r"(\d+)", txt)
                counts["종료-부여"] = int(m.group(1)) if m else 0
            except Exception:
                pass

            try:
                self._nav_open_and_iframe(BUS_END_LIST_URL)

                try:
                    pre = d.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                except Exception:
                    pre = ""

                self._set_select_value_and_fire("releaseYn", "N")
                self._click_search_manual()

                try:
                    WebDriverWait(d, 7).until(
                        lambda x: (x.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                                   if x.find_elements(By.CSS_SELECTOR, ".dataTables_info") else "") != (pre or "")
                    )
                except Exception:
                    time.sleep(0.6)

                txt = d.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip().lower()
                m = re.search(r"(\d+)", txt)
                counts["종료-제거"] = int(m.group(1)) if m else 0
            except Exception:
                pass

        finally:
            self._last_counts = counts
            try:
                self.countsReady.emit(counts)
            except Exception:
                pass

    @pyqtSlot()
    def start(self):
        if self._initializing or getattr(self,'_stopping',False):
            return
        self._initializing = True
        self._busy = True; self.busyChanged.emit(True)
        self._cancel = False
        try:
            self._start_session()
        finally:
            self._initializing = False
            self._busy = False; self.busyChanged.emit(False)

    def _start_session(self):
        def _cleanup_driver(drv):
            if not drv:
                return
            try:
                drv.quit()
            except Exception:
                try:
                    if hasattr(drv, "service") and drv.service and drv.service.process:
                        drv.service.process.kill()
                except Exception:
                    pass

        if not (self._user and self._pw):
            self._ready = False
            self.readyChanged.emit(False, "자격증명 없음")
            return

        username,password = self._user,self._pw
        if self.driver:
            if self.is_ready():
                self.readyChanged.emit(True,'이미 준비됨')
                return
            _cleanup_driver(self.driver)
            self.driver = None; self._ready = False
            self._authenticated_credentials = None

        attempt = 0
        last_err = None

        while attempt < self.MAX_INIT_RETRY and not self._cancel:
            attempt += 1
            try:
                options = make_hidden_chrome_options(self.dl_dir)

                self.driver = webdriver.Chrome(options=options)
                d = self.driver

                d.get(BUS_LOGIN_URL)
                WebDriverWait(d, 20).until(EC.presence_of_element_located((By.ID, "windowsaccount")))
                d.find_element(By.ID, "windowsaccount").send_keys(username)
                d.find_element(By.ID, "password").send_keys(password)
                d.find_element(By.ID, "btnLogin").click()
                WebDriverWait(d, 20).until(EC.url_contains("Common"))

                try:
                    d.switch_to.window(d.window_handles[-1])
                except Exception:
                    pass

                d.get(BUS_PROGRESS_LIST_URL)
                WebDriverWait(d, 20).until(EC.presence_of_all_elements_located((By.TAG_NAME, "iframe")))
                self._go_iframe()

                WebDriverWait(d, 10).until(EC.element_to_be_clickable((By.ID, "approverYn")))
                Select(d.find_element(By.ID, "approverYn")).select_by_value("1")

                self._set_request_filter(REQ_GRANT)

                if (username,password)!=(self._user,self._pw):
                    _cleanup_driver(self.driver); self.driver = None; self._ready = False
                    self._authenticated_credentials = None
                    self.readyChanged.emit(False,'로그인 중 계정이 변경되었습니다. 다시 로그인하세요.')
                    return
                self._authenticated_credentials = (username,password)
                self._ready = True
                self.readyChanged.emit(True, f"준비됨(시도 {attempt}/{self.MAX_INIT_RETRY})")
                return

            except Exception as e:
                last_err = e
                _cleanup_driver(self.driver)
                self.driver = None
                self._ready = False
                # Intermediate retries must not release login callbacks or enable another login.
                self._authenticated_credentials = None

                if attempt >= self.MAX_INIT_RETRY or self._cancel:
                    break

                delay = self.INIT_RETRY_BASE_DELAY * (2 ** (attempt - 1))
                end = time.time() + delay
                while time.time() < end and not self._cancel:
                    QApplication.processEvents()
                    time.sleep(0.05)

        self._ready = False
        self.readyChanged.emit(False, f"세션 초기화 최종 실패: {last_err}")

    @pyqtSlot()
    def stop(self):
        if getattr(self, "_stopping", False):
            return
        self._stopping = True
        try:
            self._cancel = True
            self._busy = False
            self._ready = False
            self.busyChanged.emit(False)
            self.readyChanged.emit(False, "정지됨")

            d = self.driver
            self.driver = None
            if d:
                try:
                    d.quit()
                except Exception:
                    try:
                        if hasattr(d, "service") and d.service and d.service.process:
                            d.service.process.kill()
                    except Exception:
                        pass
        finally:
            self._stopping = False
            self.stopped.emit()

    @pyqtSlot()
    def cancel_current(self):
        self._cancel = True

    def _check_cancel(self, context: str = "download"):
        if self._cancel:
            self._busy = False
            self.busyChanged.emit(False)
            self._cancel = False
            if context == "process":
                self.processed.emit([])
            else:
                self.downloaded.emit("", "사용자 취소")
            raise SystemExit

    @pyqtSlot()
    def download_list(self):
        if not self.is_ready():
            self.downloaded.emit("", "세션 준비 안됨")
            return
        if self._busy:
            self.downloaded.emit("", "다른 작업 실행중")
            return

        self._busy = True
        self.busyChanged.emit(True)
        self._cancel = False

        d = self.driver

        def _search_and_download_end(req_flag: str) -> str:
            self._goto_end_site()
            self._set_end_filter(req_flag)

            d = self.driver
            try:
                self._go_iframe()
                try:
                    WebDriverWait(d, 5).until(EC.element_to_be_clickable((By.ID, "btnSearch"))).click()
                except Exception:
                    d.switch_to.default_content()
                    WebDriverWait(d, 5).until(EC.element_to_be_clickable((By.ID, "btnSearch"))).click()
                    self._go_iframe()
            except Exception:
                pass

            before = set(glob.glob(os.path.join(self.dl_dir, "*.xls"))) | set(glob.glob(os.path.join(self.dl_dir, "*.xlsx")))

            try:
                self._go_iframe()
                WebDriverWait(d, 10).until(EC.element_to_be_clickable((By.ID, "btnExcel"))).click()
            except Exception:
                d.switch_to.default_content()
                WebDriverWait(d, 10).until(EC.element_to_be_clickable((By.ID, "btnExcel"))).click()

            try:
                WebDriverWait(d, 10).until(EC.visibility_of_element_located((By.ID, "txtExcelDownReason")))
                try:
                    reason = "종료권한리스트 수집-" + ("해제" if req_flag == END_FLAG_RELEASE else "부여")
                    el = d.find_element(By.ID, "txtExcelDownReason")
                    el.clear(); el.send_keys(reason)
                except Exception:
                    pass
                ok_btn = WebDriverWait(d, 10).until(EC.element_to_be_clickable((By.CSS_SELECTOR, "button.swal2-confirm")))
                try: ok_btn.click()
                except ElementClickInterceptedException:
                    d.execute_script("arguments[0].click();", ok_btn)
            except TimeoutException:
                pass

            end = time.time() + 120
            latest = ""
            while time.time() < end:
                self._check_cancel()
                if not list(glob.glob(os.path.join(self.dl_dir, "*.crdownload"))):
                    now = set(glob.glob(os.path.join(self.dl_dir, "*.xls"))) | set(glob.glob(os.path.join(self.dl_dir, "*.xlsx")))
                    new_files = list(now - before)
                    if new_files:
                        latest = max(new_files, key=os.path.getmtime)
                        break
                time.sleep(0.35)

            if not latest:
                return ""

            suffix = "해제" if req_flag == END_FLAG_RELEASE else "부여"
            out_path = os.path.join(self.dl_dir, f"종료권한리스트_{suffix}.xls")
            try:
                if os.path.exists(out_path): os.remove(out_path)
            except Exception:
                pass

            try:
                shutil.move(latest, out_path)
            except Exception:
                try:
                    shutil.copyfile(latest, out_path)
                    try: os.remove(latest)
                    except Exception: pass
                except Exception:
                    return ""

            return out_path

        def _search_and_download(reqtype: str) -> str:
            self._goto_progress_site()
            self._set_request_filter(reqtype)
            try:
                self._go_iframe()
                try:
                    WebDriverWait(d, 5).until(EC.element_to_be_clickable((By.ID, "btnSearch"))).click()
                except Exception:
                    d.switch_to.default_content()
                    WebDriverWait(d, 5).until(EC.element_to_be_clickable((By.ID, "btnSearch"))).click()
                    self._go_iframe()
            except Exception:
                pass

            before = set(glob.glob(os.path.join(self.dl_dir, "*.xls"))) | set(
                glob.glob(os.path.join(self.dl_dir, "*.xlsx"))
            )

            WebDriverWait(d, 10).until(EC.element_to_be_clickable((By.ID, "btnExcel"))).click()
            WebDriverWait(d, 10).until(EC.visibility_of_element_located((By.ID, "txtExcelDownReason")))
            try:
                reason = f"권한리스트 수집-{'해제' if reqtype == REQ_RELEASE else '부여'}"
                el = d.find_element(By.ID, "txtExcelDownReason")
                el.clear()
                el.send_keys(reason)
            except Exception:
                pass

            ok_btn = WebDriverWait(d, 10).until(
                EC.element_to_be_clickable((By.CSS_SELECTOR, "button.swal2-confirm"))
            )
            try:
                ok_btn.click()
            except ElementClickInterceptedException:
                d.execute_script("arguments[0].click();", ok_btn)

            end = time.time() + 120
            latest = ""
            while time.time() < end:
                self._check_cancel()
                if not list(glob.glob(os.path.join(self.dl_dir, "*.crdownload"))):
                    now = set(glob.glob(os.path.join(self.dl_dir, "*.xls"))) | set(
                        glob.glob(os.path.join(self.dl_dir, "*.xlsx"))
                    )
                    new_files = list(now - before)
                    if new_files:
                        latest = max(new_files, key=os.path.getmtime)
                        break
                time.sleep(0.35)

            if not latest:
                return ""

            suffix = "해제" if reqtype == REQ_RELEASE else "부여"
            out_path = os.path.join(self.dl_dir, f"권한리스트_{suffix}.xls")
            try:
                if os.path.exists(out_path):
                    os.remove(out_path)
            except Exception:
                pass

            try:
                shutil.move(latest, out_path)
            except Exception:
                try:
                    shutil.copyfile(latest, out_path)
                    try:
                        os.remove(latest)
                    except Exception:
                        pass
                except Exception:
                    return ""

            return out_path

        try:
            for f in glob.glob(os.path.join(self.dl_dir, "*")):
                try:
                    os.remove(f)
                except:
                    pass

            grant_path = _search_and_download(REQ_GRANT)
            release_path = _search_and_download(REQ_RELEASE)

            end_grant_path   = _search_and_download_end(END_FLAG_GRANT)
            end_release_path = _search_and_download_end(END_FLAG_RELEASE)

            header_g, data_g = _parse_html_best_table(grant_path) if grant_path else ([], [])
            header_r, data_r = _parse_html_best_table(release_path) if release_path else ([], [])
            has_normal = bool(grant_path or release_path)
            combined_path = ""
            if has_normal:
                header = header_g if len(header_g) >= len(header_r) else header_r
                def _pad(row, n): rr = list(row);  rr += [""] * max(0, n-len(rr)); return rr[:n]
                rows = []
                for r in data_g: rows.append(_pad(r, len(header)))
                for r in data_r: rows.append(_pad(r, len(header)))
                parts = ["<html><head><meta charset='utf-8'></head><body>",
                         "<table border='1'>",
                         "<thead><tr>" + "".join(f"<th>{htmllib.escape(str(h or ''))}</th>" for h in header) + "</tr></thead>",
                         "<tbody>"]
                for r in rows:
                    parts.append("<tr>" + "".join(f"<td>{htmllib.escape(str(c or ''))}</td>" for c in r) + "</tr>")
                parts.append("</tbody></table></body></html>")
                combined_path = os.path.join(self.dl_dir, "권한리스트_합본.xls")
                with open(combined_path, "w", encoding="utf-8") as f:
                    f.write("\n".join(parts))

            header_eg, data_eg = _parse_html_best_table(end_grant_path) if end_grant_path else ([], [])
            header_er, data_er = _parse_html_best_table(end_release_path) if end_release_path else ([], [])
            has_end = bool(end_grant_path or end_release_path)
            combined_end_path = ""
            if has_end:
                header_e = header_eg if len(header_eg) >= len(header_er) else header_er
                def _pad_e(row, n):
                    rr = list(row); rr += [""] * max(0, n-len(rr)); return rr[:n]
                rows_e = []
                for r in data_eg: rows_e.append(_pad_e(r, len(header_e)))
                for r in data_er: rows_e.append(_pad_e(r, len(header_e)))
                parts_e = ["<html><head><meta charset='utf-8'></head><body>",
                           "<table border='1'>",
                           "<thead><tr>" + "".join(f"<th>{htmllib.escape(str(h or ''))}</th>" for h in header_e) + "</tr></thead>",
                           "<tbody>"]
                for r in rows_e:
                    parts_e.append("<tr>" + "".join(f"<td>{htmllib.escape(str(c or ''))}</td>" for c in r) + "</tr>")
                parts_e.append("</tbody></table></body></html>")
                combined_end_path = os.path.join(self.dl_dir, "종료권한리스트_합본.xls")
                with open(combined_end_path, "w", encoding="utf-8") as f:
                    f.write("\n".join(parts_e))

            if not has_normal and not has_end:
                self.downloaded.emit("", "부여/해제 다운로드 실패")
                return

            primary = combined_path if has_normal else combined_end_path
            self.downloaded.emit(primary, "")

        except SystemExit:
            self.downloaded.emit("", "사용자 취소")
        except Exception as e:
            self.downloaded.emit("", f"오류: {e}")
        finally:
            self._busy = False
            self.busyChanged.emit(False)

    def _ensure_iframe(self):
        try:
            d = self.driver
            d.switch_to.default_content()
            frames = d.find_elements(By.TAG_NAME, "iframe")
            if frames:
                d.switch_to.frame(frames[0])
            return True
        except Exception:
            return False

    def _go_iframe(self):
        try:
            d = self.driver
            d.switch_to.default_content()
            frames = d.find_elements(By.TAG_NAME, "iframe")
            if frames:
                d.switch_to.frame(frames[0])
            return True
        except Exception:
            return False

    def _wait_overlay_gone(self, timeout=10):
        d = self.driver
        try:
            WebDriverWait(d, timeout).until(
                EC.invisibility_of_element_located(
                    (By.CSS_SELECTOR, ".swal2-container, .loading, .blockUI, .blockOverlay")
                )
            )
        except Exception:
            pass

    def _debug_dump(self, tag: str):
        if not self.debug_enabled:
            return
        try:
            d = self.driver
            ts = int(time.time())
            base = f"debug_{tag}_{ts}"
            html_path = os.path.join(self.debug_dir, f"{base}.html")
            png_path  = os.path.join(self.debug_dir, f"{base}.png")
            with open(html_path, "w", encoding="utf-8") as f:
                f.write(d.page_source)
            try:
                d.save_screenshot(png_path)
            except Exception:
                pass
        except Exception:
            pass

    def _set_request_filter(self, reqtype: str):
        self._goto_progress_site()
        d = self.driver
        self._go_iframe()

        try:
            d.execute_script("""
              for (const id of ['approverYn','processYn','releaseYn']) {
                const el = document.getElementById(id);
                if (el) { el.value=''; el.dispatchEvent(new Event('change',{bubbles:true})); }
              }
            """)
        except Exception:
            pass

        is_release = (reqtype or "").strip() in (REQ_RELEASE, "권한해제", "해제", "release", "rel")

        try:
            if is_release:
                d.execute_script("""
                  const a=document.getElementById('approverYn');
                  const p=document.getElementById('processYn');
                  const r=document.getElementById('releaseYn');
                  if (a) { a.value='';  a.dispatchEvent(new Event('change',{bubbles:true})); }
                  if (p) { p.value='';  p.dispatchEvent(new Event('change',{bubbles:true})); }
                  if (r) { r.value='N'; r.dispatchEvent(new Event('change',{bubbles:true})); }
                """)
            else:
                d.execute_script("""
                  const a=document.getElementById('approverYn');
                  const p=document.getElementById('processYn');
                  const r=document.getElementById('releaseYn');
                  if (a) { a.value='1'; a.dispatchEvent(new Event('change',{bubbles:true})); }
                  if (p) { p.value='N'; p.dispatchEvent(new Event('change',{bubbles:true})); }
                  if (r) { r.value='';  r.dispatchEvent(new Event('change',{bubbles:true})); }
                """)
        except Exception:
            pass

        time.sleep(0.3)
        try:
            vals = d.execute_script("""
              const g = id => (document.getElementById(id)||{}).value || '';
              return {a:g('approverYn'), p:g('processYn'), r:g('releaseYn')};
            """)
            if is_release:
                ok = (vals.get('a','') == '' and vals.get('p','') == '' and vals.get('r','') == 'N')
            else:
                ok = (vals.get('a','') == '1' and vals.get('p','') == 'N' and vals.get('r','') == '')
            if not ok:
                return self._set_request_filter(reqtype)
        except Exception:
            pass

        self._debug_dump("progress_release_filter" if is_release else "progress_grant_filter")

    def _set_end_filter(self, reqtype_flag: str):
        self._goto_end_site()
        d = self.driver
        self._go_iframe()

        try:
            d.execute_script("""
              for (const id of ['approverYn','processYn','releaseYn']) {
                const el = document.getElementById(id);
                if (el) { el.value=''; el.dispatchEvent(new Event('change',{bubbles:true})); }
              }
            """)
        except Exception:
            pass

        try:
            if reqtype_flag == END_FLAG_GRANT:
                d.execute_script("""
                  const a=document.getElementById('approverYn');
                  const p=document.getElementById('processYn');
                  const r=document.getElementById('releaseYn');
                  if (a) { a.value='1'; a.dispatchEvent(new Event('change',{bubbles:true})); }
                  if (p) { p.value='N'; p.dispatchEvent(new Event('change',{bubbles:true})); }
                  if (r) { r.value='';  r.dispatchEvent(new Event('change',{bubbles:true})); }
                """)
            else:
                d.execute_script("""
                  const a=document.getElementById('approverYn');
                  const p=document.getElementById('processYn');
                  const r=document.getElementById('releaseYn');
                  if (a) { a.value='';  a.dispatchEvent(new Event('change',{bubbles:true})); }
                  if (p) { p.value='';  p.dispatchEvent(new Event('change',{bubbles:true})); }
                  if (r) { r.value='N'; r.dispatchEvent(new Event('change',{bubbles:true})); }
                """)
        except Exception:
            pass

        time.sleep(0.3)
        try:
            vals = d.execute_script("""
              const g = id => (document.getElementById(id)||{}).value || '';
              return {a:g('approverYn'), p:g('processYn'), r:g('releaseYn')};
            """)
            if reqtype_flag == END_FLAG_GRANT:
                ok = (vals.get('a','') == '1' and vals.get('p','') == 'N' and vals.get('r','') == '')
            else:
                ok = (vals.get('a','') == '' and vals.get('p','') == '' and vals.get('r','') == 'N')
            if not ok:
                return self._set_end_filter(reqtype_flag)
        except Exception:
            pass

        self._debug_dump("end_grant_filter" if reqtype_flag == END_FLAG_GRANT else "end_release_filter")

    def _click_search(self):
        d = self.driver
        try:
            self._go_iframe()
            try:
                WebDriverWait(d, 4).until(EC.element_to_be_clickable((By.ID, "btnSearch"))).click()
            except Exception:
                d.switch_to.default_content()
                WebDriverWait(d, 4).until(EC.element_to_be_clickable((By.ID, "btnSearch"))).click()
            self._wait_overlay_gone(10)
            self._go_iframe()
        except Exception:
            pass

    def _for_each_page(self, work_on_page):
        d = self.driver
        tried = 0

        self._check_cancel("process")
        if work_on_page():
            return True

        while tried < 200:
            self._check_cancel("process")
            self._go_iframe()
            next_btn = None
            try:
                cands = d.find_elements(By.CSS_SELECTOR, "a[aria-label='Next'], button[aria-label='Next']")
                if not cands:
                    cands = d.find_elements(By.XPATH, "//a[normalize-space()='다음' or normalize-space()='>'] | //button[normalize-space()='다음' or normalize-space()='>']")
                if cands:
                    next_btn = cands[-1]
                    cls = (next_btn.get_attribute("class") or "").lower()
                    aria = (next_btn.get_attribute("aria-disabled") or "").lower()
                    if "disabled" in cls or aria in ("true", "1"):
                        break
            except Exception:
                next_btn = None

            if not next_btn:
                break

            try:
                d.execute_script("arguments[0].scrollIntoView({block:'center'});", next_btn)
            except Exception:
                pass

            self._check_cancel("process")
            try:
                next_btn.click()
            except ElementClickInterceptedException:
                d.execute_script("arguments[0].click();", next_btn)

            self._wait_overlay_gone(10)
            self._go_iframe()
            time.sleep(0.4)

            self._check_cancel("process")
            if work_on_page():
                return True

            tried += 1

        return False

    def _norm(self, s):
        return (s or "").strip().replace("\u00a0", " ").lower()

    def _bus_code_variants(self, proj_raw: str):
        try:
            code5 = _yyxxx_from_proj_for_groupname(proj_raw)
            c5_int = int(code5)
            code4 = None
            if c5_int < 20000 and len(code5) == 5 and code5[2] == "0":
                code4 = code5[:2] + code5[3:]
            return code5, code4
        except Exception:
            base = re.sub(r"\D", "", str(proj_raw or ""))
            if len(base) >= 5:
                base = base[-5:]
            return base, None

    def _goto_first_page(self):
        d = self.driver
        try:
            self._go_iframe()
            btns = d.find_elements(By.XPATH, "//a[normalize-space()='처음'] | //button[normalize-space()='처음']")
            if btns:
                try:
                    d.execute_script("arguments[0].scrollIntoView({block:'center'});", btns[0])
                except Exception:
                    pass
                try:
                    btns[0].click()
                except ElementClickInterceptedException:
                    d.execute_script("arguments[0].click();", btns[0])
                self._wait_overlay_gone(10)
                self._go_iframe()
                return

            page1 = d.find_elements(By.XPATH, "//a[normalize-space()='1'] | //button[normalize-space()='1']")
            if page1:
                try:
                    d.execute_script("arguments[0].scrollIntoView({block:'center'});", page1[0])
                except Exception:
                    pass
                try:
                    page1[0].click()
                except ElementClickInterceptedException:
                    d.execute_script("arguments[0].click();", page1[0])
                self._wait_overlay_gone(10)
                self._go_iframe()
        except Exception:
            pass

    def _narrow_by_path_strict(self, rows, lv2, lv3, path_hint):
        if not rows:
            return rows

        def N(s): return (s or "").replace("\u00a0"," ").strip().lower()
        lv2n = N(lv2)
        lv3n = N(lv3)
        tail = (path_hint or "").split("\\")[-1]
        tailn = N(tail)

        must_tokens = [t for t in [lv2n, lv3n, tailn] if t]

        narrowed = []
        for r in rows:
            try:
                txt = N(r.text).replace("\n", " ")
                if all(tok in txt for tok in must_tokens):
                    narrowed.append(r)
            except Exception:
                pass

        return narrowed

    @pyqtSlot(object)
    def process(self, targets):
        if not self.is_ready():
            self.processed.emit([{'row': t.get('row'), 'ok': False, 'msg': '세션 준비 안됨'} for t in (targets or [])])
            return
        if self._busy:
            self.processed.emit([{'row': t.get('row'), 'ok': False, 'msg': '다른 작업 실행중'} for t in (targets or [])])
            return

        self._busy = True
        self.busyChanged.emit(True)
        self._cancel = False

        self._check_cancel("process")

        results = []
        try:
            d = self.driver

            def count_user_rows(user_text: str) -> int:
                self._go_iframe()
                user_eq = f".//td[normalize-space()='{user_text}']"
                code_pred = f".//td[contains(normalize-space(), '{code5}')]"
                if code4:
                    code_pred = f"({code_pred} or .//td[contains(normalize-space(), '{code4}')])"
                xpath = f"//tbody/tr[{user_eq} and {code_pred}]"
                rows = d.find_elements(By.XPATH, xpath)
                rows = self._narrow_by_path_strict(rows, lv2, lv3, path_hint)
                return len(rows)

            self._go_iframe()

            def _row_signature(tr):
                try:
                    tds = tr.find_elements(By.TAG_NAME, "td")
                    txts = []
                    for td in tds:
                        txt = (td.text or "").replace("\u00a0"," ").strip()
                        txts.append(re.sub(r"\s+", " ", txt))
                    return " | ".join(txts)
                except Exception:
                    return ""

            def _signature_exists(sig: str) -> bool:
                found = False
                def _find_on_page():
                    nonlocal found
                    self._go_iframe()
                    rows = d.find_elements(By.CSS_SELECTOR, "tbody tr")
                    for rnode in rows:
                        if _row_signature(rnode) == sig:
                            found = True
                            return True
                    return False
                self._for_each_page(_find_on_page)
                return found

            for t in (targets or []):
                self._check_cancel("process")
                row_idx = t.get('row')
                user = (t.get('user') or '').strip()
                lv2 = (t.get('lv2') or '').strip()
                lv3 = (t.get('lv3') or '').strip()
                path_hint = (t.get('path') or '').strip()
                kind = (t.get('kind') or '진행').strip()
                req = (t.get('req') or REQ_GRANT).strip()
                proj = (t.get('proj') or '').strip()
                code5, code4 = self._bus_code_variants(proj)

                ok = False
                msg = ""

                try:
                    if kind == "종료":
                        self._set_end_filter(END_FLAG_RELEASE if req == REQ_RELEASE else END_FLAG_GRANT)
                    else:
                        self._set_request_filter(req)

                    self._click_search()
                    time.sleep(0.6)
                    self._goto_first_page()

                    self._stabilize_grid()

                    before_total = 0
                    def sum_on_page():
                        nonlocal before_total
                        before_total += count_user_rows(user)
                        return False
                    self._for_each_page(sum_on_page)
                    target_sig_holder = {"sig": ""}

                    def work_on_page():
                        self._check_cancel("process")
                        self._go_iframe()

                        user_eq = f".//td[normalize-space()='{user}']"
                        code_pred = f".//td[contains(normalize-space(), '{code5}')]"
                        if code4:
                            code_pred = f"({code_pred} or .//td[contains(normalize-space(), '{code4}')])"

                        xpath = f"//tbody/tr[{user_eq} and {code_pred}]"
                        cand_rows = d.find_elements(By.XPATH, xpath)

                        if (t.get('kind') or '').strip() != '종료':
                            cand_rows = self._narrow_by_path_strict(cand_rows, lv2, lv3, path_hint)

                        if not cand_rows:
                            self._debug_dump("no_rows_after_filter")
                            return False

                        target_row = cand_rows[0]

                        req_is_release = (t.get('req') or REQ_GRANT) == REQ_RELEASE
                        if req_is_release:
                            btn_xpath = (
                                ".//button[normalize-space()='해제처리' "
                                "or contains(@class,'btn') and (contains(normalize-space(),'해제') or contains(@onclick,'Release'))]"
                            )
                        else:
                            btn_xpath = (
                                ".//button[normalize-space()='완료처리' "
                                "or normalize-space()='처리' "
                                "or contains(@class,'btn-outline-aurora')]"
                            )

                        target_sig_holder["sig"] = _row_signature(target_row)

                        self._check_cancel("process")

                        try:
                            d.execute_script("arguments[0].scrollIntoView({block:'center'});", target_row)
                        except Exception:
                            pass

                        def _btn_ready(_):
                            try:
                                el = target_row.find_element(By.XPATH, btn_xpath)
                                return el if el.is_displayed() and el.is_enabled() else False
                            except StaleElementReferenceException:
                                return False

                        try:
                            btn = WebDriverWait(target_row, 10).until(_btn_ready)
                        except TimeoutException:
                            self._debug_dump("no_action_button_wait_timeout")
                            return False

                        try:
                            d.execute_script("arguments[0].scrollIntoView({block:'center'});", btn)
                        except Exception:
                            pass

                        try:
                            btn.click()
                        except ElementClickInterceptedException:
                            d.execute_script("arguments[0].click();", btn)
                        except Exception:
                            pass

                        if self._cancel:
                            try:
                                self._wait_overlay_gone(0.2)
                                try:
                                    cancel_btn = WebDriverWait(d, 1).until(
                                        EC.element_to_be_clickable((By.CSS_SELECTOR, "button.swal2-cancel"))
                                    )
                                    try:
                                        cancel_btn.click()
                                    except ElementClickInterceptedException:
                                        d.execute_script("arguments[0].click();", cancel_btn)
                                except TimeoutException:
                                    d.execute_script("if (window.Swal) { try{Swal.close();}catch(e){} }")
                            except Exception:
                                pass
                            self._check_cancel("process")
                            return False

                        try:
                            WebDriverWait(d, 3).until(
                                EC.visibility_of_element_located((By.CSS_SELECTOR, ".swal2-container"))
                            )
                            self._check_cancel("process")
                            if self._cancel:
                                try:
                                    cancel_btn = WebDriverWait(d, 1).until(
                                        EC.element_to_be_clickable((By.CSS_SELECTOR, "button.swal2-cancel"))
                                    )
                                    try:
                                        cancel_btn.click()
                                    except ElementClickInterceptedException:
                                        d.execute_script("arguments[0].click();", cancel_btn)
                                except TimeoutException:
                                    d.execute_script("if (window.Swal) { try{Swal.close();}catch(e){} }")
                                self._check_cancel("process")
                                return False

                            okbtn = WebDriverWait(d, 5).until(
                                EC.element_to_be_clickable((By.CSS_SELECTOR, "button.swal2-confirm"))
                            )
                            try:
                                okbtn.click()
                            except ElementClickInterceptedException:
                                d.execute_script("arguments[0].click();", okbtn)
                        except TimeoutException:
                            pass

                        self._wait_overlay_gone(10)
                        self._stabilize_grid()
                        return True

                    self._goto_first_page()
                    did_click = self._for_each_page(work_on_page)
                    if not did_click:
                        ok, msg = False, "사번/프로젝트코드 불일치"
                    else:
                        timeout = time.time() + 12
                        disappeared = False
                        while time.time() < timeout:
                            self._check_cancel("process")
                            self._click_search()
                            time.sleep(0.8)
                            if target_sig_holder["sig"] and not _signature_exists(target_sig_holder["sig"]):
                                disappeared = True
                                break

                        if disappeared:
                            ok, msg = True, "처리 완료"
                        else:
                            after_total = 0
                            def sum_after_page():
                                nonlocal after_total
                                after_total += count_user_rows(user)
                                return False
                            self._for_each_page(sum_after_page)

                            ok = after_total < before_total
                            msg = f"처리 {'완료' if ok else '미확인'}: 전={before_total}, 후={after_total}"

                except SystemExit:
                    return
                except Exception as e:
                    ok, msg = False, f"오류: {e}"

                if not ok and self.debug_enabled:
                    try:
                        ts = int(time.time())
                        tag_user = re.sub(r'[^0-9A-Za-z_-]+', '_', user)[:30]
                        tag_lv3  = re.sub(r'[^0-9A-Za-z_-]+', '_', lv3)[:20]
                        base = f"fail_{ts}_{tag_user}_{tag_lv3}"
                        with open(os.path.join(self.debug_dir, f"{base}.html"), "w", encoding="utf-8") as f:
                            f.write(d.page_source)
                        try:
                            d.save_screenshot(os.path.join(self.debug_dir, f"{base}.png"))
                        except Exception:
                            pass
                    except Exception:
                        pass

                results.append({'row': row_idx, 'ok': ok, 'msg': msg})

            self._busy = False
            self.busyChanged.emit(False)
            self.processed.emit(results)

        except SystemExit:
            self._busy = False
            self.busyChanged.emit(False)
            if results:
                self.processed.emit(results)
            return
        except Exception as e:
            self._busy = False
            self.busyChanged.emit(False)
            self.processed.emit([{'row': t.get('row'), 'ok': False, 'msg': f'오류: {e}'} for t in (targets or [])])

# 알림/배지 기능 비활성화: 복원 가능하도록 기존 구현을 주석으로 보존합니다.
# class BusWatcher(QObject):
#     stopped = pyqtSignal()
#     readyChanged = pyqtSignal(bool, str)
#     countsReady = pyqtSignal(dict)

#     def __init__(self, dl_dir):
#         super().__init__()
#         self._user = ""
#         self._pw = ""
#         self._last_counts = {"신규미완료":0, "진행-부여":0, "진행-제거":0, "종료-부여":0, "종료-제거":0}
#         self._mgr = BusSessionManager(dl_dir)

#     def set_creds(self, user, pw):
#         self._user, self._pw = (user or "").strip(), (pw or "").strip()

#     def set_debug(self, on, debug_dir):
#         try:
#             self._mgr.set_debug(bool(on), debug_dir)
#         except Exception:
#             pass

#     def is_ready(self):
#         try:
#             return self._mgr.is_ready()
#         except Exception:
#             return False

#     @pyqtSlot()
#     def start(self):
#         if not (self._user and self._pw):
#             self.readyChanged.emit(False, "워처 자격증명 없음")
#             return
#         try:
#             self._mgr.set_creds(self._user, self._pw)
#             self._mgr.start()
#             self.readyChanged.emit(True, "워처 준비됨")
#         except Exception:
#             self.readyChanged.emit(False, "워처 시작 실패")

#     @pyqtSlot()
#     def stop(self):
#         try:
#             self._mgr.stop()
#         except Exception:
#             pass
#         self.readyChanged.emit(False, "워처 정지됨")
#         self.stopped.emit()

#     @pyqtSlot()
#     def collect_counts(self):
#         try:
#             self._mgr.collect_request_counts()
#             self._last_counts = getattr(self._mgr, "_last_counts", self._last_counts)
#             self.countsReady.emit(self._last_counts)
#         except Exception:
#             self.countsReady.emit(self._last_counts)

class CheckBoxHeader(QHeaderView):
    stateChanged = pyqtSignal(bool)

    def __init__(self, orientation, parent=None):
        super().__init__(orientation, parent)
        self.setSectionsClickable(True)
        self.setHighlightSections(False)

        self._cb = QCheckBox(self)
        self._cb.setTristate(False)
        self._cb.setChecked(True)
        self._cb.setStyleSheet("QCheckBox{background:transparent; padding:0; margin:0;}")

        self._cb.stateChanged.connect(lambda st: self.stateChanged.emit(st == Qt.Checked))

        self.sectionResized.connect(self._reposition)
        self.sectionMoved.connect(self._reposition)
        self.geometriesChanged.connect(self._reposition)
        self.sectionCountChanged.connect(lambda *_: self._reposition())

        self._reposition()

    def setChecked(self, on: bool):
        self._cb.blockSignals(True)
        self._cb.setChecked(bool(on))
        self._cb.blockSignals(False)
        self.updateSection(0)

    def _reposition(self, *args):
        if self.count() == 0:
            return

        x0 = self.sectionViewportPosition(0)
        w0 = self.sectionSize(0)
        h  = self.height()

        sz = self._cb.sizeHint()
        x  = x0 + max(0, (w0 - sz.width()) // 2)
        y  = max(0, (h  - sz.height()) // 2)

        if w0 < sz.width():
            x = x0 + 2

        self._cb.setGeometry(x, y, sz.width(), sz.height())
        self._cb.raise_()
        self._cb.show()

    def mousePressEvent(self, e):
        if self.logicalIndexAt(e.pos()) == 0:
            self._cb.toggle()
            e.accept()
            return
        super().mousePressEvent(e)


class StatusMessage(QTextEdit):
    """Readable short status with bounded scrolling for long error details."""
    def __init__(self,text='',parent=None):
        super().__init__(parent)
        self.setReadOnly(True); self.setAcceptRichText(False)
        self.setFrameShape(QTextEdit.NoFrame)
        self.setLineWrapMode(QTextEdit.WidgetWidth)
        self.setWordWrapMode(QTextOption.WrapAnywhere)
        self.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.document().setDocumentMargin(2)
        self.document().documentLayout().documentSizeChanged.connect(self._fit_height)
        self.setText(text)

    def setText(self,text):
        text = str(text or '')
        if text==self.toPlainText():
            return
        self.setPlainText(text)
        self.verticalScrollBar().setValue(0)
        self._fit_height()

    def text(self):
        return self.toPlainText()

    def _fit_height(self,*args):
        # Leave a few pixels below the final baseline for Windows/DPI font metrics.
        needed = int(self.document().size().height()+9)
        height = min(112,max(24,needed))
        if self.height()!=height or self.minimumHeight()!=height:
            self.setFixedHeight(height)


class BusySpinner(QWidget):
    def __init__(self,parent=None):
        super().__init__(parent)
        self.setFixedSize(30,30)
        self.angle = 0
        self.timer = QTimer(self); self.timer.setInterval(40)
        self.timer.timeout.connect(self._tick)

    def _tick(self):
        self.angle = (self.angle+18)%360
        self.update()

    def showEvent(self,event):
        super().showEvent(event); self.timer.start()

    def hideEvent(self,event):
        self.timer.stop(); super().hideEvent(event)

    def paintEvent(self,event):
        painter = QPainter(self); painter.setRenderHint(QPainter.Antialiasing)
        pen = QPen(QColor('#dbe5f5'),3); painter.setPen(pen)
        painter.drawEllipse(5,5,self.width()-10,self.height()-10)
        pen.setColor(QColor('#2864ed')); pen.setCapStyle(Qt.RoundCap); painter.setPen(pen)
        painter.drawArc(5,5,self.width()-10,self.height()-10,-self.angle*16,105*16)


class ProcessStepIcon(QLabel):
    """A stage number/checkmark with a running animation in the same space."""
    def __init__(self,text):
        super().__init__(text)
        self.spinner = BusySpinner(self)
        self.spinner.setFixedSize(24,24)
        self.spinner.hide()

    def setProperty(self,name,value):
        result = super().setProperty(name,value)
        if name == 'state':
            self.spinner.setVisible(value == 'running')
            if value == 'running':
                self.setText('')
        return result


class BusLoginDialog(QDialog):
    """Startup-only credentials prompt with an explicit remember choice."""
    def __init__(self,username='',parent=None):
        super().__init__(parent)
        self.setWindowTitle('BUS 로그인')
        self.setWindowFlags(self.windowFlags() & ~Qt.WindowContextHelpButtonHint)
        self.setMinimumWidth(360)
        self.setStyleSheet("""
            QDialog {background:#ffffff;}
            QLabel {background:transparent;color:#263245;font-size:13px;}
            QLabel#loginTitle {font-size:18px;font-weight:600;}
            QLineEdit {background:#ffffff;color:#263245;border:1px solid #cbd6ea;border-radius:5px;padding:7px;font-size:13px;}
            QLineEdit:focus {border-color:#2864ed;}
            QCheckBox {background:transparent;color:#263245;font-size:13px;}
            QCheckBox::indicator {background:#ffffff;border:1px solid #b9c6d8;border-radius:3px;width:15px;height:15px;}
            QCheckBox::indicator:checked {background:#2864ed;border-color:#2864ed;}
            QPushButton {background:#ffffff;color:#263245;border:1px solid #cbd6ea;border-radius:5px;padding:7px 16px;}
            QPushButton#loginButton {background:#2864ed;color:#ffffff;border-color:#2864ed;}
            QPushButton#loginButton:disabled {background:#edf1f7;color:#9da8b8;border-color:#dce3ec;}
        """)
        layout = QVBoxLayout(self); layout.setContentsMargins(24,22,24,20); layout.setSpacing(16)
        title = QLabel('BUS 로그인'); title.setObjectName('loginTitle'); layout.addWidget(title)
        form = QFormLayout(); form.setHorizontalSpacing(12); form.setVerticalSpacing(12)
        self.username = QLineEdit(username); self.username.setPlaceholderText('아이디')
        self.password = QLineEdit(); self.password.setPlaceholderText('비밀번호')
        self.password.setEchoMode(QLineEdit.Password)
        form.addRow('아이디',self.username); form.addRow('비밀번호',self.password)
        layout.addLayout(form)
        self.remember = QCheckBox('계정 정보 저장'); self.remember.setChecked(False)
        layout.addWidget(self.remember)
        buttons = QDialogButtonBox(self)
        self.login_button = buttons.addButton('로그인',QDialogButtonBox.AcceptRole)
        self.login_button.setObjectName('loginButton'); self.login_button.setDefault(True)
        buttons.addButton('취소',QDialogButtonBox.RejectRole)
        buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject)
        self.username.textChanged.connect(self._validate)
        self.password.textChanged.connect(self._validate)
        layout.addWidget(buttons)
        self._validate()
        (self.password if username else self.username).setFocus()

    def values(self):
        return self.username.text().strip(), self.password.text()

    def _validate(self):
        username,password = self.values()
        self.login_button.setEnabled(bool(username and password))

    def accept(self):
        if all(self.values()):
            super().accept()


class StartupLoadingDialog(QDialog):
    """Asynchronous startup coordinator. No permission jobs are submitted here."""
    def __init__(self,host):
        super().__init__(host)
        self.host = host
        host._startup_dialog = self
        self.stage = -1; self.active = False; self.api_thread = None; self.api_task = None
        self._login_failed = False
        self.setWindowTitle('Folder Grant 시작'); self.setFixedSize(420,360)
        self.setWindowFlags(self.windowFlags() & ~Qt.WindowContextHelpButtonHint)
        self.setStyleSheet('QWidget{background:transparent;} QDialog{background:#fff;} QLabel{color:#263245;background:transparent;font-size:13px;} QLabel#startupTitle{font-size:18px;font-weight:700;} QPushButton{background:#fff;color:#2864ed;border:1px solid #cbd6ea;border-radius:5px;padding:7px 12px;} QPushButton:hover{background:#eef4ff;} QPushButton:disabled{color:#98a2b3;}')
        layout = QVBoxLayout(self); layout.setContentsMargins(28,24,28,22); layout.setSpacing(12)
        title = QLabel('Folder Grant 준비 중'); title.setObjectName('startupTitle'); title.setFont(QFont('Segoe UI',15,QFont.Bold)); layout.addWidget(title)
        self.step_rows = []; self.spinners = []; self.icons = []
        for text in ('API 연결 확인','BUS 로그인','권한 요청 조회','신규 프로젝트 조회'):
            row = QHBoxLayout(); row.setSpacing(10)
            slot = QWidget(); slot.setFixedSize(30,30)
            spinner = BusySpinner(slot); spinner.hide()
            icon = QLabel('○',slot); icon.setStyleSheet('color:#8a99b2;'); icon.setAlignment(Qt.AlignCenter); icon.setGeometry(0,0,30,30)
            label = QLabel(text); row.addWidget(slot); row.addWidget(label,1); layout.addLayout(row)
            self.step_rows.append(label); self.spinners.append(spinner); self.icons.append(icon)
        self.message = QLabel('연결과 요청 목록을 준비합니다.'); self.message.setWordWrap(True); layout.addWidget(self.message)
        buttons = QHBoxLayout(); buttons.addStretch()
        self.retry = QPushButton('다시 시도'); self.retry.clicked.connect(self.start); self.retry.hide()
        self.enter = QPushButton('메인 화면 열기'); self.enter.clicked.connect(self.open_main)
        buttons.addWidget(self.retry); buttons.addWidget(self.enter); layout.addLayout(buttons)
        self.deadline = QTimer(self); self.deadline.setSingleShot(True); self.deadline.setInterval(120000)
        self.deadline.timeout.connect(self._timed_out)
        self.host.startup_result.connect(self._loaded)
        self.host.session.busyChanged.connect(self._update_retry)

    def _timed_out(self):
        self.fail('응답이 지연되고 있습니다. 다시 시도하거나 메인 화면에서 확인하세요.')

    def _update_retry(self):
        self.retry.setEnabled(self.api_thread is None and not self.host.session.is_busy() and not getattr(self.host,'_pending_bus_callbacks',[]))

    def start(self):
        if self.api_thread is not None or self.host.session.is_busy() or getattr(self.host,'_pending_bus_callbacks',[]):
            return
        self.active = True; self.host._startup_active = True; self.retry.hide()
        for icon,spinner,label in zip(self.icons,self.spinners,self.step_rows):
            spinner.hide(); icon.setText('○'); icon.setStyleSheet('color:#8a99b2;'); icon.show(); label.setStyleSheet('')
        self.stage = -1
        self._stage(0)
        if not self.host.api_config.get('api_key'):
            self.deadline.stop()
            self.host._open_api_settings()
            self.deadline.start()
        if not self.host.api_config.get('api_key'):
            self.fail('API 연결 설정이 필요합니다.'); return
        thread = QThread(self); task = ApiConnectionTask(self.host._api_client())
        self.api_thread,self.api_task = thread,task; task.moveToThread(thread)
        thread.started.connect(task.run)
        task.succeeded.connect(self._api_ready); task.failed.connect(self._api_failed)
        task.finished.connect(task.deleteLater); task.finished.connect(thread.quit,type=Qt.DirectConnection)
        thread.finished.connect(self._api_finished); thread.finished.connect(thread.deleteLater)
        thread.start()

    def _stage(self,index):
        if self.stage>=0:
            self.spinners[self.stage].hide(); self.icons[self.stage].setText('✓'); self.icons[self.stage].setStyleSheet('color:#16a34a;'); self.icons[self.stage].show()
            self.step_rows[self.stage].setStyleSheet('color:#263245;')
        self.stage = index; self.icons[index].hide(); self.spinners[index].show()
        self.step_rows[index].setStyleSheet('color:#2864ed;font-weight:600;')
        self.message.setText(self.step_rows[index].text()+' 중…'); self.deadline.start()
        if index==2:
            self.host.table.set_empty_message('요청을 조회하고 있습니다…')
        elif index==3:
            self.host._ensure_project_viewer()
            self.host._newdlg.tbl.set_empty_message('요청을 조회하고 있습니다…')

    def _api_ready(self,result):
        if not self.active or self.stage!=0:
            return
        self.host._set_api_connection_state(True,result.get('ready') or {})
        self._stage(1)
        if not self.host.creds.get('id') or not self.host.creds.get('pw') or self._login_failed:
            self.deadline.stop()
            login = BusLoginDialog(self.host.creds.get('id',''),self)
            if login.exec_()!=QDialog.Accepted:
                self.fail('BUS 로그인을 취소했습니다. 다시 시도하거나 메인 화면을 여세요.'); return
            username,password = login.values()
            self.host._apply_login_credentials(username,password,login.remember.isChecked())
            self._login_failed = False
            self.deadline.start()
        self.host.ensure_bus_session_async('시작 준비',self._bus_ready)

    def _api_failed(self,message):
        if not self.active:
            return
        self.host._set_api_connection_state(False)
        self.host._log(message,summary='시작 준비: API 연결 실패')
        self.fail('API 연결에 실패했습니다. 설정과 서버 상태를 확인하세요.')

    def _api_finished(self):
        self.api_thread = None; self.api_task = None; self._update_retry()

    def _bus_ready(self,ok):
        if self.stage==1:
            self._login_failed = not ok
        self._update_retry()
        if not self.active or self.stage!=1:
            return
        if not ok:
            self._login_failed = True
            self.fail('BUS 로그인에 실패했습니다. 계정과 연결 상태를 확인하세요.'); return
        self._stage(2)
        self._load_access_when_idle()

    def _loaded(self,kind,ok,message):
        if not self.active:
            return
        expected = 'access' if self.stage==2 else 'project' if self.stage==3 else None
        if kind!=expected:
            return
        if not ok:
            self.fail(message); return
        if kind=='access':
            self._stage(3); self._load_project_when_idle()
        else:
            self.spinners[3].hide(); self.icons[3].setText('✓'); self.icons[3].show()
            self.open_main()

    def _load_access_when_idle(self):
        if not self.active or self.stage!=2:
            return
        if self.host.session.is_busy():
            QTimer.singleShot(50,self._load_access_when_idle)
            return
        self.host._request_and_import()

    def _load_project_when_idle(self):
        if not self.active or self.stage!=3:
            return
        if self.host.session.is_busy():
            QTimer.singleShot(50,self._load_project_when_idle)
            return
        self.host._load_new_items()

    def fail(self,message):
        if not self.active:
            return
        self.active = False; self.host._startup_active = False; self.deadline.stop()
        if self.stage in (2,3):
            table = self.host.table if self.stage==2 else self.host._newdlg.tbl
            table.set_empty_message('조회하지 못했습니다. BUS 조회로 다시 시도하세요.')
        if self.stage>=0:
            self.spinners[self.stage].hide(); self.icons[self.stage].setText('!'); self.icons[self.stage].setStyleSheet('color:#dc2626;'); self.icons[self.stage].show()
        self.host._log('시작 준비 실패: '+message)
        self.message.setText(message); self.retry.show(); self._update_retry()

    def open_main(self):
        self.active = False; self.host._startup_active = False; self.deadline.stop()
        self.host.show(); self.hide()

    def closeEvent(self,event):
        self.open_main(); event.accept()


class FocuslessItemDelegate(QStyledItemDelegate):
    """Hide the cell focus outline while preserving keyboard navigation."""
    def initStyleOption(self, option, index):
        super().initStyleOption(option, index)
        option.state &= ~QStyle.State_HasFocus


class RequestTable(QTableWidget):
    """A viewport label distinguishes an empty result from a pending fetch."""
    def __init__(self,*args,**kwargs):
        super().__init__(*args,**kwargs)
        self.setItemDelegate(FocuslessItemDelegate(self))
        self.empty_label = QLabel('아직 조회하지 않았습니다.',self.viewport())
        self.empty_label.setObjectName('mutedLabel')
        self.empty_label.setStyleSheet('background:transparent;')
        self.empty_label.setAlignment(Qt.AlignCenter)
        self.empty_label.setAttribute(Qt.WA_TransparentForMouseEvents)
        self.model().rowsInserted.connect(self._update_empty)
        self.model().rowsRemoved.connect(self._update_empty)
        self.model().modelReset.connect(self._update_empty)

    def set_empty_message(self,text):
        self.empty_label.setText(text); self._update_empty()

    def _update_empty(self,*args):
        self.empty_label.setGeometry(self.viewport().rect())
        self.empty_label.setVisible(self.rowCount()==0)

    def resizeEvent(self,event):
        super().resizeEvent(event)
        if hasattr(self,'empty_label'):
            self._update_empty()


# 알림/배지 기능 비활성화: 복원 가능하도록 기존 구현을 주석으로 보존합니다.
# class WorkspaceTabBar(QTabBar):
#     """Keep notification badges attached to the upper right of each label."""
#     def __init__(self, parent=None):
#         super().__init__(parent)
#         self.count_badges = []

#     def add_count_badges(self):
#         for index in range(self.count()):
#             badge = QLabel('0',self); badge.setObjectName('requestCountBadge')
#             badge.setAlignment(Qt.AlignCenter); badge.setFixedSize(26,18); badge.hide()
#             badge.setAttribute(Qt.WA_TransparentForMouseEvents)
#             self.count_badges.append(badge)
#         self._position_badges()

#     def tabSizeHint(self,index):
#         size = super().tabSizeHint(index)
#         size.setWidth(size.width()+24)
#         return size

#     def tabLayoutChange(self):
#         super().tabLayoutChange()
#         self._position_badges()

#     def resizeEvent(self,event):
#         super().resizeEvent(event)
#         self._position_badges()

#     def _position_badges(self):
#         for index,badge in enumerate(getattr(self,'count_badges',[])):
#             option = QStyleOptionTab(); self.initStyleOption(option,index)
#             rect = self.style().subElementRect(QStyle.SE_TabBarTabText,option,self)
#             width = option.fontMetrics.horizontalAdvance(self.tabText(index))
#             x = rect.center().x()+width//2+4
#             y = rect.center().y()-option.fontMetrics.height()//2-10
#             badge.move(min(x,self.tabRect(index).right()-badge.width()-3),max(3,y))
#             badge.raise_()


class CopyTable(RequestTable):
    def keyPressEvent(self, event):
        if event.matches(QKeySequence.Copy):
            ranges = self.selectedRanges()
            if not ranges:
                super().keyPressEvent(event); return
            blocks = []
            for rng in ranges:
                rows = []
                for r in range(rng.topRow(), rng.bottomRow()+1):
                    cols = []
                    for c in range(rng.leftColumn(), rng.rightColumn()+1):
                        it = self.item(r, c)
                        cols.append("" if it is None else it.text())
                    rows.append("\t".join(cols))
                blocks.append("\n".join(rows))
            QGuiApplication.clipboard().setText("\n\n".join(blocks))
        else:
            super().keyPressEvent(event)

class ManualEntryDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("수동 입력")
        self.setModal(True)

        self.cb_kind = QComboBox()
        self.cb_kind.addItems(["진행", "종료"])
        self.cb_kind.setCurrentText("진행")
        self.cb_kind.currentTextChanged.connect(self._on_kind_changed)

        self.cb_reqtype = QComboBox()
        self.le_user = QLineEdit()
        self.le_proj = QLineEdit()
        self.cb_lv2  = QComboBox()
        self.cb_lv3  = QComboBox()
        self.cb_role = QComboBox()

        self.cb_reqtype.addItems([REQ_GRANT, REQ_RELEASE])
        self.cb_reqtype.setCurrentText(REQ_GRANT)

        self.cb_lv2.addItems(["Study", "Isolated"])
        self.cb_lv2.setCurrentIndex(-1)

        self.cb_lv3.addItems(LEVEL3_CHOICES)
        self.cb_lv3.setCurrentIndex(-1)

        self.cb_role.addItem("")
        self.cb_role.setEnabled(False)

        self.cb_lv2.currentTextChanged.connect(self._refresh_role_candidates)
        self.cb_lv3.currentTextChanged.connect(self._refresh_role_candidates)
        self.le_proj.textChanged.connect(self._refresh_role_candidates)

        form = QFormLayout()
        form.addRow("구분*", self.cb_kind)
        form.addRow("요청사항*", self.cb_reqtype)
        form.addRow("사번*", self.le_user)
        form.addRow("프로젝트코드*", self.le_proj)
        form.addRow("Level2*", self.cb_lv2)
        form.addRow("Level3*", self.cb_lv3)
        form.addRow("STATROLE", self.cb_role)

        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel, parent=self)
        btns.accepted.connect(self._on_ok)
        btns.rejected.connect(self.reject)

        lay = QVBoxLayout(self)
        lay.addLayout(form)
        lay.addWidget(btns)

        self.result_row = None

    def _is_new_by_proj_text(self) -> bool:
        txt = self.le_proj.text().strip()
        if not txt:
            return True
        try:
            return is_new_template(txt)
        except Exception:
            return True

    def _refresh_role_candidates(self):
        try:
            lv2 = self.cb_lv2.currentText().strip() if self.cb_lv2.currentIndex() != -1 else ""
            lv3 = self.cb_lv3.currentText().strip() if self.cb_lv3.currentIndex() != -1 else ""
            is_new = self._is_new_by_proj_text()

            self.cb_role.blockSignals(True)
            self.cb_role.clear()
            self.cb_role.addItem("")
            self.cb_role.blockSignals(False)
            self.cb_role.setCurrentIndex(0)
            self.cb_role.setEnabled(False)

            if not is_stat_lv3(lv3):
                return

            candidates = []
            if lv2 == "Study":
                candidates = sorted(STUDY_ROLES)
            elif lv2 == "Isolated":
                if is_stat_idmc_lv3(lv3):
                    if is_stat_idmc_new_policy(self.le_proj.text().strip()):
                        candidates = sorted(STUDY_ROLES)
                else:
                    candidates = sorted(ISOLATED_ROLES) if is_new else ["Randomization Statistician"]

            if candidates:
                self.cb_role.blockSignals(True)
                for r in candidates:
                    self.cb_role.addItem(r)
                self.cb_role.blockSignals(False)
                self.cb_role.setEnabled(True)
        except Exception as e:
            parent = self.parent()
            if parent and hasattr(parent, "_log"):
                parent._log(f"[역할후보 갱신 예외] {e}")

    def _on_kind_changed(self, text: str):
        """구분이 '종료'면 Level2/Level3 입력을 비활성화하고 값을 비움."""
        is_end = (text or "").strip() == "종료"

        for cb in (self.cb_lv2, self.cb_lv3):
            try:
                cb.blockSignals(True)
                if is_end:
                    cb.setCurrentIndex(-1)
                cb.setEnabled(not is_end)
            finally:
                cb.blockSignals(False)

        try:
            self._refresh_role_candidates()
        except Exception:
            pass

    def _on_ok(self):
        try:
            kind   = self.cb_kind.currentText().strip()
            reqtype = self.cb_reqtype.currentText().strip()
            user = self.le_user.text().strip()
            proj = self.le_proj.text().strip()
            lv2  = self.cb_lv2.currentText().strip() if self.cb_lv2.currentIndex() != -1 else ""
            lv3  = self.cb_lv3.currentText().strip() if self.cb_lv3.currentIndex() != -1 else ""
            role = self.cb_role.currentText().strip() if self.cb_role.isEnabled() else ""

            missing = []
            if not user: missing.append("사번")
            if not proj: missing.append("프로젝트코드")
            if kind != "종료":
                if not lv2:  missing.append("Level2")
                if not lv3:  missing.append("Level3")
            if missing:
                QMessageBox.warning(self, "입력 누락", f"다음 항목을 입력해 주세요: {', '.join(missing)}")
                return

            if not is_stat_lv3(lv3):
                role = ""

            self.result_row = (reqtype, user, proj, lv2, lv3, "", role, kind)
            self.accept()
        except Exception as e:
            try:
                parent = self.parent()
                if parent and hasattr(parent, "_log"):
                    parent._log(f"[수동입력 예외] {e}")
            except:
                pass
            QMessageBox.critical(self, "오류", f"수동 입력 처리 중 오류:\n{e}")

    def _set_roles_for_lv2(self, lv2_text):
        self.cb_role.clear()
        self.cb_role.addItem("")
        if lv2_text == "Study":
            for r in sorted(STUDY_ROLES):
                self.cb_role.addItem(r)
        elif lv2_text == "Isolated":
            for r in sorted(ISOLATED_ROLES):
                self.cb_role.addItem(r)
        self.cb_role.setCurrentIndex(0)

class CheckBoxHeaderAt(QHeaderView):
    stateChanged = pyqtSignal(bool)

    def __init__(self, orientation, parent=None, target_index: int = 0, initially_checked: bool = True):
        super().__init__(orientation, parent)
        self._target = int(target_index)

        self._cb = QCheckBox(self)
        self._cb.setTristate(False)
        self._cb.setChecked(bool(initially_checked))
        self._cb.setStyleSheet("QCheckBox{background:transparent; padding:0; margin:0;}")
        self._cb.stateChanged.connect(lambda st: self.stateChanged.emit(st == Qt.Checked))

        self.setSectionsClickable(True)
        self.setHighlightSections(False)

        self.sectionResized.connect(self._reposition)
        self.sectionMoved.connect(self._reposition)
        self.sectionCountChanged.connect(lambda *_: self._reposition())

    def showEvent(self, e):
        super().showEvent(e)
        self._reposition()

    def setTargetSection(self, idx: int):
        self._target = int(idx)
        self._reposition()

    def setChecked(self, on: bool):
        self._cb.blockSignals(True)
        self._cb.setChecked(bool(on))
        self._cb.blockSignals(False)
        self.updateSection(self._target)

    def _reposition(self, *args):
        if self.count() == 0 or self._target < 0 or self._target >= self.count():
            self._cb.hide(); return
        x0 = self.sectionViewportPosition(self._target)
        w0 = self.sectionSize(self._target)
        h  = self.height()
        sz = self._cb.sizeHint()
        x  = x0 + max(0, (w0 - sz.width()) // 2)
        y  = max(0, (h  - sz.height()) // 2)
        self._cb.setGeometry(x, y, sz.width(), sz.height())
        self._cb.raise_()
        self._cb.show()

    def mousePressEvent(self, e):
        if self.logicalIndexAt(e.pos()) == self._target:
            self._cb.toggle()
            e.accept()
            return
        super().mousePressEvent(e)

class ManualNewRequestDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("수동 추가")
        lay = QVBoxLayout(self)

        form = QFormLayout()
        self.edt_proj = QLineEdit(self)
        self.edt_proj.setPlaceholderText("예: 25001")
        self.edt_proj.setMaxLength(5)
        self.edt_name = QLineEdit(self)
        self.edt_name.setPlaceholderText("프로젝트명")
        form.addRow("프로젝트 코드", self.edt_proj)
        form.addRow("프로젝트명", self.edt_name)
        lay.addLayout(form)

        btn_box = QDialogButtonBox(Qt.Horizontal, self)
        self.btn_ok = btn_box.addButton("생성", QDialogButtonBox.AcceptRole)
        self.btn_cancel = btn_box.addButton("취소", QDialogButtonBox.RejectRole)
        self.btn_ok.clicked.connect(self.accept)
        self.btn_cancel.clicked.connect(self.reject)
        lay.addWidget(btn_box)

    def values(self) -> tuple[str, str]:
        return self.edt_proj.text().strip(), self.edt_name.text().strip()


class DetailTextValue(QTextEdit):
    """Wrap long IDs/UNC paths without changing the text users copy."""
    def __init__(self, text):
        super().__init__()
        self.setReadOnly(True); self.setFrameShape(QTextEdit.NoFrame)
        self.setStyleSheet('background:transparent;border:0;padding:0;')
        self.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.setWordWrapMode(QTextOption.WrapAtWordBoundaryOrAnywhere)
        self.document().setDocumentMargin(0)
        self.setSizePolicy(QSizePolicy.Expanding,QSizePolicy.Fixed)
        self.setMinimumWidth(0); self.setPlainText(text)
        QTimer.singleShot(0,self._fit_text)

    def text(self):
        return self.toPlainText()

    def sizeHint(self):
        return QSize(90,self.height())

    def resizeEvent(self,event):
        super().resizeEvent(event)
        self._fit_text()

    def _fit_text(self):
        self.document().setTextWidth(max(1,self.viewport().width()))
        height = max(self.fontMetrics().height()+2,int(self.document().size().height()+2))
        if self.height()!=height:
            self.setFixedHeight(height)


class NewItemsViewer(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.host = parent
        self.setWindowTitle("신규 폴더 생성요청 목록")
        self.resize(800, 520)

        lay = QVBoxLayout(self)
        self.tbl = RequestTable(0, 0, self)
        self.tbl.setContextMenuPolicy(Qt.CustomContextMenu)
        self.tbl.customContextMenuRequested.connect(self._open_table_menu)

        hdr = CheckBoxHeader(Qt.Horizontal, self.tbl)
        hdr.stateChanged.connect(self._toggle_all_rows)
        self.tbl.setHorizontalHeader(hdr)

        self.tbl.setAlternatingRowColors(True)
        self.tbl.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.tbl.horizontalHeader().setStretchLastSection(True)

        lay.setContentsMargins(0,0,0,0)
        lay.setSpacing(0)
        self.toolbar = QWidget(); self.toolbar.setObjectName('requestToolbar')
        top = QHBoxLayout(self.toolbar); top.setContentsMargins(15,12,12,12)
        heading = QLabel('요청 목록'); heading.setObjectName('sectionTitle')
        top.addWidget(heading)
        # self.new_request_notice = QLabel('새 요청이 있습니다 · BUS 신규조회'); self.new_request_notice.setObjectName('mutedLabel'); self.new_request_notice.hide()
        # top.addWidget(self.new_request_notice); top.addStretch()
        top.addStretch()
        self.btn_refresh = QPushButton('BUS 신규조회'); self.btn_refresh.setObjectName('primaryButton')
        self.btn_manual = QPushButton('수동 추가')
        top.addWidget(self.btn_manual); top.addWidget(self.btn_refresh)
        self.btn_manual.clicked.connect(self._on_manual_clicked)
        lay.addWidget(self.toolbar)
        content = QHBoxLayout(); content.setSpacing(10)
        content.setContentsMargins(0,0,0,0)
        content.addWidget(self.tbl,1)
        self.detail_panel = QWidget(); self.detail_panel.setObjectName('detailPanel')
        self.detail_panel.setFixedWidth(258)
        detail = QVBoxLayout(self.detail_panel); detail.setContentsMargins(17,14,17,14)
        title = QLabel('선택 요청 상세'); title.setObjectName('sectionTitle')
        detail.addWidget(title)
        fields = QWidget(); fields.setObjectName('projectDetailFields')
        self.detail_form = QFormLayout(fields); self.detail_form.setVerticalSpacing(9)
        self.detail_form.setContentsMargins(0,0,0,0)
        self.detail_form.setFieldGrowthPolicy(QFormLayout.AllNonFixedFieldsGrow)
        self.detail_form.setRowWrapPolicy(QFormLayout.WrapLongRows)
        self.detail_scroll = QScrollArea(); self.detail_scroll.setObjectName('projectDetailScroll'); self.detail_scroll.setWidgetResizable(True)
        self.detail_scroll.setFrameShape(QScrollArea.NoFrame)
        body = QWidget(); body.setObjectName('projectDetailBody'); body_layout = QVBoxLayout(body)
        body_layout.setContentsMargins(0,0,0,0); body_layout.setSpacing(12)
        body_layout.addWidget(fields)
        process_title = QLabel('처리 진행'); process_title.setObjectName('sectionTitle')
        body_layout.addWidget(process_title)
        self.project_step_icons = []; self.project_step_labels = []
        self.project_step_types = ['request','ensure_ad_group','ensure_project_folder','cleanup_unknown_acl',
                                   'grant_project_root_acl','grant_study_all_acl','verify_project_creation','bus']
        self.project_step_names = ['요청 확인','보안 그룹 준비','프로젝트 폴더 준비','기본 ACL 정리',
                                   '프로젝트 권한 적용','기본 폴더 권한 적용','생성 결과 확인','BUS 완료 처리']
        stages = QWidget(); stages.setObjectName('projectProcess'); stage_layout = QVBoxLayout(stages); stage_layout.setContentsMargins(0,0,0,0); stage_layout.setSpacing(9)
        for index,name in enumerate(self.project_step_names):
            row_layout = QHBoxLayout(); row_layout.setContentsMargins(0,0,0,0); row_layout.setSpacing(12)
            icon = ProcessStepIcon(str(index+1)); icon.setObjectName('stepIcon'); icon.setAlignment(Qt.AlignCenter); icon.setFixedSize(24,24)
            label = QLabel(name); label.setObjectName('stepLabel'); label.setWordWrap(True)
            row_layout.addWidget(icon); row_layout.addWidget(label,1)
            stage_layout.addLayout(row_layout)
            self.project_step_icons.append(icon); self.project_step_labels.append(label)
        body_layout.addWidget(stages)
        extra_title = QLabel('서버 작업 상세'); extra_title.setObjectName('sectionTitle'); body_layout.addWidget(extra_title)
        self.server_detail_form = QFormLayout(); self.server_detail_form.setContentsMargins(0,0,0,0); self.server_detail_form.setVerticalSpacing(9)
        self.server_detail_form.setFieldGrowthPolicy(QFormLayout.AllNonFixedFieldsGrow)
        self.server_detail_form.setRowWrapPolicy(QFormLayout.WrapLongRows)
        body_layout.addLayout(self.server_detail_form); body_layout.addStretch()
        self.detail_scroll.setWidget(body)
        self._detail_generation = 0; self._detail_token = None
        detail.addWidget(self.detail_scroll,1)
        self.detail_result = QLabel('요청을 선택하세요')
        self.detail_result.setObjectName('detailBanner'); self.detail_result.setWordWrap(True)
        self.detail_result.setContentsMargins(12,10,12,10)
        self.detail_result.setTextFormat(Qt.PlainText)
        self.detail_result_scroll = QScrollArea(); self.detail_result_scroll.setObjectName('projectResultScroll'); self.detail_result_scroll.setWidgetResizable(True)
        self.detail_result_scroll.setFrameShape(QScrollArea.NoFrame)
        self.detail_result_scroll.setMinimumHeight(90); self.detail_result_scroll.setMaximumHeight(120)
        self.detail_result.setSizePolicy(QSizePolicy.Ignored,QSizePolicy.Preferred)
        self.detail_result_scroll.setWidget(self.detail_result)
        detail.addWidget(self.detail_result_scroll)
        self.detail_values = {}
        content.addWidget(self.detail_panel)
        lay.addLayout(content,1)
        self.tbl.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.tbl.setSelectionMode(QAbstractItemView.SingleSelection)
        self.tbl.itemSelectionChanged.connect(self._refresh_project_detail)
        self.tbl.verticalHeader().setDefaultSectionSize(47)
        self.tbl.horizontalHeader().setFixedHeight(43)
        self.tbl.setAlternatingRowColors(False)
        self.actions_bar = QWidget(); self.actions_bar.setObjectName('actionsBar')
        btns = QHBoxLayout(self.actions_bar); btns.setContentsMargins(20,12,20,10); btns.setSpacing(9)
        self.btn_create = QPushButton('선택 실행'); self.btn_create.setObjectName('primaryButton')
        self.btn_create.clicked.connect(lambda: self._on_create_clicked(dry_run=False))
        self.btn_preview = QPushButton('요청 검증')
        self.btn_preview.clicked.connect(lambda: self._on_create_clicked(dry_run=True))
        self.btn_complete = QPushButton('BUS 완료 재확인')
        self.btn_complete.clicked.connect(lambda: self._on_create_clicked(completion_only=True))

        self.status_lbl = QLabel('',self); self.status_lbl.hide()
        self.prg = QProgressBar(self)
        self.prg.setRange(0, 0)
        self.prg.setFixedSize(0, 0)
        self.prg.hide()

        self._last_worker_msg = ""
        self.worker_thread = None
        self.worker = None
        self.chk_auto_complete = QCheckBox('성공 후 BUS 자동 완료')
        self.chk_auto_complete.setChecked(True)
        self.btn_stop = QPushButton('중지')
        self.btn_stop.clicked.connect(lambda: self.worker.stop() if self.worker else None)
        self.selection_label = QLabel('선택 0건'); self.selection_label.setObjectName('selectionLabel')
        btns.addWidget(self.selection_label)
        btns.addWidget(self.chk_auto_complete); btns.addStretch()
        btns.addWidget(self.btn_preview); btns.addWidget(self.btn_create); btns.addWidget(self.btn_complete); btns.addWidget(self.btn_stop)
        self.btn_stop.setObjectName('dangerButton')
        self.btn_stop.setEnabled(False)
        self.tbl.model().rowsRemoved.connect(self._refresh_selection)
        self.tbl.model().modelReset.connect(self._refresh_selection)
        self._refresh_selection()

    def _refresh_selection(self, *args):
        selected = sum(self._is_row_checked(row) for row in range(self.tbl.rowCount()))
        self.selection_label.setText(f'선택 {selected}건')
        self.host._sync_selection_label_width(self.selection_label)
        self.tbl.horizontalHeader().setChecked(bool(self.tbl.rowCount()) and selected == self.tbl.rowCount())

    def _can_delete_rows(self):
        return (self.tbl.contextMenuPolicy() == Qt.CustomContextMenu
                and self.btn_manual.isEnabled() and self.worker_thread is None
                and not self.host._closing and self.host.api_thread is None
                and not self.host._waiting_for_bus and not self.host.session.is_busy())

    def _open_table_menu(self, pos):
        if not self._can_delete_rows():
            return
        row = self.tbl.indexAt(pos).row()
        menu = QMenu(self)
        delete_row = menu.addAction('행 삭제')
        delete_checked = menu.addAction('선택 행 삭제')
        delete_all = menu.addAction('전체 삭제')
        delete_row.setEnabled(row >= 0)
        delete_checked.setEnabled(any(self._is_row_checked(r) for r in range(self.tbl.rowCount())))
        delete_all.setEnabled(self.tbl.rowCount() > 0)
        action = menu.exec_(self.tbl.viewport().mapToGlobal(pos))
        if action == delete_row:
            self._delete_row(row)
        elif action == delete_checked:
            self._delete_checked_rows()
        elif action == delete_all:
            self._delete_all_rows()

    def _is_row_checked(self, row):
        holder = self.tbl.cellWidget(row, 0)
        checkbox = holder.findChild(QCheckBox) if holder else None
        return bool(checkbox and checkbox.isChecked())

    def _after_delete(self):
        if not self.tbl.rowCount():
            self.tbl.setCurrentCell(-1, -1)
            self.tbl.set_empty_message('표시할 요청이 없습니다. BUS 신규조회 또는 수동 추가를 사용하세요.')
        self._refresh_project_detail()

    def _delete_row(self, row):
        if self._can_delete_rows() and 0 <= row < self.tbl.rowCount():
            self.tbl.removeRow(row)
            self._after_delete()

    def _delete_checked_rows(self):
        if not self._can_delete_rows():
            return
        for row in range(self.tbl.rowCount()-1, -1, -1):
            if self._is_row_checked(row):
                self.tbl.removeRow(row)
        self._after_delete()

    def _delete_all_rows(self):
        if not self._can_delete_rows():
            return
        self.tbl.setRowCount(0)
        self._after_delete()

    def _set_busy(self, on: bool, msg: str = ""):
        self.status_lbl.setText(msg or "")
        self.prg.hide()
        self.btn_manual.setEnabled(not on)
        self.btn_create.setEnabled(not on)
        self.btn_refresh.setEnabled(not on)
        self.btn_preview.setEnabled(not on)
        self.btn_complete.setEnabled(not on)
        self.tbl.setContextMenuPolicy(Qt.PreventContextMenu if on else Qt.CustomContextMenu)
        self.chk_auto_complete.setEnabled(not on)
        self.btn_stop.setEnabled(on)

    def _on_worker_progress(self, cur: int, total: int, message: str):
        self.status_lbl.setText(message)
        self._last_worker_msg = message
        self.host._set_progress_status(message,done=cur,total=total,running=True)

    def _on_worker_finished(self,ok_cnt,fail_cnt):
        prefix = (f'BUS 완료 {ok_cnt} / 보류·실패 {fail_cnt}' if getattr(self,'_completion_running',False)
                  else '요청 검증 완료' if getattr(self,'_validation_running',False) else f'실제 성공 {ok_cnt} / 실패 {fail_cnt}')
        self.status_lbl.setText(prefix+' - '+self._last_worker_msg)
        self._refresh_project_detail()


    def _manual_add(self,proj,name):
        if not self.tbl.columnCount():
            self.set_data(['프로젝트코드','프로젝트명'],[])
        row = self.tbl.rowCount()
        self.tbl.insertRow(row)
        checkbox = QCheckBox(); checkbox.setChecked(True)
        checkbox.toggled.connect(self._refresh_selection)
        holder = QWidget(); layout = QHBoxLayout(holder)
        layout.addWidget(checkbox); layout.setContentsMargins(0,0,0,0)
        self.tbl.setCellWidget(row,0,holder)
        self.tbl.setItem(row,self._hidx['프로젝트코드'],QTableWidgetItem(proj))
        name_col = self._hidx.get('프로젝트명',self._hidx.get('과제명',self._hidx.get('제목')))
        if name_col is None:
            name_col = self.tbl.columnCount()
            self.tbl.setColumnCount(name_col+1)
            self.tbl.setHorizontalHeaderItem(name_col,QTableWidgetItem('프로젝트명'))
            self._hidx['프로젝트명'] = name_col
        self.tbl.setItem(row,name_col,QTableWidgetItem(name))
        self.tbl.item(row,self._hidx['프로젝트코드']).setData(Qt.UserRole,
            dict(source='test-client',request_id=new_client_request_id(),bus_done=False))
        self.tbl.setItem(row,self._status_col,QTableWidgetItem('실행 대기'))
        self.tbl.setCurrentCell(row,self._hidx['프로젝트코드'])
        self._refresh_project_detail()
        self._refresh_selection()
        self.status_lbl.setText('수동 요청을 추가했습니다. 요청 검증 또는 선택실행을 누르세요.')


    def _on_manual_clicked(self):
        dlg = ManualNewRequestDialog(self)
        if dlg.exec_() != QDialog.Accepted:
            return

        proj, name = dlg.values()
        if not proj:
            QMessageBox.warning(self, "알림", "프로젝트 코드를 입력하세요.")
            return

        self._manual_add(proj, name)

    def _hide_empty_columns(self, always_show: set[str] | None = None, skip_col_idx: int | None = None):
        t = self.tbl
        if t is None:
            return
        rows, cols = t.rowCount(), t.columnCount()
        always_show = always_show or set()

        for c in range(cols):
            t.setColumnHidden(c, False)

        for c in range(cols):
            if skip_col_idx is not None and c == skip_col_idx:
                continue
            header = t.horizontalHeaderItem(c).text() if t.horizontalHeaderItem(c) else ""
            if header in always_show:
                continue
            has_data = False
            for r in range(rows):
                it = t.item(r, c)
                if it and it.text().strip():
                    has_data = True
                    break
            t.setColumnHidden(c, not has_data)

    def set_data(self, header, rows):
        try:
            self.tbl.blockSignals(True)

            col_select = 0
            col_count = 1 + len(header)

            self.tbl.clear()
            self.tbl.setColumnCount(col_count)
            self.tbl.setRowCount(len(rows))

            it_sel = QTableWidgetItem("")
            it_sel.setToolTip("선택")
            self.tbl.setHorizontalHeaderItem(col_select, it_sel)
            self.tbl.horizontalHeader().setSectionResizeMode(col_select, QHeaderView.Fixed)
            self.tbl.setColumnWidth(col_select, 28)

            for i, h in enumerate(header, start=1):
                self.tbl.setHorizontalHeaderItem(i, QTableWidgetItem(h or ""))

            for r, row in enumerate(rows):
                cb = QCheckBox(self.tbl)
                cb.setChecked(True)
                cb.toggled.connect(self._refresh_selection)
                wrap = QWidget(self.tbl)
                box = QHBoxLayout(wrap); box.setContentsMargins(0,0,0,0)
                box.addWidget(cb, 0, Qt.AlignCenter)
                self.tbl.setCellWidget(r, col_select, wrap)

                for c, v in enumerate(row, start=1):
                    it = QTableWidgetItem(str(v) if v is not None else "")
                    it.setTextAlignment(Qt.AlignCenter)
                    self.tbl.setItem(r, c, it)

            for c in range(1, col_count):
                self.tbl.horizontalHeader().setSectionResizeMode(c, QHeaderView.ResizeToContents)
            self.tbl.horizontalHeader().setStretchLastSection(True)

            self._hide_empty_columns(always_show={"프로젝트코드","프로젝트명"}, skip_col_idx=0)
        finally:
            self.tbl.blockSignals(False)

        self._hidx = { (self.tbl.horizontalHeaderItem(i).text() or ""): i
                       for i in range(1, self.tbl.columnCount()) }
        col = self._hidx.get('프로젝트코드',-1)
        name_col = self._hidx.get('프로젝트명',self._hidx.get('과제명',self._hidx.get('제목',-1)))
        for row in range(self.tbl.rowCount()):
            item = self.tbl.item(row,col) if col >= 0 else None
            name_item = self.tbl.item(row,name_col) if name_col >= 0 else None
            if item:
                values = dict(proj=item.text(),project_name=name_item.text() if name_item else '')
                identity = bus_row_identity(header,rows[row])
                try:
                    request_id = make_bus_request_id(values,identity)
                except ValueError:
                    request_id = new_client_request_id()
                item.setData(Qt.UserRole,dict(source='bus',request_id=request_id,bus_done=False))
        self._configure_project_summary()
        self._refresh_selection()

    def _configure_project_summary(self):
        self._status_col = self.tbl.columnCount()
        self.tbl.setColumnCount(self._status_col+1)
        self.tbl.setHorizontalHeaderItem(self._status_col,QTableWidgetItem('상태'))
        visible = {self._hidx.get('프로젝트코드'),
                   self._hidx.get('프로젝트명',self._hidx.get('과제명',self._hidx.get('제목')))}
        for column in range(1,self._status_col):
            self.tbl.setColumnHidden(column,column not in visible)
            if column in visible:
                self.tbl.horizontalHeader().setSectionResizeMode(column,QHeaderView.Stretch)
        self.tbl.horizontalHeader().setSectionResizeMode(self._status_col,QHeaderView.Interactive)
        self.tbl.setColumnWidth(self._status_col,130)
        for row in range(self.tbl.rowCount()):
            self.tbl.setItem(row,self._status_col,QTableWidgetItem('실행 대기'))
        if self.tbl.rowCount():
            self.tbl.setCurrentCell(0,self._hidx.get('프로젝트코드',1))
        self._refresh_project_detail()

    @staticmethod
    def _project_result_label(result):
        if result.get('bus_done'):
            return 'BUS 완료'
        if result.get('bus_processing'):
            return 'BUS 완료 처리 중'
        if result.get('completion_check_failed'):
            return 'BUS 완료 확인 불가'
        if result.get('completion_only'):
            return 'BUS 완료 보류'
        return {'submitting':'요청 접수 중','validating':'요청 검증 중','completion_checking':'서버 재조회 중',
                'partially_succeeded':'부분 성공','simulated':'모의 실행 완료','preview':'실행 계획','succeeded':'서버 성공','failed':'실패',
                'cancelled':'중지됨','running':'실행 중','queued':'실행 대기'}.get(result.get('status'),'실행 대기')

    def _refresh_project_process(self, meta, job):
        steps = {step.get('type'):step for step in (job.get('steps') or []) if isinstance(step,dict)}
        states = ['queued']*8
        names = list(self.project_step_names)
        if job.get('job_id'):
            states[0] = 'succeeded'
        action = meta.get('project_action')
        if action in {'validating','submitting'}:
            states[0] = 'running'
        validation = meta.get('validation_result') if meta.get('show_validation') else None
        if validation:
            names[0] = '요청 검증'
            states[0] = {'passed':'succeeded','failed':'failed','warning':'warning','unknown':'warning'}.get(validation.get('validation_status'),'queued')
        if action == 'validating':
            names[0] = '요청 검증'
        for index,kind in enumerate(self.project_step_types[1:-1],1):
            step = steps.get(kind,{})
            state = step.get('status','queued')
            if job.get('executor_mode')=='mock' and state in {'succeeded','simulated'}:
                state = 'simulated'; names[index] += ' · 모의 실행'
            states[index] = state if state in {'succeeded','running','failed','cancelled','skipped','simulated'} else 'queued'
        if meta.get('bus_done') and meta.get('source')=='bus':
            states[-1] = 'succeeded'
        elif meta and meta.get('source')!='bus':
            states[-1] = 'not_applicable'; names[-1] += ' · 수동 요청 제외'
        elif action == 'completion_checking' or job.get('bus_processing'):
            states[-1] = 'running'
        elif meta.get('completion_check_error') or (job.get('bus_completion_attempted') and not job.get('bus_done')):
            states[-1] = 'failed'
        for index,(icon,label,state) in enumerate(zip(self.project_step_icons,self.project_step_labels,states)):
            icon.setText('✓' if state=='succeeded' else '-' if state in {'skipped','cancelled','simulated','not_applicable'} else str(index+1))
            suffix = ' · 미실행' if state=='skipped' else ' · 중지됨' if state=='cancelled' else ''
            label.setText(names[index]+suffix)
            for widget in (icon,label):
                widget.setProperty('state',state)
                widget.style().unpolish(widget); widget.style().polish(widget)

    def _refresh_project_detail(self):
        saved_scroll = self.detail_scroll.verticalScrollBar().value()
        for form in (self.detail_form,self.server_detail_form):
            while form.rowCount():
                form.removeRow(0)
        self.detail_values = {}
        row = self.tbl.currentRow()
        code_item = self.tbl.item(row,getattr(self,'_hidx',{}).get('프로젝트코드',-1)) if row>=0 else None
        meta = dict(code_item.data(Qt.UserRole) or {}) if code_item else {}
        token = meta.get('request_id') or (code_item.text() if code_item else None)
        if token != self._detail_token:
            saved_scroll = 0
        self._detail_token = token
        self._detail_generation += 1; generation = self._detail_generation
        def add_field(form,name,text):
            text = str(text) if text is not None and str(text) else '-'
            if form is self.server_detail_form:
                value = DetailTextValue(text)
            else:
                value = QLabel(text)
                value.setWordWrap(True); value.setTextFormat(Qt.PlainText)
                value.setMinimumWidth(0); value.setSizePolicy(QSizePolicy.Ignored,QSizePolicy.Preferred)
                value.setTextInteractionFlags(Qt.TextSelectableByMouse)
            label = QLabel(name); label.setObjectName('mutedLabel'); label.setWordWrap(True)
            form.addRow(label,value); self.detail_values[name] = value
        for name,column in getattr(self,'_hidx',{}).items():
            item = self.tbl.item(row,column) if row>=0 else None
            if row>=0:
                add_field(self.detail_form,name,item.text() if item else '-')
        job = meta.get('live_result') if meta.get('project_action') else meta.get('result') or {}
        job = job or {}
        self._refresh_project_process(meta,job)
        steps = {step.get('type'):step for step in (job.get('steps') or []) if isinstance(step,dict)}
        extra = [('원본 요청번호',meta.get('request_id')),('작업 ID',job.get('job_id')),
                 ('보안 그룹',steps.get('ensure_ad_group',{}).get('target')),
                 ('프로젝트 경로',steps.get('ensure_project_folder',{}).get('target')),
                 ('템플릿',steps.get('ensure_project_folder',{}).get('details',{}).get('template_path')),
                 ('서버 상태',self._project_result_label(dict(status=job['status'])) if job.get('status') else None),('시작 시각',job.get('started_at')),('완료 시각',job.get('finished_at'))]
        for name,value in extra:
            if value:
                if name in {'시작 시각','완료 시각'}:
                    try:
                        when = datetime.datetime.fromisoformat(str(value).replace('Z','+00:00'))
                        if when.tzinfo is None: when = when.replace(tzinfo=datetime.timezone.utc)
                        value = when.astimezone().strftime('%Y-%m-%d %H:%M:%S')
                    except (ValueError,TypeError): pass
                add_field(self.server_detail_form,name,value)
        result = (meta.get('validation_result') or {}) if meta.get('show_validation') else job
        if row<0:
            self.detail_result.setText('요청을 선택하세요')
        elif meta.get('completion_check_error'):
            self.detail_result.setText('BUS 완료 확인 불가\n'+meta['completion_check_error'])
        else:
            title = ('요청 검증 중' if meta.get('project_action')=='validating' else
                     '서버 결과 재조회 중' if meta.get('project_action')=='completion_checking' else
                     validation_label(result) if result.get('status')=='validated' else self._project_result_label(result))
            lines = [title]
            running = next((step for step in steps.values() if step.get('status')=='running'),None)
            if running:
                lines.append('현재 단계: '+str(running.get('description') or running.get('type')))
            for key,label in [('error_code','오류'),('error_message','상세'),('message','메시지')]:
                if result.get(key): lines.append(f'{label}: {result[key]}')
            if result.get('status')=='validated':
                lines.append('검증 결과는 작업 로그에서 확인하세요. 권한은 변경하지 않았습니다.')
            self.detail_result.setText('\n'.join(lines))
        def restore_scroll():
            if generation==self._detail_generation:
                self.detail_scroll.verticalScrollBar().setValue(saved_scroll)
        QTimer.singleShot(0,restore_scroll)

    def _on_project_progress(self,row,result):
        item = self.tbl.item(row,self._hidx.get('프로젝트코드',-1))
        if not item:
            return
        meta = dict(item.data(Qt.UserRole) or {})
        meta.pop('show_validation',None); meta.pop('completion_check_error',None)
        meta.update(live_result=dict(result),project_action=result.get('status','running'))
        item.setData(Qt.UserRole,meta)
        self.tbl.setItem(row,self._status_col,QTableWidgetItem(self._project_result_label(result)))
        if result.get('status') in {'submitting','validating','completion_checking'}:
            self.tbl.setCurrentCell(row,self._hidx['프로젝트코드'])
        if self.tbl.currentRow()==row:
            self._refresh_project_detail()

    def _toggle_all_rows(self, checked: bool):
        t = self.tbl
        for r in range(t.rowCount()):
            w = t.cellWidget(r, 0)
            if not w:
                continue
            cb = w.findChild(QCheckBox)
            if cb:
                cb.setChecked(checked)

    def _selected_targets(self):
        t = self.tbl
        items = []
        code_col = self._hidx.get("프로젝트코드", -1)
        name_col = self._hidx.get("프로젝트명", -1)

        if code_col < 0:
            QMessageBox.warning(self, "알림", "헤더 '프로젝트코드'를 찾을 수 없습니다.")
            return items

        for r in range(t.rowCount()):
            w = t.cellWidget(r, 0)
            cb = w.findChild(QCheckBox) if w else None
            if not (cb and cb.isChecked()):
                continue
            proj = (t.item(r, code_col).text().strip() if t.item(r, code_col) else "")
            pname = (t.item(r, name_col).text().strip() if (name_col >= 0 and t.item(r, name_col)) else "")
            if proj:
                items.append((proj, pname))
        return items


    def _on_create_clicked(self, dry_run=False, completion_only=False):
        parent = self.host
        if parent._closing or parent._waiting_for_bus or self.worker_thread is not None or parent.api_thread is not None or parent.session.is_busy():
            return
        if not parent.api_config.get('api_key'):
            parent._open_api_settings()
            if not parent.api_config.get('api_key'):
                return
        items = []
        code_col = self._hidx.get('프로젝트코드',-1)
        name_col = self._hidx.get('프로젝트명',self._hidx.get('과제명',self._hidx.get('제목',-1)))
        if code_col < 0:
            return
        for row in range(self.tbl.rowCount()):
            holder = self.tbl.cellWidget(row,0)
            checkbox = holder.findChild(QCheckBox) if holder else None
            code_item = self.tbl.item(row,code_col)
            if not checkbox or not checkbox.isChecked() or not code_item:
                continue
            meta = dict(code_item.data(Qt.UserRole) or {})
            if completion_only and (meta.get('source')!='bus' or meta.get('bus_done') or not meta.get('request_id')):
                parent._log(f'{row+1}행: 미완료 BUS 요청이 아니므로 완료 재확인에서 제외했습니다.')
                continue
            if meta.get('bus_done') and not dry_run:
                continue
            project = code_item.text().strip()
            name_item = self.tbl.item(row,name_col) if name_col >= 0 else None
            name = name_item.text().strip() if name_item else ''
            try:
                canonical_project_code(project)
            except ValueError as exc:
                QMessageBox.warning(self,'입력 확인',str(exc)); return
            items.append(dict(proj=project,name=name,row=row,**meta))
        if not items:
            return
        parent.stop_requested = False
        parent._progress_started_at = time.monotonic()
        parent._progress_elapsed = 0.0
        parent._progress_detail = 'BUS 완료 재확인 중' if completion_only else '요청 검증 중' if dry_run else '프로젝트 생성 중'
        parent.done_jobs = 0; parent.total_jobs = len(items)
        parent._set_running_ui(True)
        self.worker_thread = QThread(self)
        self.worker = CreateWorker(items,dict(parent.creds),parent._api_client(),
                                   dry_run=dry_run, validate_only=dry_run, completion_only=completion_only,
                                   auto_complete=self.chk_auto_complete.isChecked(),
                                   requested_by=parent._requester())
        self.worker.moveToThread(self.worker_thread)
        self.worker_thread.started.connect(self.worker.run)
        self.worker.progress.connect(self._on_worker_progress)
        self.worker.itemResult.connect(self._on_project_result)
        self.worker.itemProgress.connect(self._on_project_progress)
        self.worker.finished.connect(self._on_worker_finished)
        self.worker.finished.connect(self.worker.deleteLater)
        self.worker.finished.connect(self.worker_thread.quit,type=Qt.DirectConnection)
        self.worker_thread.finished.connect(self._project_thread_finished)
        self.worker_thread.finished.connect(self.worker_thread.deleteLater)
        self._validation_running = dry_run
        self._completion_running = completion_only
        self._set_busy(True,'BUS 완료 재확인 시작' if completion_only else '요청 검증 시작' if dry_run else 'API 프로젝트 작업 시작')
        if parent.stop_requested or parent._closing:
            self.worker.stop()
        self.worker_thread.start()


    def _on_project_result(self,row,result):
        verified = result.get('executor_mode') in {'mock','powershell'} or result.get('status')=='preview'
        self.host._set_api_connection_state(verified,result)
        col = self._hidx.get('프로젝트코드',-1)
        item = self.tbl.item(row,col) if col >= 0 else None
        if item:
            meta = dict(item.data(Qt.UserRole) or {})
            meta.pop('completion_check_error',None)
            meta.pop('project_action',None); meta.pop('live_result',None)
            if result.get('completion_check_failed'):
                meta.pop('show_validation',None)
                meta['completion_check_error'] = result.get('error_message','서버 결과 확인 불가')
            elif result.get('status') == 'validated':
                meta.update(validation_result=result,show_validation=True)
            else:
                meta.pop('show_validation',None)
                meta.update(result=result,bus_done=result.get('bus_done',False))
            item.setData(Qt.UserRole,meta)
            item.setToolTip(json.dumps(result,ensure_ascii=False,indent=2))
        self.tbl.setItem(row,self._status_col,QTableWidgetItem(validation_label(result) if result.get('status')=='validated' else self._project_result_label(result)))
        self._refresh_project_detail()
        if result.get('status') == 'validated':
            self.host._log_validation(result,f'프로젝트 {item.text() if item else row+1}')
            return
        summary = (f'프로젝트 {item.text() if item else row+1}: '+self._project_result_label(result)+' · '+(result.get('message') or result.get('error_message') or '')
                   if result.get('completion_only') else self.host._api_log_summary(result,f'프로젝트 {item.text() if item else row+1}'))
        self.host._log('프로젝트 결과: '+json.dumps(result,ensure_ascii=False,indent=2),summary=summary)


    def _project_thread_finished(self):
        for row in range(self.tbl.rowCount()):
            item = self.tbl.item(row,self._hidx.get('프로젝트코드',-1))
            meta = dict(item.data(Qt.UserRole) or {}) if item else {}
            if meta.pop('project_action',None):
                meta.pop('live_result',None)
                item.setData(Qt.UserRole,meta)
                self.tbl.setItem(row,self._status_col,QTableWidgetItem(self._project_result_label(meta.get('result') or {})))
        self._refresh_project_detail()
        self.worker = None
        self.worker_thread = None
        self._set_busy(False,self.status_lbl.text())
        self.host._set_running_ui(False)


    def closeEvent(self,event):
        if self.worker_thread is not None and self.worker_thread.isRunning():
            self.worker.stop()
            event.ignore()
        else:
            event.accept()


    def reject(self):
        if self.worker_thread is not None:
            self.worker.stop()
            return
        if not self.isWindow():
            self.host.tabs.setCurrentIndex(0)
            return
        super().reject()


class CreateWorker(QObject):
    itemProgress = pyqtSignal(int,dict)
    itemResult = pyqtSignal(int,dict)
    progress = pyqtSignal(int, int, str)
    finished = pyqtSignal(int, int)
    error = pyqtSignal(str)

    def __init__(self,items,creds,client,dry_run=False,auto_complete=True,requested_by='client-operator',parent=None,validate_only=False,completion_only=False):
        super().__init__(parent)
        self.items,self.creds,self.client = items,creds,client
        self.dry_run,self.auto_complete,self.requested_by = dry_run,auto_complete,requested_by
        self.validate_only = validate_only
        if completion_only and (dry_run or validate_only):
            raise ValueError('BUS 완료 재확인은 요청 검증·Preview와 함께 실행할 수 없습니다.')
        self.completion_only = completion_only
        self._stop = False
        self._active_job = None


    def stop(self):
        self._stop = True
        if self._active_job is not None:
            self._active_job.stop()


    def _bus_click_process(self, proj_code: str) -> (bool, str):
        if self._stop:
            return False, '사용자 중지'
        d = None
        stage = "INIT"
        bus_msg_title = ""
        bus_msg_body = ""
        try:
            options = make_hidden_chrome_options()
            d = webdriver.Chrome(options=options)

            def normalize_text(s: str) -> str:
                s = (s or "").replace("\u00a0", " ").strip()
                return re.sub(r"\s+", " ", s)

            def set_select_value_and_fire(driver, select_id, value):
                driver.execute_script("""
                    var s = document.getElementById(arguments[0]);
                    if (s) {
                        s.value = arguments[1];
                        try { s.dispatchEvent(new Event('change', {bubbles:true})); } catch(e) {}
                        if (typeof AllBtnYn === 'function') { try { AllBtnYn(); } catch(e) {} }
                    }
                """, select_id, value)

            def click_search_manual(driver):
                try:
                    WebDriverWait(driver, 5).until(
                        EC.element_to_be_clickable((By.ID, "btnSearch"))
                    ).click()
                except Exception:
                    driver.switch_to.default_content()
                    WebDriverWait(driver, 5).until(
                        EC.element_to_be_clickable((By.ID, "btnSearch"))
                    ).click()
                    WebDriverWait(driver, 10).until(
                        EC.presence_of_all_elements_located((By.TAG_NAME, "iframe"))
                    )
                    driver.switch_to.frame(driver.find_elements(By.TAG_NAME, "iframe")[0])

            def find_swal_popup(driver):
                try:
                    driver.switch_to.default_content()
                except Exception:
                    pass

                popups = driver.find_elements(By.CSS_SELECTOR, "div.swal2-container.swal2-shown")
                if popups:
                    return popups[0]

                try:
                    frames = driver.find_elements(By.TAG_NAME, "iframe")
                except Exception:
                    frames = []

                for fr in frames:
                    try:
                        driver.switch_to.frame(fr)
                        popups = driver.find_elements(By.CSS_SELECTOR, "div.swal2-container.swal2-shown")
                        if popups:
                            return popups[0]
                    except Exception:
                        pass
                    try:
                        driver.switch_to.default_content()
                    except Exception:
                        pass

                try:
                    driver.switch_to.default_content()
                except Exception:
                    pass
                return None

            def enter_first_iframe(driver, timeout=15, reload_url=None):
                for attempt in range(2):
                    try:
                        driver.switch_to.default_content()
                    except Exception:
                        pass

                    try:
                        WebDriverWait(driver, timeout).until(
                            lambda drv: len(drv.find_elements(By.TAG_NAME, "iframe")) > 0
                        )
                        frames = driver.find_elements(By.TAG_NAME, "iframe")
                        if frames:
                            driver.switch_to.frame(frames[0])
                            return True
                    except Exception:
                        pass

                    if reload_url and attempt == 0:
                        try:
                            driver.get(reload_url)
                        except Exception:
                            pass
                        time.sleep(0.5)

                return False

            def refocus_first_iframe(driver, reload_url=None):
                return enter_first_iframe(driver, timeout=5, reload_url=reload_url)

            stage = "OPEN_LOGIN_PAGE"
            d.get(BUS_LOGIN_URL)
            WebDriverWait(d, 20).until(EC.presence_of_element_located((By.ID, "windowsaccount")))

            stage = "INPUT_IDPW"
            d.find_element(By.ID, "windowsaccount").clear()
            d.find_element(By.ID, "windowsaccount").send_keys(self.creds["id"])
            d.find_element(By.ID, "password").clear()
            d.find_element(By.ID, "password").send_keys(self.creds["pw"])

            stage = "CLICK_LOGIN"
            d.find_element(By.ID, "btnLogin").click()
            WebDriverWait(d, 20).until(EC.url_contains("Common"))

            stage = "OPEN_NEW_PAGE"
            d.get(BUS_NEW_URL)

            stage = "ENTER_IFRAME"
            if not enter_first_iframe(d, reload_url=BUS_NEW_URL):
                return False, f"IFRAME_ENTER_FAIL @ {stage}"

            try:
                pre_info = d.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
            except Exception:
                pre_info = ""

            stage = "SET_FILTER"
            try:
                set_select_value_and_fire(d, "processYn", "N")
            except Exception:
                pass

            stage = "CLICK_SEARCH"
            if not refocus_first_iframe(d, reload_url=BUS_NEW_URL):
                return False, f"IFRAME_LOST_BEFORE_SEARCH @ {stage}"
            click_search_manual(d)

            stage = "WAIT_FILTER_APPLY"
            try:
                WebDriverWait(d, 7).until(
                    lambda x: (
                        x.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                        if x.find_elements(By.CSS_SELECTOR, ".dataTables_info") else ""
                    ) != (pre_info or "")
                )
            except Exception:
                time.sleep(0.8)

            stage = "WAIT_ROWS_BEFORE"

            def wait_rows_and_sample(driver):
                WebDriverWait(driver, 20).until(
                    EC.presence_of_all_elements_located((By.CSS_SELECTOR, "tbody tr"))
                )
                rows_inner = driver.find_elements(By.CSS_SELECTOR, "tbody tr")
                return rows_inner

            def find_target_row(driver, proj_code_val):
                proj_digits_val = re.sub(r"\D", "", proj_code_val or "")
                samples_inner = []
                target_inner = None
                for attempt_inner in range(3):
                    try:
                        rows_inner = driver.find_elements(By.CSS_SELECTOR, "tbody tr")
                        samples_inner = []
                        for idx, tr in enumerate(rows_inner):
                            tds = tr.find_elements(By.CSS_SELECTOR, "td")
                            full_txt = " ".join([normalize_text(td.text) for td in tds])
                            full_digits = re.sub(r"\D", "", full_txt)
                            if idx < 3:
                                samples_inner.append(full_txt or "(empty)")
                            if (proj_code_val and proj_code_val in full_txt) or (proj_digits_val and proj_digits_val in full_digits):
                                target_inner = tr
                                break
                        if target_inner is not None:
                            break
                    except StaleElementReferenceException:
                        target_inner = None
                        samples_inner = []
                        refocus_first_iframe(driver)
                        time.sleep(0.3)
                return target_inner, samples_inner, proj_digits_val

            rows = wait_rows_and_sample(d)
            if not rows:
                return False, f"ROWS_EMPTY_BEFORE @ {stage}"

            stage = "FIND_TARGET_ROW_BEFORE"
            target, samples, proj_digits = find_target_row(d, proj_code)

            if not target:
                stage = "RETRY_SEARCH_BEFORE"
                try:
                    pre_info_retry = d.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                except Exception:
                    pre_info_retry = ""
                if refocus_first_iframe(d, reload_url=BUS_NEW_URL):
                    click_search_manual(d)
                try:
                    WebDriverWait(d, 7).until(
                        lambda x: (
                            x.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                            if x.find_elements(By.CSS_SELECTOR, ".dataTables_info") else ""
                        ) != (pre_info_retry or "")
                    )
                except Exception:
                    time.sleep(0.8)
                if not enter_first_iframe(d, reload_url=BUS_NEW_URL):
                    return False, f"IFRAME_REENTER_FAIL @ {stage}"
                rows = wait_rows_and_sample(d)
                if not rows:
                    return False, f"ROWS_EMPTY_BEFORE @ {stage}"
                stage = "FIND_TARGET_ROW_BEFORE"
                target, samples, proj_digits = find_target_row(d, proj_code)

            if not target:
                return False, f"ROW_NOT_FOUND @ {stage}: proj={proj_code}, digits={proj_digits}, sample={samples}"

            stage = "FIND_PROCESS_BUTTON"
            try:
                btn = target.find_element(By.XPATH, ".//button[contains(.,'처리')]")
            except Exception:
                return False, f"BTN_NOT_FOUND @ {stage}"

            stage = "CLICK_PROCESS_BUTTON"
            if self._stop:
                return False, '사용자 중지 (BUS 완료 미실행)'
            try:
                d.execute_script("arguments[0].scrollIntoView({block:'center'});", btn)
                try:
                    btn.click()
                except ElementClickInterceptedException:
                    d.execute_script("arguments[0].click();", btn)
            except Exception as e:
                return False, f"BTN_CLICK_FAIL @ {stage}: {e}"

            stage = "WAIT_CONFIRM_POPUP"

            popup = None
            for _ in range(40):  # 최대 8초 (0.2 * 40)
                try:
                    popup = find_swal_popup(d)
                    if popup:
                        break
                except Exception:
                    popup = None
                time.sleep(0.2)

            if popup is None:
                return False, f"CONFIRM_POPUP_NOT_FOUND @ {stage}"

            stage = "READ_CONFIRM_MESSAGE"
            try:
                bus_msg_title = (popup.find_element(By.CSS_SELECTOR, ".swal2-title").text or "").strip()
            except Exception:
                bus_msg_title = ""
            try:
                bus_msg_body = (popup.find_element(By.CSS_SELECTOR, ".swal2-html-container").text or "").strip()
            except Exception:
                bus_msg_body = ""

            stage = "CLICK_CONFIRM_POPUP"
            if self._stop:
                return False, '사용자 중지 (BUS 확인 미실행)'
            try:
                ok_btn = WebDriverWait(popup, 5).until(
                    EC.element_to_be_clickable((By.CSS_SELECTOR, "button.swal2-confirm"))
                )
                try:
                    ok_btn.click()
                except ElementClickInterceptedException:
                    d.execute_script("arguments[0].click();", ok_btn)
            except Exception as e:
                return False, f"CONFIRM_POPUP_CLICK_FAIL @ {stage}: {e}"

            stage = "WAIT_CONFIRM_CLOSE"
            try:
                WebDriverWait(d, 10).until(lambda _:
                    find_swal_popup(d) is None
                )
            except Exception:
                pass

            stage = "RELOAD_IFRAME_AFTER_PROCESS"
            if not enter_first_iframe(d, reload_url=BUS_NEW_URL):
                return False, f"IFRAME_REENTER_FAIL @ {stage}"

            try:
                pre_info2 = d.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
            except Exception:
                pre_info2 = ""

            stage = "SET_FILTER_AFTER"
            try:
                set_select_value_and_fire(d, "processYn", "N")
            except Exception:
                pass

            stage = "CLICK_SEARCH_AFTER"
            if not refocus_first_iframe(d, reload_url=BUS_NEW_URL):
                return False, f"IFRAME_LOST_AFTER_SEARCH @ {stage}"
            click_search_manual(d)

            stage = "WAIT_FILTER_APPLY_AFTER"
            try:
                WebDriverWait(d, 7).until(
                    lambda x: (
                        x.find_element(By.CSS_SELECTOR, ".dataTables_info").text.strip()
                        if x.find_elements(By.CSS_SELECTOR, ".dataTables_info") else ""
                    ) != (pre_info2 or "")
                )
            except Exception:
                time.sleep(0.8)

            stage = "WAIT_ROWS_AFTER"
            try:
                WebDriverWait(d, 20).until(
                    EC.presence_of_all_elements_located((By.CSS_SELECTOR, "tbody tr"))
                )
                rows2 = d.find_elements(By.CSS_SELECTOR, "tbody tr")
            except Exception:
                rows2 = []

            stage = "VERIFY_ROW_REMOVED"
            proj_digits = re.sub(r"\D", "", proj_code or "")
            still_exists = False
            for tr in rows2:
                tds = tr.find_elements(By.CSS_SELECTOR, "td")
                full_txt = " ".join([normalize_text(td.text) for td in tds])
                full_digits = re.sub(r"\D", "", full_txt)
                if (proj_code and proj_code in full_txt) or (proj_digits and proj_digits in full_digits):
                    still_exists = True
                    break

            if still_exists:
                if bus_msg_title or bus_msg_body:
                    return False, f"ROW_STILL_EXISTS_AFTER_PROCESS @ {stage}: title={bus_msg_title}, body={bus_msg_body}"
                return False, f"ROW_STILL_EXISTS_AFTER_PROCESS @ {stage}"

            if bus_msg_title or bus_msg_body:
                return True, f"OK @ {stage}: title={bus_msg_title}, body={bus_msg_body}"
            return True, f"OK @ {stage}"

        except Exception as e:
            return False, f"EXCEPTION @ {stage}: {e}"

        finally:
            if d is not None:
                try:
                    d.quit()
                except Exception:
                    pass


    def _run_completion_only(self):
        completed = pending = 0
        try:
            for index,item in enumerate(self.items):
                if self._stop:
                    break
                request_id = item.get('request_id')
                if item.get('source')!='bus' or item.get('bus_done') or not request_id:
                    continue
                self.progress.emit(index,len(self.items),f"{item['proj']} 서버 결과 재조회")
                self.itemProgress.emit(item.get('row',index),dict(status='completion_checking'))
                try:
                    result = self.client.request('GET','/api/v1/project-jobs/by-request/'+quote(request_id,safe=''))
                    if (result.get('request_id')!=request_id or not result.get('job_id')
                            or result.get('project_code')!=canonical_project_code(item['proj'])
                            or result.get('project_name')!=item.get('name','').strip()):
                        raise ApiClientError('서버 프로젝트 작업이 현재 BUS 요청과 일치하지 않습니다.')
                    if self._stop:
                        break
                    result = dict(result,completion_only=True,bus_done=False)
                    if can_complete_bus(result,'bus',self._stop):
                        if not self.creds.get('id') or not self.creds.get('pw'):
                            message = 'BUS 계정 누락: 계정 설정 후 완료만 다시 시도하세요.'
                        else:
                            self.itemProgress.emit(item.get('row',index),dict(result,bus_processing=True))
                            ok,message = self._bus_click_process(item['proj'])
                            result['bus_completion_attempted'] = True
                            result['bus_done'] = bool(ok)
                    else:
                        message = '서버 실제 성공이 확인되지 않아 BUS 완료를 보류했습니다.'
                    result['message'] = message
                    completed += int(result['bus_done'])
                    pending += int(not result['bus_done'])
                except Exception as exc:
                    message = str(exc)
                    if self.client.api_key:
                        message = message.replace(self.client.api_key,'[API Key]')
                    result = dict(completion_only=True,completion_check_failed=True,
                                  status='completion_unknown',bus_done=False,error_message=message)
                    pending += 1
                self.itemResult.emit(item.get('row',index),result)
                self.progress.emit(index+1,len(self.items),f"{item['proj']}: "+('BUS 완료' if result.get('bus_done') else 'BUS 완료 보류'))
        except Exception as exc:
            message = str(exc).replace(self.client.api_key,'[API Key]') if self.client.api_key else str(exc)
            self.error.emit(message)
            pending += 1
        finally:
            self.finished.emit(completed,pending)

    def run(self):
        if self.completion_only:
            self._run_completion_only()
            return
        ok_count = fail_count = 0
        try:
            for index,item in enumerate(self.items):
                if self._stop:
                    break
                self.progress.emit(index,len(self.items),f"{item['proj']} API 작업 중")
                self.itemProgress.emit(item.get('row',index),dict(status='validating' if self.validate_only else 'submitting'))
                source = item.get('source','test-client')
                request_id = item.get('request_id') or new_client_request_id()
                payload = make_project_payload(item['proj'],item.get('name',''),request_id,source,self.requested_by)
                cached = item.get('result',{})
                if not self.dry_run and not self.validate_only and can_complete_bus(cached,source) and not item.get('bus_done'):
                    result = dict(cached)
                else:
                    self._active_job = ApiJobWorker(self.client,'project',payload,self.dry_run,validate_only=self.validate_only)
                    self._active_job.progress.connect(lambda result,row=item.get('row',index):self.itemProgress.emit(row,result))
                    if self._stop:
                        self._active_job.stop()
                    try:
                        result = self._active_job.execute()
                    except Exception as exc:
                        result = (dict(status='validated',validation_status='unknown',request_id=request_id,
                                       checks=[dict(id='request',label='실제 자원 검증',target='',status='unknown',message=str(exc))])
                                  if self.validate_only else dict(status='failed',error_message=str(exc)))
                    self._active_job = None
                result = dict(result)
                real_success = result.get('status')=='succeeded' and result.get('executor_mode')=='powershell'
                if real_success:
                    ok_count += 1
                elif result.get('status') not in {'preview','simulated','validated'}:
                    fail_count += 1
                message = f"{item['proj']} API {result.get('status')}: {result.get('error_message') or ''}"
                if self.auto_complete and not self.validate_only and can_complete_bus(result,source,self._stop or result.get('client_stopped',False)):
                    if not self.creds.get('id') or not self.creds.get('pw'):
                        result['bus_done'] = False
                        message += ' / BUS 계정 누락: 완료만 재시도 가능'
                    else:
                        self.itemProgress.emit(item.get('row',index),dict(result,bus_processing=True))
                        bus_ok,bus_message = self._bus_click_process(item['proj'])
                        result['bus_completion_attempted'] = True
                        result['bus_done'] = bus_ok
                        message += ' / BUS '+bus_message
                self.itemResult.emit(item.get('row',index),result)
                self.progress.emit(index+1,len(self.items),message)
        except Exception as exc:
            message = str(exc)
            if self.client.api_key:
                message = message.replace(self.client.api_key,'[API Key]')
            self.error.emit(message)
            fail_count += 1
        finally:
            self._active_job = None
            self.finished.emit(ok_count,fail_count)


# 알림/배지 기능 비활성화: 복원 가능하도록 기존 구현을 주석으로 보존합니다.
# class BadgeToolButton(QWidget):
#     def __init__(self, *args, **kwargs):
#         target = None
#         parent = None
#         if len(args) == 1:
#             parent = args[0]
#         elif len(args) >= 2:
#             if isinstance(args[0], QAbstractButton):
#                 target = args[0]
#                 parent = args[1]
#             else:
#                 parent = args[1]
#         else:
#             parent = kwargs.get("parent", None)

#         super().__init__(parent)
#         self._target = None
#         self._offset = QPoint(0, 1)
#         self._value = 0
#         self._lbl = QLabel(self)
#         self._lbl.setObjectName("notifBadge")
#         self._lbl.setAlignment(Qt.AlignCenter)
#         f = self._lbl.font()
#         f.setPointSize(7)
#         f.setBold(True)
#         self._lbl.setFont(f)
#         self._base_px = 7
#         self._apply_style(self._base_px)
#         self._lbl.setVisible(False)

#         self.setAttribute(Qt.WA_TransparentForMouseEvents, True)
#         self._lbl.setAttribute(Qt.WA_TransparentForMouseEvents, True)

#         if target is not None:
#             self.setTarget(target)

#     def setTarget(self, btn: QAbstractButton):
#         if self._target is not None:
#             try:
#                 self._target.removeEventFilter(self)
#             except Exception:
#                 pass
#         self._target = btn
#         self._lbl.setParent(self._target)
#         self._lbl.raise_()
#         if self._target is not None:
#             self._target.installEventFilter(self)
#         self._reposition()

#     def setBadge(self, n: int):
#         self._value = int(n or 0)
#         if self._value <= 0:
#             self._lbl.hide()
#             return
#         self._lbl.setText("99+" if self._value >= 100 else str(self._value))
#         self._lbl.adjustSize()
#         self._lbl.show()
#         self._lbl.raise_()
#         self._reposition()

#     def setOffset(self, dx: int, dy: int):
#         self._offset = QPoint(dx, dy)
#         self._reposition()

#     def setBadgeSize(self, px: int = 1):
#         self._base_px = px
#         self._apply_style(px)
#         self._lbl.adjustSize()
#         self._reposition()

#     def eventFilter(self, obj, ev):
#         if obj is self._target and ev.type() in (ev.Resize, ev.Move, ev.Show):
#             self._reposition()
#         return super().eventFilter(obj, ev)

#     def _reposition(self):
#         if not (self._target and self._lbl.isVisible()):
#             return
#         bw = max(self._base_px, self._lbl.width())
#         x = self._target.width() - bw + self._offset.x()
#         y = self._offset.y()
#         self._lbl.move(x, y)

#     def _apply_style(self, px: int):
#         r = px // 2
#         pad = max(1, px // 4)
#         maxw = px * 3

#         self._lbl.setStyleSheet(f"""
#             QLabel#notifBadge {{
#                 color: white;
#                 background: #E23;
#                 border-radius: {r}px;
#                 min-width: {px}px;
#                 min-height: {px}px;
#                 max-width: {maxw}px;
#                 padding-left: {pad}px;
#                 padding-right: {pad}px;
#             }}
#         """)

# Drawing helpers keep visual presentation separate from the API payload model.
class RequestDelegate(FocuslessItemDelegate):
    def __init__(self,host):
        super().__init__(host.table)
        self.host = host

    def paint(self,painter,option,index):
        opt = QStyleOptionViewItem(option)
        self.initStyleOption(opt,index)
        host = self.host
        if index.column() == host.COL_LV2 and host.table.isColumnHidden(host.COL_LV3):
            lv2 = host._get(index.row(),host.COL_LV2)
            lv3 = host._get(index.row(),host.COL_LV3)
            opt.text = f'{lv2} · {lv3}' if lv2 and lv3 else '-'
            host.table.style().drawControl(QStyle.CE_ItemViewItem,opt,painter,host.table)
            return
        if index.column() != host.COL_STATUS:
            super().paint(painter,opt,index)
            return
        text,tone = request_status_presentation(str(index.data() or ''))
        opt.text = ''
        host.table.style().drawControl(QStyle.CE_ItemViewItem,opt,painter,host.table)
        theme = THEMES.get(host.current_theme,THEMES['light'])
        colors = {'blue':theme['accent'],'green':theme['green'],'amber':theme['amber'],
                  'red':theme['red'],'muted':theme['muted']}
        painter.save()
        painter.setRenderHint(QPainter.Antialiasing)
        rect = option.rect
        painter.setPen(Qt.NoPen)
        painter.setBrush(QColor(colors[tone]))
        painter.drawEllipse(QPoint(rect.left()+15,rect.center().y()),4,4)
        painter.setPen(QColor(theme['text']))
        painter.setFont(option.font)
        text_rect = rect.adjusted(28,0,-7,0)
        painter.drawText(text_rect,Qt.AlignVCenter|Qt.AlignLeft,
                         option.fontMetrics.elidedText(text,Qt.ElideRight,text_rect.width()))
        painter.restore()


def validation_label(result):
    return {'passed':'요청 검증 통과','warning':'요청 검증 주의',
            'failed':'요청 검증 실패','unknown':'요청 검증 확인 불가'}.get(result.get('validation_status'),'요청 검증 확인 불가')


def request_status_presentation(status):
    if status.startswith('요청 검증'):
        return status, 'muted' if '통과' in status else 'blue' if status.endswith('중') else 'amber'
    if 'BUS 완료' == status:
        return 'BUS 완료','green'
    if '초과' in status or 'TIMEOUT' in status:
        return '시간 초과','amber'
    if '부분' in status:
        return '부분 성공','amber'
    if '실패' in status or '오류' in status:
        return '확인 필요','red'
    if '모의' in status:
        return '모의 완료','muted'
    if status == 'API 성공' or status.startswith('서버 성공'):
        return '서버 성공','green'
    if 'DryRun' in status or 'Preview' in status or '요청 검증' in status:
        return '요청 검증','muted'
    if '실행중' in status or '실행 중' in status:
        return '실행 중','blue'
    if '완료 처리중' in status:
        return 'BUS 처리 중','blue'
    if '취소' in status or '중지' in status:
        return '중지됨','muted'
    if status in {'대기',''}:
        return '실행 대기','muted'
    return status,'muted'


def folder_mark(color,size=46):
    pixmap = QPixmap(size,size)
    pixmap.fill(Qt.transparent)
    painter = QPainter(pixmap)
    painter.setRenderHint(QPainter.Antialiasing)
    painter.setPen(QPen(QColor(color),3,Qt.SolidLine,Qt.RoundCap,Qt.RoundJoin))
    painter.setBrush(Qt.NoBrush)
    path = QPainterPath()
    path.moveTo(6,11); path.lineTo(17,11); path.lineTo(22,17)
    path.lineTo(size-6,17); path.lineTo(size-6,size-8)
    path.lineTo(6,size-8); path.closeSubpath()
    painter.drawPath(path)
    painter.end()
    return pixmap


def themed_icon(name,color):
    pixmap = QIcon(resource_path('assets/'+name+'.svg')).pixmap(24,24)
    painter = QPainter(pixmap)
    painter.setCompositionMode(QPainter.CompositionMode_SourceIn)
    painter.fillRect(pixmap.rect(),QColor(color))
    painter.end()
    return QIcon(pixmap)


class ThemeSwitch(QCheckBox):
    def __init__(self):
        super().__init__()
        self.setFixedSize(36,22)
        self.setCursor(Qt.PointingHandCursor)

    def hitButton(self,position):
        return self.rect().contains(position)

    def paintEvent(self,event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing)
        painter.setPen(Qt.NoPen)
        painter.setBrush(QColor('#3b82f6' if self.isChecked() else '#4b8df8'))
        painter.drawRoundedRect(1,2,34,18,9,9)
        painter.setBrush(QColor('#ffffff'))
        painter.drawEllipse(19 if self.isChecked() else 3,4,14,14)


class AccessManager(QMainWindow):
    trigger_new_download = pyqtSignal()
    trigger_session_start = pyqtSignal()
    trigger_session_download = pyqtSignal()
    trigger_session_cancel = pyqtSignal()
    trigger_session_stop = pyqtSignal()
    trigger_session_process = pyqtSignal(object)
    # trigger_watcher_start = pyqtSignal()
    # trigger_watcher_stop  = pyqtSignal()
    startup_result = pyqtSignal(str,bool,str)
    # trigger_watcher_collect = pyqtSignal()
    COL_SELECT  = 0
    COL_KIND    = 1
    COL_REQTYPE = 2
    COL_USER    = 3
    COL_NAME    = 4
    COL_PROJ    = 5
    COL_LV2     = 6
    COL_LV3     = 7
    COL_DEPT    = 8
    COL_ROLE    = 9
    COL_STATUS  = 10

    EDITABLE_COLS = {COL_KIND, COL_REQTYPE, COL_USER, COL_NAME, COL_PROJ, COL_LV2, COL_LV3, COL_DEPT, COL_ROLE}

    def _show_help(self):
        HelpDialog(self).exec_()

    def apply_theme(self,mode):
        mode = mode if mode in THEMES else 'light'
        self.current_theme = mode
        t = THEMES[mode]
        qss = f'''
        QWidget {{background:{t['bg']};color:{t['text']};font-size:13px;}}
        QLabel,QCheckBox,QToolButton {{background:transparent;}}
        QWidget#appHeader,QWidget#connectionBar {{background:{t['panel']};border-bottom:1px solid {t['border']};}}
        QLabel#brandTitle {{font-size:23px;font-weight:700;}}
        QLabel#brandSubtitle {{font-size:12px;}}
        QLabel#mutedLabel,QLabel#summaryLabel {{color:{t['muted']};font-size:12px;}}
        QLabel#versionBadge {{background:{t['badge']};border:1px solid {t['border']};border-radius:5px;padding:4px 9px;font-size:11px;font-weight:600;}}
        QLabel#modeBadge {{background:{t['badge']};color:{t['muted']};border:1px solid {t['border']};border-radius:5px;padding:5px 10px;font-size:12px;}}
        QLabel#modeBadge[mode="powershell"] {{background:{t['amber_bg']};color:{t['amber']};border-color:{t['amber_border']};}}
        QLabel#modeBadge[mode="mock"] {{background:{t['blue_bg']};color:{t['accent']};}}
        QPushButton {{background:{t['panel']};border:1px solid {t['button_border']};border-radius:5px;padding:7px 12px;min-height:20px;}}
        QPushButton:hover:enabled {{background:{t['hover']};border-color:{t['accent']};}}
        QPushButton:pressed:enabled {{background:{t['select']};}}
        QPushButton#primaryButton {{background:{t['accent']};color:#ffffff;border-color:{t['accent']};font-weight:500;}}
        QPushButton#primaryButton:hover:enabled {{background:{t['accent_hover']};}}
        QPushButton#dangerButton {{color:{t['red']};border-color:{t['red']};min-width:52px;}}
        QPushButton#textButton {{border:0;background:transparent;color:{t['muted']};font-size:12px;padding:3px 4px;}}
        QPushButton:disabled, QPushButton#primaryButton:disabled, QPushButton#dangerButton:disabled, QPushButton#textButton:disabled {{color:{t['disabled']};background:{t['badge']};border:1px solid {t['border']};font-weight:400;}}
        QToolButton {{border:0;padding:4px;}}
        QToolButton#utilityButton {{font-size:23px;min-width:28px;}}
        QToolButton:hover:enabled {{color:{t['accent']};}}
        QToolButton:disabled, QToolButton#utilityButton:disabled, QToolButton#logToggle:disabled {{color:{t['disabled']};background:{t['badge']};}}
        QToolButton::menu-indicator {{image:none;}}
        QToolButton#logToggle {{font-size:13px;font-weight:600;padding:3px 0;}}
        QCheckBox {{spacing:7px;}}
        QCheckBox::indicator {{width:15px;height:15px;border:1px solid {t['button_border']};border-radius:3px;background:{t['panel']};}}
        QCheckBox::indicator:checked {{background:{t['accent']};border-color:{t['accent']};image:url({self._checkmark_path()});}}
        QCheckBox:disabled {{color:{t['disabled']};}}
        QCheckBox::indicator:disabled {{background:{t['badge']};border-color:{t['border']};}}
        QCheckBox::indicator:checked:disabled {{background:{t['button_border']};border-color:{t['button_border']};}}
        QCheckBox#themeSwitch::indicator {{width:31px;height:17px;border-radius:9px;background:{t['button_border']};image:none;}}
        QCheckBox#themeSwitch::indicator:checked {{background:{t['accent']};image:none;}}
        QWidget#workspaceTabRow {{background:{t['panel']};border-bottom:1px solid {t['border']};}}
        QTabBar#workspaceTabs {{background:{t['panel']};border:0;font-size:15px;font-weight:600;}}
        QTabBar::tab {{background:transparent;color:{t['text']};padding:15px 22px 13px;min-height:24px;border-bottom:3px solid transparent;}}
        QTabBar::tab:selected {{color:{t['accent']};border-bottom-color:{t['accent']};}}
        QTabBar::tab:disabled {{color:{t['disabled']};}}
        QWidget#requestToolbar {{background:{t['bg']};border:0;}}
        QWidget#projectDetailFields, QWidget#projectDetailBody, QWidget#projectProcess, QScrollArea#projectDetailScroll, QScrollArea#projectResultScroll {{background:{t['panel']};}}
        QWidget#projectDetailFields QLabel {{background:transparent;}}
        QLabel#requestCountBadge {{background:#ef3340;color:white;border:0;border-radius:5px;padding:0;font-size:11px;font-weight:600;}}
        QLabel#sectionTitle {{font-size:16px;font-weight:700;}}
        QTableWidget {{background:{t['panel']};alternate-background-color:{t['bg']};border:1px solid {t['border']};border-radius:5px;gridline-color:{t['grid']};selection-background-color:{t['select']};selection-color:{t['text']};}}
        QTableWidget QWidget {{background:transparent;}}
        QTableWidget::item {{padding:5px;border:0;}}
        QTableWidget::item:selected {{background:{t['select']};color:{t['text']};}}
        QHeaderView::section {{background:{t['header']};color:{t['text']};border:0;border-bottom:1px solid {t['border']};padding:7px 4px;font-size:12px;font-weight:600;}}
        QTableCornerButton::section {{background:{t['header']};border:1px solid {t['border']};}}
        QWidget#detailPanel {{background:{t['panel']};border:1px solid {t['border']};border-radius:5px;}}
        QLabel#stepIcon {{background:{t['badge']};color:{t['muted']};border-radius:12px;font-size:12px;}}
        QLabel#stepIcon[state="succeeded"] {{background:{t['green']};color:#ffffff;}}
        QLabel#stepIcon[state="running"] {{background:{t['accent']};color:#ffffff;}}
        QLabel#stepIcon[state="failed"] {{background:{t['red']};color:#ffffff;}}
        QLabel#stepLabel[state="running"] {{color:{t['accent']};font-weight:600;}}
        QLabel#stepLabel[state="failed"] {{color:{t['red']};}}
        QLabel#stepIcon[state="warning"] {{background:{t['amber']};color:#ffffff;}}
        QLabel#stepLabel[state="warning"] {{color:{t['amber']};}}
        QWidget#detailBanner {{background:{t['blue_bg']};border-radius:5px;}}
        QLabel#bannerTitle {{color:{t['accent']};font-weight:600;font-size:12px;}}
        QTextEdit#bannerBody {{background:{t['blue_bg']};color:{t['muted']};font-size:11px;border:0;padding:0;}}
        QWidget#logsPanel {{background:{t['panel']};border:1px solid {t['border']};border-radius:6px;}}
        QTextEdit#jobLog {{background:transparent;border:0;border-top:1px solid {t['border']};padding:7px 2px;color:{t['text']};font-size:12px;}}
        QWidget#actionsBar {{background:{t['bg']};}}
        QLabel#selectionLabel {{color:{t['accent']};font-weight:700;font-size:14px;}}
        QLineEdit,QComboBox,QSpinBox {{background:{t['panel']};border:1px solid {t['button_border']};border-radius:4px;padding:5px;selection-background-color:{t['select']};}}
        QMenu {{background:{t['panel']};border:1px solid {t['border']};padding:5px;}}
        QMenu::item {{padding:7px 18px;border-radius:3px;}}
        QMenu::item:selected {{background:{t['select']};}}
        QMenu::item:disabled {{color:{t['disabled']};background:transparent;}}
        QMenu::separator {{height:1px;background:{t['border']};margin:5px;}}
        QScrollBar:vertical {{background:{t['bg']};width:9px;margin:0;}}
        QScrollBar::handle:vertical {{background:{t['button_border']};border-radius:4px;min-height:25px;}}
        QScrollBar::add-line:vertical,QScrollBar::sub-line:vertical {{height:0;}}
        QScrollBar:horizontal {{background:{t['bg']};height:9px;margin:0;}}
        QScrollBar::handle:horizontal {{background:{t['button_border']};border-radius:4px;min-width:25px;}}
        QScrollBar::add-line:horizontal,QScrollBar::sub-line:horizontal {{width:0;}}
        '''
        self.setStyleSheet(qss)
        self.brand_icon.setPixmap(folder_mark(t['accent']))
        self.settings_button.setText('')
        self.settings_button.setIcon(themed_icon('settings',t['muted']))
        self.settings_button.setIconSize(QSize(24,24))
        self.btn_theme.setText('')
        self.btn_theme.setIcon(themed_icon('moon' if mode=='light' else 'sun',t['muted']))
        self.btn_theme.setIconSize(QSize(24,24))
        self.btn_theme.setToolTip('다크 테마로 전환' if mode=='light' else '라이트 테마로 전환')
        self._refresh_connection_strip()
        self.table.viewport().update()


    def _toggle_theme(self):
        next_mode = "dark" if getattr(self, "current_theme", "light") == "light" else "light"
        self.apply_theme(next_mode)
        try:
            data = {}
            if os.path.exists(CONF_FILE):
                with open(CONF_FILE, "r", encoding="utf-8") as f:
                    data = json.load(f)
            data["theme"] = next_mode
            with open(CONF_FILE, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def __init__(self,startup=False):
        super().__init__()
        self._startup_active = startup
        self._list_loaded = {'access':False,'project':False}
        # self._known_counts = {'access':0,'project':0}
        for folder in (DL_DIR, LOG_DIR, CONF_DIR, DEBUG_DIR):
            os.makedirs(folder, exist_ok=True)
        # self._notif_counts = {"신규미완료": 0, "진행-부여": 0, "진행-제거": 0, "종료-부여": 0, "종료-제거": 0}
        self.setWindowTitle(f"{APP_NAME}")
        self.resize(1060, 900)
        self._init_ui()
        # self._notify_refresh_mins = 10
        # try:
            # if os.path.exists(CONF_FILE):
                # with open(CONF_FILE, "r", encoding="utf-8") as f:
                    # data = json.load(f)
                    # self._notify_refresh_mins = int(data.get("notify_refresh_min", self._notify_refresh_mins))
        # except Exception:
            # pass
        # self.watch_thread = QThread(self)
        # self.watch_session = BusWatcher(DL_DIR)
        # self.watch_session.moveToThread(self.watch_thread)
        # self.watch_thread.start()
        # self.trigger_watcher_start.connect(self.watch_session.start, type=Qt.QueuedConnection)
        # self.trigger_watcher_stop.connect(self.watch_session.stop, type=Qt.QueuedConnection)
        # self.trigger_watcher_collect.connect(self.watch_session.collect_counts, type=Qt.QueuedConnection)
        # self.watch_session.countsReady.connect(self._on_counts_ready)
        # self.notify_timer = QTimer(self)
        # self.notify_timer.setInterval(max(5, int(self._notify_refresh_mins)) * 60 * 1000)
        # self.notify_timer.timeout.connect(self.refresh_notifications)
        # self.notify_timer.start()
        self._ensure_logo_and_set_title()
        self.session_thread = QThread(self)
        self.session = BusSessionManager(DL_DIR)
        self.session.moveToThread(self.session_thread)
        self.session.readyChanged.connect(self._on_session_ready)
        self.session.downloaded.connect(self._on_session_downloaded)
        self.session.processed.connect(self._on_session_processed)
        self.session.busyChanged.connect(self._on_session_busy)
        self.session_thread.start()
        # self.watch_session.readyChanged.connect(self._on_watcher_ready)
        self.trigger_session_start.connect(self.session.start, type=Qt.QueuedConnection)
        self.trigger_session_download.connect(self.session.download_list, type=Qt.QueuedConnection)
        self.trigger_new_download.connect(self.session.download_new_list,type=Qt.QueuedConnection)
        self.trigger_session_process.connect(self.session.process, type=Qt.QueuedConnection)
        self.trigger_session_cancel.connect(self.session.cancel_current, type=Qt.QueuedConnection)
        self.trigger_session_stop.connect(self.session.stop, type=Qt.QueuedConnection)
        # QTimer.singleShot(1000, self.refresh_notifications)
        self.api_config = load_api_config()
        self.api_thread = None
        self.api_worker = None
        self._closing = False
        self._shutdown_started = False
        if not startup:
            QTimer.singleShot(0,self._show_initial_api_settings)
        self.run_queue = []
        self.total_jobs = 0
        self.done_jobs = 0
        self.current_row = -1
        self.current_seq = 0
        self.current_mode = ""
        self.setAcceptDrops(True)
        self._load_creds()
        # if self.creds.get("id") and self.creds.get("pw"):
            # self.watch_session.set_creds(self.creds["id"], self.creds["pw"])
            # self.trigger_watcher_start.emit()
        self._log(f"{APP_NAME} 실행")
        self._buf_out = {}
        self.remove_fail_tolerance = int(self.creds.get("remove_fail_tol", 5))
        self.auto_complete_after_add = True
        self._waiting_for_bus = False
        self._pending_after_add_row = None
        self.bus_seq = 0
        self._ignore_bus_results = False
        self.stop_requested = False
        self._auto_create_group_decided = None
        self._retrying_after_group_create = False
        self._bus_mode = False
        self._bus_queue = []
        self._bus_done = 0
        self._bus_total = 0

        saved_theme = self._read_saved_theme()
        self.current_theme = saved_theme or "light"
        self.apply_theme(self.current_theme)

        self.debug_enabled = bool(self.creds.get("debug", False))
        self.session.set_debug(self.debug_enabled, DEBUG_DIR)

    def _init_ui(self):
        self._ui_refresh_pending = False
        self._api_verified = False
        self._executor_mode = None
        self.current_theme = 'light'
        self.setMinimumSize(1000,820)
        self.setFont(QFont('Segoe UI' if os.name=='nt' else 'Noto Sans CJK KR',10))
        root = QWidget()
        root.setObjectName('appRoot')
        main = QVBoxLayout(root)
        main.setContentsMargins(0,0,0,0)
        main.setSpacing(0)

        header = QWidget(); header.setObjectName('appHeader')
        self.header_bar = QHBoxLayout(header)
        self.header_bar.setContentsMargins(22,17,22,17)
        self.header_bar.setSpacing(15)
        self.brand_icon = QLabel(); self.brand_icon.setFixedSize(46,46)
        self.header_bar.addWidget(self.brand_icon)
        title = QVBoxLayout(); title.setSpacing(1)
        self.brand_title = QLabel('Folder Grant'); self.brand_title.setObjectName('brandTitle')
        self.brand_title.setCursor(Qt.PointingHandCursor)
        self.brand_title.mousePressEvent = lambda event:self._show_help()
        subtitle = QLabel('폴더 권한 관리'); subtitle.setObjectName('brandSubtitle')
        title.addWidget(self.brand_title); title.addWidget(subtitle)
        self.header_bar.addLayout(title)
        self.header_bar.addStretch()
        self.settings_button = QToolButton(); self.settings_button.setObjectName('utilityButton')
        self.settings_button.setText('⚙'); self.settings_button.setToolTip('설정 및 도움말')
        self.settings_button.setPopupMode(QToolButton.InstantPopup)
        menu = QMenu(self.settings_button)
        self.btn_api_settings = menu.addAction('API 연결 설정',self._open_api_settings)
        self.btn_settings = menu.addAction('BUS 계정 설정',self._open_settings)
        menu.addSeparator()
        self.columns_action = menu.addAction('상세 열 표시 (이름 · 부서 · Role)')
        self.columns_action.setCheckable(True)
        self.columns_action.toggled.connect(self._toggle_detail_columns)
        menu.addAction('도움말',self._show_help)
        # menu.addAction('BUS 요청 알림',self._open_notif_popup)
        self.settings_button.setMenu(menu)
        self.header_bar.addWidget(self.settings_button)
        self.btn_theme = QToolButton(); self.btn_theme.setObjectName('utilityButton')
        self.btn_theme.clicked.connect(self._toggle_theme)
        self.header_bar.addWidget(self.btn_theme)
        # self.btn_notif = QToolButton(self); self.btn_notif.hide()
        self._title_widget = self.brand_title
        main.addWidget(header)

        self.connection_bar = QWidget(); self.connection_bar.setObjectName('connectionBar')
        connection = QHBoxLayout(self.connection_bar)
        connection.setContentsMargins(24,10,24,10); connection.setSpacing(18)
        self.api_status_label = QLabel('○  API 미확인')
        self.bus_status_label = QLabel('○  BUS 미연결')
        self.server_label = QLabel('-'); self.server_label.setObjectName('mutedLabel')
        for widget in (self.api_status_label,self.bus_status_label,self.server_label):
            connection.addWidget(widget)
        connection.addStretch()
        main.addWidget(self.connection_bar)

        # self.tabs = WorkspaceTabBar(); self.tabs.setObjectName('workspaceTabs')
        self.tabs = QTabBar(); self.tabs.setObjectName('workspaceTabs')
        self.tabs.setExpanding(False)
        self.tabs.setElideMode(Qt.ElideNone)
        self.tabs.addTab('권한 요청'); self.tabs.addTab('신규 프로젝트')
        tab_row = QWidget(); tab_row.setObjectName('workspaceTabRow')
        tab_layout = QHBoxLayout(tab_row)
        tab_layout.setContentsMargins(18,0,18,0); tab_layout.setSpacing(0)
        tab_layout.addWidget(self.tabs)
        tab_layout.addStretch()
        # self.tabs.add_count_badges()
        # self.tab_count_badges = self.tabs.count_badges
        main.addWidget(tab_row)
        self.workspaces = QStackedWidget()
        self.request_page = QWidget()
        request_layout = QVBoxLayout(self.request_page)
        request_layout.setContentsMargins(18,0,18,0); request_layout.setSpacing(0)
        toolbar = QWidget(); toolbar.setObjectName('requestToolbar')
        toolbar_layout = QHBoxLayout(toolbar); toolbar_layout.setContentsMargins(15,12,12,12)
        heading = QLabel('요청 목록'); heading.setObjectName('sectionTitle')
        toolbar_layout.addWidget(heading)
        # self.new_request_notice = QLabel('새 요청이 있습니다 · BUS 요청 조회'); self.new_request_notice.setObjectName('mutedLabel'); self.new_request_notice.hide()
        # toolbar_layout.addWidget(self.new_request_notice); toolbar_layout.addStretch()
        toolbar_layout.addStretch()
        self.btn_request = QPushButton('BUS 요청 조회'); self.btn_request.setObjectName('primaryButton'); self.btn_request.clicked.connect(self._request_and_import)
        self.btn_file = QPushButton('엑셀 불러오기'); self.btn_file.clicked.connect(self.choose_file)
        self.btn_manual = QPushButton('수동 입력')
        self.btn_manual.clicked.connect(self.open_manual_dialog)
        for button in (self.btn_manual,self.btn_file,self.btn_request):
            toolbar_layout.addWidget(button)
        request_layout.addWidget(toolbar)

        content = QHBoxLayout(); content.setContentsMargins(0,0,0,0); content.setSpacing(10)
        self.table = CopyTable(); self.table.setObjectName('requestTable')
        self.table.setProperty('hasRows',False)
        self.table.setColumnCount(11)
        self.table.setHorizontalHeaderLabels(['','구분','요청','사번','이름','프로젝트','Level2 / Level3','Level3','부서','Role','상태'])
        header_check = CheckBoxHeader(Qt.Horizontal,self.table)
        self.table.setHorizontalHeader(header_check)
        header_check.stateChanged.connect(self.toggle_all_rows)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.table.setEditTriggers(QAbstractItemView.DoubleClicked|QAbstractItemView.EditKeyPressed)
        self.table.setAlternatingRowColors(False)
        self.table.setShowGrid(True)
        self.table.setWordWrap(False)
        self.table.verticalHeader().show()
        self.table.verticalHeader().setDefaultSectionSize(47)
        self.table.horizontalHeader().setFixedHeight(43)
        self.table.horizontalHeader().setMinimumSectionSize(30)
        self.table.horizontalHeader().setStretchLastSection(True)
        widths = {self.COL_SELECT:34,self.COL_KIND:52,self.COL_REQTYPE:78,self.COL_USER:67,
                  self.COL_PROJ:83,self.COL_LV2:159,self.COL_STATUS:138}
        for column,width in widths.items():
            self.table.horizontalHeader().setSectionResizeMode(column,QHeaderView.Interactive)
            self.table.setColumnWidth(column,width)
        for column in (self.COL_NAME,self.COL_LV3,self.COL_DEPT,self.COL_ROLE):
            self.table.setColumnHidden(column,True)
        self.table.setItemDelegate(RequestDelegate(self))
        self.table.cellChanged.connect(self._on_cell_changed)
        self.table.itemSelectionChanged.connect(self._refresh_detail)
        self.table.model().rowsInserted.connect(self._schedule_ui_refresh)
        self.table.model().rowsRemoved.connect(self._schedule_ui_refresh)
        self.table.setContextMenuPolicy(Qt.CustomContextMenu)
        self.table.customContextMenuRequested.connect(self._open_table_menu)
        content.addWidget(self.table,1)

        self.detail_panel = QWidget(); self.detail_panel.setObjectName('detailPanel')
        self.detail_panel.setFixedWidth(258)
        detail = QVBoxLayout(self.detail_panel); detail.setContentsMargins(17,14,17,14); detail.setSpacing(12)
        detail_title = QLabel('선택 요청 상세'); detail_title.setObjectName('sectionTitle'); detail.addWidget(detail_title)
        form = QFormLayout(); form.setHorizontalSpacing(12); form.setVerticalSpacing(9)
        self.detail_project = QLabel('-'); self.detail_employee = QLabel('-')
        self.detail_access = QLabel('요청을 선택하세요'); self.detail_access.setWordWrap(True)
        self.detail_role = QLabel('-'); self.detail_role.setWordWrap(True)
        self.detail_name = QLabel('-'); self.detail_department = QLabel('-')
        for text,widget in [('프로젝트',self.detail_project),('대상 사번',self.detail_employee),
                            ('권한',self.detail_access),('Role',self.detail_role)]:
            label = QLabel(text); label.setObjectName('mutedLabel'); form.addRow(label,widget)
        detail.addLayout(form)
        self.detail_extra = QLabel(''); self.detail_extra.setObjectName('mutedLabel'); self.detail_extra.setWordWrap(True)
        self.detail_extra.hide(); detail.addWidget(self.detail_extra)
        self.step_icons = []; self.step_labels = []
        for index,text in enumerate(['요청 검증','그룹 권한 적용','폴더 권한 적용','BUS 완료 처리']):
            step_layout = QHBoxLayout(); step_layout.setSpacing(12)
            icon = ProcessStepIcon(str(index+1)); icon.setObjectName('stepIcon'); icon.setAlignment(Qt.AlignCenter); icon.setFixedSize(24,24)
            label = QLabel(text); label.setObjectName('stepLabel'); label.setWordWrap(True)
            step_layout.addWidget(icon); step_layout.addWidget(label,1)
            detail.addLayout(step_layout)
            self.step_icons.append(icon); self.step_labels.append(label)
        detail.addStretch()
        self.progress_panel = QWidget(); self.progress_panel.setObjectName('detailBanner')
        banner = QVBoxLayout(self.progress_panel); banner.setContentsMargins(12,10,12,10); banner.setSpacing(3)
        self.detail_banner_title = QLabel('선택 요청 대기'); self.detail_banner_title.setObjectName('bannerTitle'); self.detail_banner_title.setWordWrap(True)
        self.status_label = StatusMessage('목록에서 요청을 선택하세요.')
        self.status_label.setObjectName('bannerBody')
        banner.addWidget(self.detail_banner_title); banner.addWidget(self.status_label)
        detail.addWidget(self.progress_panel)
        content.addWidget(self.detail_panel)
        request_layout.addLayout(content,1)
        self.workspaces.addWidget(self.request_page)
        self.project_page = QWidget(); self.project_page_layout = QVBoxLayout(self.project_page)
        self.project_page_layout.setContentsMargins(18,0,18,0)
        self.workspaces.addWidget(self.project_page)
        main.addWidget(self.workspaces,1)
        self.tabs.currentChanged.connect(self._switch_workspace)
        self.btn_newcheck = QAction('신규 프로젝트',self)
        self.btn_newcheck.triggered.connect(self.open_new_viewer)

        self.logs_panel = QWidget(); self.logs_panel.setObjectName('logsPanel')
        logs = QVBoxLayout(self.logs_panel); logs.setContentsMargins(12,0,12,0); logs.setSpacing(0)
        log_header = QHBoxLayout(); log_header.setContentsMargins(0,7,0,7)
        self.log_toggle = QToolButton(); self.log_toggle.setText('⌄  작업 로그')
        self.log_toggle.setObjectName('logToggle'); self.log_toggle.setCheckable(True); self.log_toggle.setChecked(True)
        self.log_toggle.toggled.connect(self._toggle_log)
        self.btn_clear_log = QPushButton('로그 지우기'); self.btn_clear_log.setObjectName('textButton')
        self.btn_clear_log.clicked.connect(self._clear_log)
        log_header.addWidget(self.log_toggle); log_header.addStretch(); log_header.addWidget(self.btn_clear_log)
        logs.addLayout(log_header)
        self.log = QTextEdit(); self.log.setReadOnly(True); self.log.setAcceptRichText(False)
        self.log.setObjectName('jobLog'); self.log.setFixedHeight(103)
        self.log.setFont(QFont('Consolas' if os.name=='nt' else 'Noto Sans Mono CJK KR',9))
        self.log.setLineWrapMode(QTextEdit.WidgetWidth); self.log.setWordWrapMode(QTextOption.WrapAnywhere)
        logs.addWidget(self.log)
        logs_outer = QWidget()
        logs_outer_layout = QVBoxLayout(logs_outer)
        logs_outer_layout.setContentsMargins(18,10,18,0)
        logs_outer_layout.addWidget(self.logs_panel)
        main.addWidget(logs_outer)

        self.actions_bar = QWidget(); self.actions_bar.setObjectName('actionsBar')
        actions = QHBoxLayout(self.actions_bar); actions.setContentsMargins(20,12,20,10); actions.setSpacing(9)
        self.selection_label = QLabel('선택 0건'); self.selection_label.setObjectName('selectionLabel')
        self.chk_auto_complete = QCheckBox('성공 후 BUS 자동 완료'); self.chk_auto_complete.setChecked(True)
        self.chk_auto_complete.toggled.connect(lambda checked:setattr(self,'auto_complete_after_add',checked))
        self.chk_dry = QCheckBox('Dry Run',self); self.chk_dry.hide()
        self.btn_preview = QPushButton('요청 검증'); self.btn_preview.clicked.connect(self._preview_selected)
        self.btn_run_execute = QPushButton('선택 실행'); self.btn_run_execute.setObjectName('primaryButton')
        self.btn_run_execute.clicked.connect(self._execute_selected)
        self.btn_run_complete = QPushButton('BUS 완료 재확인'); self.btn_run_complete.clicked.connect(self.run_complete)
        self.btn_stop = QPushButton('중지'); self.btn_stop.setObjectName('dangerButton'); self.btn_stop.setEnabled(False)
        self.btn_stop.clicked.connect(self._stop_all)
        actions.addWidget(self.selection_label); actions.addWidget(self.chk_auto_complete); actions.addStretch()
        for button in (self.btn_preview,self.btn_run_execute,self.btn_run_complete,self.btn_stop):
            actions.addWidget(button)
        main.addWidget(self.actions_bar)
        self.summary_label = QLabel('전체 0건 · 실행 중 0 · 완료 0 · 확인 필요 0')
        self.summary_label.setObjectName('summaryLabel')
        footer = QWidget()
        footer_layout = QHBoxLayout(footer)
        footer_layout.setContentsMargins(20,0,20,12)
        footer_layout.addWidget(self.summary_label,1)
        self.release_label = QLabel(f"V{APP_VERSION} Build:{APP_BUILD.replace('-','')}")
        self.release_label.setObjectName('mutedLabel')
        self.release_label.setAlignment(Qt.AlignRight|Qt.AlignVCenter)
        footer_layout.addWidget(self.release_label)
        main.addWidget(footer)
        self.setCentralWidget(root)
        self._statusbar = self.statusBar(); self._statusbar.hide()
        self.prg = QProgressBar(self); self.prg.hide()
        self._progress_elapsed = 0.0
        self._progress_started_at = None; self._progress_detail = ''; self._progress_project = ''
        self._progress_done = 0; self._progress_total = 0; self._progress_running = False
        self._progress_timer = QTimer(self); self._progress_timer.setInterval(1000)
        self._progress_timer.timeout.connect(self._refresh_progress_status)
        self._progress_timer.timeout.connect(self._refresh_detail)


    def _format_elapsed(self) -> str:
        if getattr(self, "_progress_started_at", None) is None:
            return "00:00"
        elapsed = max(0, int(time.monotonic() - self._progress_started_at)) if self._progress_running else int(self._progress_elapsed)
        minutes, seconds = divmod(elapsed, 60)
        hours, minutes = divmod(minutes, 60)
        if hours:
            return f"{hours:02d}:{minutes:02d}:{seconds:02d}"
        return f"{minutes:02d}:{seconds:02d}"

    def _set_progress_status(self, detail: str, project: str = "", done: int | None = None,
                             total: int | None = None, running: bool = True):
        if done is not None:
            self._progress_done = max(0, int(done))
        if total is not None:
            self._progress_total = max(0, int(total))
        self._progress_detail = detail or ""
        self._progress_project = project or ""
        was_running = self._progress_running
        self._progress_running = bool(running)

        if running and (not was_running or self._progress_started_at is None):
            self._progress_elapsed = 0.0
            self._progress_started_at = time.monotonic()
        if running and not self._progress_timer.isActive():
            self._progress_timer.start()
        self._refresh_progress_status()

    def _refresh_progress_status(self):
        total = max(0, int(getattr(self, "_progress_total", 0)))
        done = max(0, int(getattr(self, "_progress_done", 0)))
        running = bool(getattr(self, "_progress_running", False))
        detail = getattr(self, "_progress_detail", "") or "처리 중"
        project = getattr(self, "_progress_project", "") or ""

        if total > 0:
            if running:
                current = min(total, done + 1)
                prefix = f"처리 중 {current}/{total}"
                pct = int((current / total) * 100)
            else:
                current = min(total, done)
                prefix = f"완료 {current}/{total}" if done >= total else f"중지됨 {current}/{total}"
                pct = int((current / total) * 100)
        else:
            prefix = detail
            pct = 0 if running else 100

        parts = [prefix]
        if project:
            parts.append(f"프로젝트 {project}")
        if detail and detail != prefix:
            parts.append(detail)
        if running or getattr(self, "_progress_started_at", None):
            parts.append(self._format_elapsed())

        self.status_label.setText(" · ".join(parts))
        self.prg.setValue(max(0, min(100, pct)))
        self.progress_panel.setVisible(running or bool(getattr(self, "_progress_started_at", None)))
        self.prg.hide()

    def _finish_progress_status(self, detail: str = "완료", stopped: bool = False):
        if self._progress_running and self._progress_started_at is not None:
            self._progress_elapsed = max(0, time.monotonic() - self._progress_started_at)
        self._progress_running = False
        if self._progress_timer.isActive():
            self._progress_timer.stop()
        if not stopped and self._progress_total:
            self._progress_done = self._progress_total
        self._progress_detail = detail or ("중지됨" if stopped else "완료")
        self._refresh_progress_status()

    def _hide_progress_status(self):
        self._progress_running = False
        self._progress_started_at = None
        if self._progress_timer.isActive():
            self._progress_timer.stop()
        self.status_label.setText("")
        self.prg.setValue(0)
        self.progress_panel.setVisible(False)
        self.prg.hide()

    def _set_plain_status(self, message: str):
        self._progress_running = False
        self._progress_started_at = None
        if self._progress_timer.isActive():
            self._progress_timer.stop()
        self.status_label.setText(message or "")
        self.progress_panel.setVisible(bool(message))
        self.prg.hide()

    # def refresh_notifications(self):
        # if hasattr(self, "watch_session") and self.watch_session and self.watch_session.is_ready():
            # self.trigger_watcher_collect.emit()
            # return
        # self._on_counts_ready({"신규미완료":0,"진행-부여":0,"진행-제거":0,"종료-부여":0,"종료-제거":0})

    # @pyqtSlot(dict)
    # def _on_counts_ready(self, counts: dict):
        # counts = counts or {}
        # def _toi(x):
            # try:
                # return int(str(x).strip())
            # except Exception:
                # return 0

        # self._notif_counts = {
            # "신규미완료": _toi(counts.get("신규미완료", 0)),
            # "진행-부여": _toi(counts.get("진행-부여", 0)),
            # "진행-제거": _toi(counts.get("진행-제거", 0)),
            # "종료-부여": _toi(counts.get("종료-부여", 0)),
            # "종료-제거": _toi(counts.get("종료-제거", 0)),
        # }

        # self._refresh_request_count_badge()

    # def _refresh_request_count_badge(self):
        # counts = getattr(self,'_notif_counts',{})
        # totals = [sum(counts.get(key,0) for key in ('진행-부여','진행-제거','종료-부여','종료-제거')),
                  # counts.get('신규미완료',0)]
        # for badge,total in zip(self.tab_count_badges,totals):
            # badge.setText('99+' if total>99 else str(total))
            # badge.setVisible(total>0)
        # self._refresh_new_request_notice()

    def _mark_list_loaded(self,kind,count):
        self._list_loaded[kind] = True
        # self._known_counts[kind] = count
        # if kind=='project':
            # self._notif_counts['신규미완료'] = count
        # else:
            # for key in ('진행-부여','진행-제거','종료-부여','종료-제거'):
                # self._notif_counts[key] = 0
            # for row in range(self.table.rowCount()):
                # key = self._get(row,self.COL_KIND)+'-'+('제거' if self._get(row,self.COL_REQTYPE)==REQ_RELEASE else '부여')
                # if key in self._notif_counts:
                    # self._notif_counts[key] += 1
        # self._refresh_request_count_badge()
        table = self.table if kind=='access' else self._newdlg.tbl
        table.set_empty_message('대기 중인 요청이 없습니다.')
        # self._refresh_new_request_notice()

    # def _refresh_new_request_notice(self):
        # counts = getattr(self,'_notif_counts',{})
        # access = sum(counts.get(key,0) for key in ('진행-부여','진행-제거','종료-부여','종료-제거'))
        # known = getattr(self,'_known_counts',{})
        # self.new_request_notice.setVisible(access>known.get('access',0))
        # viewer = getattr(self,'_newdlg',None)
        # if viewer is not None:
            # viewer.new_request_notice.setVisible(counts.get('신규미완료',0)>known.get('project',0))

    # def _on_watcher_ready(self, ok: bool, msg: str):
        # self._refresh_connection_strip()
        # if hasattr(self, "btn_notif"):
            # self.btn_notif.setEnabled(bool(ok))
            # self.btn_notif.setToolTip("요청건 알림" if ok else "세션 준비 안됨")

        # if ok:
            # QTimer.singleShot(0, lambda: self.trigger_watcher_collect.emit())
        # else:
            # self._on_counts_ready({"신규미완료":0,"진행-부여":0,"진행-제거":0,"종료-부여":0,"종료-제거":0})

    # def _open_notif_popup(self):
        # if not hasattr(self, "_notif_counts") or not isinstance(self._notif_counts, dict):
            # self._notif_counts = {"신규미완료":0,"진행-부여":0,"진행-제거":0,"종료-부여":0,"종료-제거":0}
        # c = self._notif_counts
        # menu = QMenu(self)
        # n_new  = int(self._notif_counts.get("신규미완료", 0))
        # n_pg_g = int(self._notif_counts.get("진행-부여", 0))
        # n_pg_r = int(self._notif_counts.get("진행-제거", 0))
        # n_ed_g = int(self._notif_counts.get("종료-부여", 0))
        # n_ed_r = int(self._notif_counts.get("종료-제거", 0))
        # menu.addAction(f"폴더 생성 요청: {n_new}")
        # menu.addAction(f"진행-부여: {n_pg_g}")
        # menu.addAction(f"진행-제거: {n_pg_r}")
        # menu.addAction(f"종료-부여: {n_ed_g}")
        # menu.addAction(f"종료-제거: {n_ed_r}")
        # pos = self.settings_button.mapToGlobal(self.settings_button.rect().bottomRight())
        # menu.popup(pos)

    def open_new_viewer(self):
        if self.tabs.currentIndex()!=1:
            self.tabs.setCurrentIndex(1)
            return
        if self._closing:
            return
        existing = getattr(self,'_newdlg',None)
        if existing is not None:
            existing.show()
            if existing.isWindow():
                existing.raise_()
                existing.activateWindow()
            return
        self._ensure_project_viewer()
        def _after_ready(ok: bool):
            if not ok or self._closing:
                return
            self._load_new_items()

        self.ensure_bus_session_async("폴더 생성 요청 조회", _after_ready)

    def _ensure_project_viewer(self):
        if getattr(self,'_newdlg',None) is not None:
            return
        # Register the one persistent viewer before asynchronous session preparation.
        self._newdlg = NewItemsViewer(self)
        self._newdlg.setWindowFlags(Qt.Widget)
        main_layout = self.centralWidget().layout()
        main_layout.insertWidget(main_layout.indexOf(self.actions_bar)+1,self._newdlg.actions_bar)
        self._newdlg.actions_bar.setVisible(self.tabs.currentIndex()==1)
        self.project_page_layout.addWidget(self._newdlg)
        self._newdlg.btn_refresh.clicked.connect(self._load_new_items)
        self._newdlg.show()
        # self._refresh_new_request_notice()

    def _load_new_items(self):
        if self.api_thread is not None or self.session.is_busy():
            return
        try:
            self.session.newDownloaded.disconnect(self._on_new_downloaded)
        except Exception:
            pass
        self.session.newDownloaded.connect(self._on_new_downloaded)
        self._set_running_ui(True)
        self._newdlg.tbl.set_empty_message('요청을 조회하고 있습니다…')
        self.trigger_new_download.emit()


    def _on_new_downloaded(self, path: str, err: str):
        self._set_running_ui(False)
        if err:
            was_startup = self._startup_active or bool(getattr(self,'_startup_dialog',None) and self._startup_dialog.isVisible())
            self.startup_result.emit('project',False,'신규 프로젝트 조회에 실패했습니다.')
            if was_startup:
                return
            self._set_plain_status("폴더 생성 요청 로드 실패")
            QMessageBox.critical(self, "오류", f"폴더 생성 요청 로드 실패: {err}")
            return
        try:
            header, data = _parse_html_best_table(path)
        except Exception as exc:
            self._log(str(exc),summary='신규 프로젝트 조회 파일 확인 실패')
            self._on_new_downloaded('',str(exc))
            return
        if not header:
            was_startup = self._startup_active or bool(getattr(self,'_startup_dialog',None) and self._startup_dialog.isVisible())
            self.startup_result.emit('project',False,'신규 프로젝트 표 데이터를 확인하지 못했습니다.')
            if was_startup:
                return
            self._set_plain_status("표 데이터 없음")
            QMessageBox.warning(self, "안내", "표 데이터를 찾지 못했습니다.")
            return
        if getattr(self, "_newdlg", None):
            self._newdlg.set_data(header, data)
        self._set_plain_status(f"폴더 생성 요청 {len(data)}건")
        self._mark_list_loaded('project',len(data))
        self.startup_result.emit('project',True,'')


    def _refresh_has_rows(self):
        has_rows = self.table.rowCount() > 0

        self.table.setProperty("hasRows", has_rows)
        hh = self.table.horizontalHeader()
        vh = self.table.verticalHeader()
        for w in (self.table, hh, vh):
            try:
                w.setProperty("hasRows", has_rows)
                st = w.style()
                st.unpolish(w); st.polish(w)
                w.update()
            except Exception:
                pass

    def _refresh_theme_button_emoji(self):
        if hasattr(self, "btn_theme") and self.btn_theme:
            self.btn_theme.setText("🌙" if self.current_theme == "light" else "🌞")
            self.btn_theme.setToolTip("다크 모드로 전환" if self.current_theme == "light" else "라이트 모드로 전환")

    def _read_saved_theme(self) -> str:
        try:
            if os.path.exists(CONF_FILE):
                with open(CONF_FILE, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    t = (data.get("theme") or "").strip().lower()
                    if t in ("light", "dark"):
                        return t
        except Exception:
            pass
        return "light"

    def _ensure_logo_and_set_title(self):
        icon = resource_path('assets/fgt.ico')
        if os.path.exists(icon):
            self.setWindowIcon(QIcon(icon))


    def _apply_title_logo(self, logo_path: str | None):
        try:
            if hasattr(self, "_title_widget") and self._title_widget is not None:
                self._title_widget.setParent(None)
        except Exception:
            pass

        if logo_path and os.path.exists(logo_path):
            lbl = QLabel()
            pm = QPixmap(logo_path)
            if not pm.isNull():
                pm = pm.scaledToHeight(48, Qt.SmoothTransformation)
                lbl.setPixmap(pm)
            else:
                lbl.setText(APP_NAME)
                lbl.setFont(QFont("Segoe UI", 18, QFont.Bold))
            lbl.setCursor(Qt.PointingHandCursor)
            lbl.mousePressEvent = lambda e: self._show_help()
            widget = lbl
        else:
            txt = QLabel(APP_NAME)
            txt.setFont(QFont("Segoe UI", 18, QFont.Bold))
            txt.setAlignment(Qt.AlignCenter)
            txt.setCursor(Qt.PointingHandCursor)
            txt.mousePressEvent = lambda e: self._show_help()
            widget = txt

        self.header_bar.insertWidget(1, widget, 0, Qt.AlignVCenter)
        self._title_widget = widget

    def _ensure_bus_session(self, purpose: str = "완료 처리") -> bool:
        if not self.creds.get("id") or not self.creds.get("pw"):
            self._open_settings()
            if not self.creds.get("id") or not self.creds.get("pw"):
                QMessageBox.warning(self, "알림", "아이디/비밀번호가 필요합니다.")
                return False

        self.session.set_creds(self.creds["id"], self.creds["pw"])

        if self.session.is_ready():
            return True

        self._set_running_ui(True)
        self.status_label.setText(f"BUS 세션 준비 중 ({purpose})")
        self.trigger_session_start.emit()

        end = time.time() + 60
        while time.time() < end:
            QApplication.processEvents()
            if self.session.is_ready():
                self.status_label.setText("세션 준비됨")
                return True
            time.sleep(0.05)

        self._set_running_ui(False)
        QMessageBox.critical(self, "오류", "세션 초기화 실패")
        return False

    def ensure_bus_session_async(self, purpose: str, on_ready):
        if not self.creds.get("id") or not self.creds.get("pw"):
            self._open_settings()
            if not self.creds.get("id") or not self.creds.get("pw"):
                QMessageBox.warning(self, "알림", "아이디/비밀번호가 필요합니다.")
                if on_ready:
                    try:
                        on_ready(False)
                    except Exception:
                        pass
                return

        self.session.set_creds(self.creds["id"], self.creds["pw"])

        if self.session.is_ready():
            if on_ready:
                try:
                    on_ready(True)
                except Exception:
                    pass
            return

        if not hasattr(self, "_pending_bus_callbacks"):
            self._pending_bus_callbacks = []
        if on_ready:
            self._pending_bus_callbacks.append(on_ready)

        self._set_running_ui(True)
        self.status_label.setText(f"BUS 세션 준비 중 ({purpose})")
        self.trigger_session_start.emit()

    def _start_next_bus_item(self):
        if self.stop_requested or not self._bus_queue:
            self._bus_mode = False
            self._finish_progress_status("완료" if not self.stop_requested else "중지됨", stopped=self.stop_requested)
            self._set_running_ui(False)
            self.stop_requested = False
            return

        t = self._bus_queue.pop(0)
        row = t.get('row', -1)
        self._set_progress_status("BUS 완료 처리 중", t.get('proj', ''), self._bus_done, self._bus_total, running=True)
        if 0 <= row < self.table.rowCount():
            it = self.table.item(row, self.COL_STATUS)
            if it:
                it.setText("완료 처리중")

        self.trigger_session_process.emit([t])

    def run_complete(self):
        if self._closing or self._waiting_for_bus or self.api_thread is not None or self.session.is_busy():
            return
        self.stop_requested = False
        self._ignore_bus_results = False
        requests = []
        for row in range(self.table.rowCount()):
            if not self._is_row_checked(row):
                continue
            meta = self._row_metadata(row)
            if meta.get('source') == 'bus' and not meta.get('bus_done') and meta.get('request_id'):
                requests.append((row,meta['request_id']))
            else:
                self._log(f'{row+1}행: 미완료 BUS 요청번호가 없어 완료 처리에서 제외했습니다.')
        if not requests:
            return
        if not self.api_config.get('api_key'):
            self._open_api_settings()
            if not self.api_config.get('api_key'):
                return
        self._completion_targets = []
        self._progress_started_at = time.monotonic()
        self._progress_elapsed = 0.0
        thread = QThread(self)
        worker = ApiJobRefreshWorker(self._api_client(),requests)
        self.api_thread,self.api_worker = thread,worker
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        worker.refreshed.connect(self._on_completion_refreshed)
        worker.failed.connect(self._on_completion_refresh_error)
        worker.finished.connect(worker.deleteLater)
        worker.finished.connect(thread.quit,type=Qt.DirectConnection)
        thread.finished.connect(self._completion_refresh_finished)
        thread.finished.connect(thread.deleteLater)
        self._set_running_ui(True)
        self._set_progress_status('BUS 완료 전 서버 작업 상태 재조회', '',0,len(requests),running=True)
        thread.start()

    def _on_completion_refreshed(self,row,result):
        self._set_api_connection_state(True,result)
        meta = self._row_metadata(row)
        meta.pop('show_validation',None)
        meta['result'] = dict(result)
        meta.pop('_retry_parent_job_id',None); meta.pop('_failed_job_id',None)
        self._set_row_metadata(row,meta)
        self._log(f'{row+1}행 서버 재조회: '+json.dumps(result,ensure_ascii=False),
                  summary=self._api_log_summary(result,f'{row+1}행 서버 재조회'))
        if can_complete_bus(result,meta.get('source'),self.stop_requested or self._closing):
            self._completion_targets.append(self._bus_target(row))
            self.table.item(row,self.COL_STATUS).setText('서버 성공 확인 (BUS 완료 대기)')
        else:
            status = result.get('status','unknown')
            reason = result.get('error_code') or status
            self.table.item(row,self.COL_STATUS).setText('BUS 완료 보류: '+reason)
            if reason == 'POWERSHELL_TIMEOUT':
                self._log(f'{row+1}행: 서버가 PowerShell을 시간 초과로 종료했습니다. '
                          '서버 실행 제한을 변경해도 기존 실패 Job은 재개되지 않습니다. '
                          '부분 적용 상태 확인 및 별도 재실행이 필요합니다.')
            elif status in {'queued','running','cancel_requested'}:
                self._log(f'{row+1}행: 서버 작업이 아직 끝나지 않았습니다. 나중에 완료 처리를 다시 누르면 재조회합니다.')

    def _on_completion_refresh_error(self,row,message):
        self._set_api_connection_state(False)
        meta = self._row_metadata(row)
        meta['result'] = {}
        self._set_row_metadata(row,meta)
        self.table.item(row,self.COL_STATUS).setText('BUS 완료 보류: 서버 재조회 실패')
        self._log(f'{row+1}행: {message}. 이전 성공 결과로 BUS 완료를 진행하지 않습니다.',
                  summary=f'{row+1}행: 서버 재조회 실패 · BUS 완료 보류 (상세 내용은 로그 파일 확인)')

    def _completion_refresh_finished(self):
        self.api_worker = None
        self.api_thread = None
        if self._closing:
            QTimer.singleShot(0,self.close)
            return
        targets = self._completion_targets
        self._completion_targets = []
        if self.stop_requested or not targets:
            self._set_running_ui(False)
            self._finish_progress_status('BUS 완료 보류' if not self.stop_requested else '중지됨',stopped=self.stop_requested)
            return
        self._waiting_for_bus = True
        ready = self._ensure_bus_session('완료 처리')
        self._waiting_for_bus = False
        if not ready or self.stop_requested or self._closing:
            self._set_running_ui(False)
            return
        self._bus_mode = True
        self._bus_queue = targets
        self._bus_total = len(targets)
        self._bus_done = 0
        self._progress_started_at = time.monotonic()
        self._set_running_ui(True)
        self._start_next_bus_item()


    def _on_session_processed(self, results):
        if self._ignore_bus_results:
            return
        for result in results or []:
            row = result.get('row',-1)
            if 0 <= row < self.table.rowCount():
                meta = self._row_metadata(row)
                meta['bus_done'] = bool(result.get('ok'))
                self._set_row_metadata(row,meta)
                self.table.item(row,self.COL_STATUS).setText('BUS 완료' if result.get('ok') else 'BUS 실패 (완료만 재시도)')
                self._log(f"{row+1}행 BUS 결과: {result.get('msg','')}")
        if self._bus_mode:
            self._bus_done += len(results or [])
            self._start_next_bus_item()
        elif self._waiting_for_bus:
            if any(r.get('row')==self._pending_after_add_row for r in results or []):
                self._waiting_for_bus = False
                self._pending_after_add_row = None
                self._start_next_job()


    def _on_session_ready(self, ok: bool, msg: str):
        self._refresh_connection_strip()
        if ok:
            self._set_plain_status("세션 준비됨")
        else:
            self._set_plain_status(f"세션 미준비: {msg}")

        self._set_running_ui(False)
        cbs = getattr(self, "_pending_bus_callbacks", [])
        self._pending_bus_callbacks = []

        for cb in cbs:
            try:
                cb(bool(ok))
            except Exception as e:
                self._log(f"[세션 준비 콜백 오류] {e}")

    def _on_session_busy(self, b: bool):
        self._set_running_ui(b or self.api_thread is not None or self._waiting_for_bus)

    def _on_session_downloaded(self, path: str, err: str):
        self._set_running_ui(False)
        self.btn_request.setEnabled(True)
        self.btn_settings.setEnabled(True)

        if err:
            was_startup = self._startup_active or bool(getattr(self,'_startup_dialog',None) and self._startup_dialog.isVisible())
            self.startup_result.emit('access',False,'권한 요청 조회에 실패했습니다.')
            if was_startup:
                return
            if err.strip() == "사용자 취소":
                self._set_plain_status("요청 확인 취소됨")
                return
            self._set_plain_status("로드 실패")
            QMessageBox.critical(self, "오류", f"로드 실패: {err}")
            return

        self._set_plain_status("로드 완료")
        if not self.load_excel(path, append=False, silent=False, source="bus"):
            self.table.set_empty_message('요청 파일을 확인하지 못했습니다.')
            self.startup_result.emit('access',False,'권한 요청 파일을 확인하지 못했습니다.')
            return

        try:
            end_combined = os.path.join(DL_DIR, "종료권한리스트_합본.xls")
            normal_combined = os.path.join(DL_DIR, "권한리스트_합본.xls")
            if os.path.basename(path) == "권한리스트_합본.xls" and os.path.exists(end_combined):
                if not self.load_excel(end_combined, append=True, silent=True, source="bus"):
                    raise ValueError("종료 요청 파일 확인 실패")
            elif os.path.basename(path) == "종료권한리스트_합본.xls" and os.path.exists(normal_combined):
                if not self.load_excel(normal_combined, append=True, silent=True, source="bus"):
                    raise ValueError("진행 요청 파일 확인 실패")
        except Exception as exc:
            self._log(str(exc),summary='권한 요청 파일 확인 실패')
            self.startup_result.emit('access',False,'권한 요청 파일을 확인하지 못했습니다.')
            return

        prog_grant = prog_rel = end_grant = end_rel = 0
        for r in range(self.table.rowCount()):
            kind = (self._get(r, self.COL_KIND) or "").strip()
            req  = (self._get(r, self.COL_REQTYPE) or REQ_GRANT).strip()
            if   kind == "진행" and req == REQ_GRANT:   prog_grant += 1
            elif kind == "진행" and req == REQ_RELEASE: prog_rel   += 1
            elif kind == "종료" and req == REQ_GRANT:   end_grant  += 1
            elif kind == "종료" and req == REQ_RELEASE: end_rel    += 1

        parts = []
        if prog_grant: parts.append(f"진행-부여 {prog_grant}건")
        if prog_rel:   parts.append(f"진행-해제 {prog_rel}건")
        if end_grant:  parts.append(f"종료-부여 {end_grant}건")
        if end_rel:    parts.append(f"종료-해제 {end_rel}건")
        self._log("로드 완료: " + (" , ".join(parts) if parts else "항목 없음"))
        self._mark_list_loaded('access',self.table.rowCount())
        self.startup_result.emit('access',True,'')

    def _load_creds(self):
        self.creds = {"id":"", "pw":""}
        try:
            if os.path.exists(CONF_FILE):
                with open(CONF_FILE, "r", encoding="utf-8") as f:
                    self.creds = json.load(f)
                    self.creds.setdefault("debug", False)
        except:
            pass
        self.remove_fail_tolerance = int(self.creds.get("remove_fail_tol", 5))

        self._remember_credentials = bool(self.creds.get('id') and self.creds.get('pw'))

    def _save_creds(self, id_, pw_, debug_):
        with open(CONF_FILE, "w", encoding="utf-8") as f:
            json.dump({"id": id_, "pw": pw_, "debug": bool(debug_), "theme": getattr(self, "current_theme", "light")},
                      f, ensure_ascii=False, indent=2)

    def closeEvent(self,event):
        self._closing = True
        startup = getattr(self,'_startup_dialog',None)
        if startup is not None:
            startup.active = False; startup.deadline.stop()
            if startup.api_thread is not None:
                event.ignore(); QTimer.singleShot(100,self.close); return
        if self.api_thread is not None and self.api_thread.isRunning():
            self.api_worker.stop()
            event.ignore()
            return
        newdlg = getattr(self,'_newdlg',None)
        if newdlg and getattr(newdlg,'worker_thread',None) and newdlg.worker_thread.isRunning():
            newdlg.worker.stop()
            event.ignore()
            QTimer.singleShot(200,self.close)
            return
        # self.notify_timer.stop()
        self._progress_timer.stop()
        self.session.cancel_current()
        # if self.session.is_busy() or self.watch_session._mgr.is_busy():
        if self.session.is_busy():
            event.ignore()
            QTimer.singleShot(200,self.close)
            return
        if not self._shutdown_started:
            self._shutdown_started = True
            self.session.stopped.connect(self.session_thread.quit, type=Qt.DirectConnection)
            # self.watch_session.stopped.connect(self.watch_thread.quit, type=Qt.DirectConnection)
            self.trigger_session_stop.emit()
            # self.trigger_watcher_stop.emit()
        # if self.session_thread.isRunning() or self.watch_thread.isRunning():
        if self.session_thread.isRunning():
            event.ignore()
            QTimer.singleShot(100,self.close)
            return
        event.accept()


    def _open_settings(self):
        current_debug = bool(self.creds.get("debug", False))
        remembered = getattr(self,"_remember_credentials",False)
        current_tol = int(self.creds.get("remove_fail_tol", 5))

        dlg = SettingsDialog(self, self.creds, remembered=remembered, debug_on=current_debug, fail_tol_default=current_tol)

        if dlg.exec_() == QDialog.Accepted:
            uid, pw, do_save, debug_on, fail_tol, notify_mins = dlg.result()
            tol = dlg.get_fail_tol()

            self.creds = {
                "id": uid,
                "pw": pw,
                "debug": bool(debug_on),
                "remove_fail_tol": int(tol),
                "notify_refresh_min": int(notify_mins),
            }

            # if uid and pw:
                # self.watch_session.set_creds(uid, pw)
                # self.trigger_watcher_start.emit()
            # else:
                # self.trigger_watcher_stop.emit()

            data = {
                "theme": getattr(self, "current_theme", "light"),
                "debug": bool(debug_on),
                "remove_fail_tol": int(tol),
                "notify_refresh_min": int(notify_mins),
            }
            self._apply_login_credentials(uid,pw,do_save,data)

            # self._notify_refresh_mins = int(notify_mins)
            # self.notify_timer.setInterval(max(5, self._notify_refresh_mins) * 60 * 1000)
            self.debug_enabled = bool(debug_on)
            self.session.set_debug(self.debug_enabled, DEBUG_DIR)
            self.remove_fail_tolerance = int(tol)

    def _apply_login_credentials(self,uid,pw,remember,settings_data=None):
        """Shared credential handling for settings and the compact startup prompt."""
        self.creds.update(id=uid,pw=pw)
        try:
            data = settings_data
            if data is None:
                data = {}
                if os.path.exists(CONF_FILE):
                    with open(CONF_FILE,encoding='utf-8') as handle:
                        data = json.load(handle)
            if remember:
                data.update(id=uid,pw=pw)
            else:
                data.pop('id',None); data.pop('pw',None)
            with open(CONF_FILE,'w',encoding='utf-8') as handle:
                json.dump(data,handle,ensure_ascii=False,indent=2)
            self._remember_credentials = bool(remember)
        except Exception:
            self._log('BUS 계정 저장 설정을 기록하지 못했습니다.')

    def _request_and_import(self):
        if self.api_thread is not None or self.session.is_busy():
            return
        def ready(ok):
            if not ok or self._closing:
                return
            self.btn_request.setEnabled(False)
            self.btn_settings.setEnabled(False)
            self.table.set_empty_message('요청을 조회하고 있습니다…')
            self._set_running_ui(True)
            self._set_progress_status('요청 확인 중','',0,0,running=True)
            self.trigger_session_download.emit()
        self.ensure_bus_session_async('권한 요청 조회',ready)

    def _clear_log(self):
        self.log.clear()

    def _stop_all(self):
        self.run_queue = []
        self.stop_requested = True
        self._ignore_bus_results = True
        self._waiting_for_bus = False
        self._bus_queue = []
        self._bus_mode = False
        self.session.cancel_current()
        if self.api_worker is not None:
            self.api_worker.stop()
            self._log('API 취소 요청 후 최종 상태를 기다립니다.')
        newdlg = getattr(self,'_newdlg',None)
        if newdlg and getattr(newdlg,'worker',None):
            newdlg.worker.stop()
        if self.api_thread is None and not self.session.is_busy():
            self._set_running_ui(False)
        if self.api_thread is not None or (newdlg and newdlg.worker_thread is not None) or self.session.is_busy():
            self._set_progress_status('중지 요청됨',running=True)
        else:
            self._finish_progress_status('중지 요청됨',stopped=True)


    def _set_running_ui(self, running: bool):
        running = running or getattr(self,'api_thread',None) is not None or getattr(self,'_waiting_for_bus',False)
        newdlg = getattr(self,'_newdlg',None)
        running = running or bool(newdlg and getattr(newdlg,'worker_thread',None))
        if running:
            self._set_progress_status(
                getattr(self, "_progress_detail", "") or "처리 중",
                getattr(self, "_progress_project", ""),
                getattr(self, "done_jobs", 0),
                getattr(self, "total_jobs", 0),
                running=True,
            )

        elif self._progress_running:
            self._finish_progress_status('중지됨' if self.stop_requested else '완료',stopped=self.stop_requested)

        to_disable = [
            getattr(self, "btn_preview", None),
            getattr(self, "tabs", None),
            getattr(self, "btn_api_settings", None),
            getattr(self, "btn_manual", None),
            getattr(self, "btn_file", None),
            getattr(self, "btn_run_execute", None),
            getattr(self, "btn_run_complete", None),
            getattr(self, "btn_clear_log", None),
            getattr(self, "btn_request", None),
            getattr(self, "btn_settings", None),
            getattr(self, "chk_dry", None),
            getattr(self, "chk_auto_complete", None),
            getattr(self, "btn_newcheck", None),
        ]
        if newdlg is not None:
            to_disable.extend([newdlg.btn_manual,newdlg.btn_refresh,newdlg.btn_create,
                               newdlg.btn_preview,newdlg.btn_complete,newdlg.chk_auto_complete])
        if newdlg is not None:
            newdlg.tbl.setContextMenuPolicy(Qt.PreventContextMenu if running else Qt.CustomContextMenu)
        for w in to_disable:
            if w:
                w.setEnabled(not running)

        if hasattr(self, "btn_stop") and self.btn_stop:
            self.btn_stop.setEnabled(running)

        self.setAcceptDrops(not running)
        if hasattr(self,'table'):
            self.table.setEditTriggers(QAbstractItemView.NoEditTriggers if running else (QAbstractItemView.DoubleClicked | QAbstractItemView.SelectedClicked | QAbstractItemView.EditKeyPressed))
        if hasattr(self, "table") and self.table:
            self.table.setContextMenuPolicy(Qt.PreventContextMenu if running else Qt.CustomContextMenu)



    def _start_next_job(self):
        if self.api_thread is not None or self._waiting_for_bus:
            return
        if self._closing or self.stop_requested or not self.run_queue:
            self._set_running_ui(False)
            self._finish_progress_status('중지됨' if self.stop_requested else '완료', stopped=self.stop_requested)
            return
        entry = self.run_queue.pop(0)
        row,payload,dry = entry[:3]
        retry_job_id = entry[3] if len(entry)>3 else None
        self.current_row = row
        meta = self._row_metadata(row)
        meta.pop('show_validation',None)
        meta['live_result'] = dict(status='validating' if dry else 'submitting')
        self._set_row_metadata(row,meta)
        self.current_seq = self.done_jobs+1
        self.current_mode = 'add' if payload['operation']=='grant' else 'remove'
        self.table.item(row,self.COL_STATUS).setText('요청 검증 중' if dry else 'API 실행중')
        self.table.setCurrentCell(row,self.COL_PROJ)
        self._refresh_detail()
        self._log(f"API {'요청 검증' if dry else 'Job'} 요청: {payload['request_id']}",
                  summary=f"{row+1}행 · 프로젝트 {self._get(row,self.COL_PROJ)}: {'요청 검증' if dry else '권한 작업'} 요청")
        if retry_job_id:
            self._log(f'실패 작업 재실행: 이전 job_id={retry_job_id}',summary=f'{row+1}행: 실패 단계 재실행 요청')
        thread = QThread(self)
        worker = ApiJobWorker(self._api_client(),'access',
                              dict(payload,requested_by=self._requester()) if retry_job_id else payload,
                              dry,retry_job_id=retry_job_id,validate_only=dry)
        self.api_thread,self.api_worker = thread,worker
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        worker.progress.connect(self._on_api_progress)
        worker.succeeded.connect(self._on_api_result)
        worker.failed.connect(self._on_api_error)
        worker.finished.connect(worker.deleteLater)
        worker.finished.connect(thread.quit, type=Qt.DirectConnection)
        thread.finished.connect(self._api_thread_finished)
        thread.finished.connect(thread.deleteLater)
        thread.start()


    def open_manual_dialog(self):
        try:
            dlg = ManualEntryDialog(self)
            if dlg.exec_() == QDialog.Accepted and dlg.result_row:
                reqtype, user, proj, lv2, lv3, dept, role, kind = dlg.result_row
                self.add_table_row(reqtype, user, proj, lv2, lv3, dept, role, kind)
        except Exception as e:
            self._log(f"[수동입력 오픈 예외] {e}")
            QMessageBox.critical(self, "오류", f"수동 입력 다이얼로그 실행 오류:\n{e}")

    def add_table_row(self, reqtype: str, user: str, proj: str, lv2: str, lv3: str, dept: str, role: str, kind: str = "진행"):
        self.table.blockSignals(True)
        try:
            r = self.table.rowCount()
            self.table.insertRow(r)

            chk = QCheckBox()
            chk.setChecked(True)

            wrapper = QWidget()
            layout = QHBoxLayout(wrapper)
            layout.addWidget(chk)
            layout.setAlignment(Qt.AlignCenter)
            layout.setContentsMargins(0,0,0,0)
            self.table.setCellWidget(r, self.COL_SELECT, wrapper)

            vals = [kind, reqtype or REQ_GRANT, user, "", proj, lv2, lv3, dept or "", role or "", "대기"]

            for j, val in enumerate(vals, start=1):
                col = j
                it = QTableWidgetItem(str(val))
                it.setTextAlignment(Qt.AlignCenter)

                if col in self.EDITABLE_COLS:
                    it.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled | Qt.ItemIsEditable)
                else:
                    it.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)
                self.table.setItem(r, col, it)
        finally:
            self.table.blockSignals(False)

        self._set_row_metadata(r,dict(source='test-client',request_id=new_client_request_id(),bus_done=False))
        if (kind or "").strip() == "종료":
            self._log(f"수동 입력 추가: 구분={kind}, 요청={reqtype}, 사번={user}, 프로젝트={proj}")
        else:
            self._log(f"수동 입력 추가: 구분={kind}, 요청={reqtype}, 사번={user}, 프로젝트={proj}, L2={lv2}, L3={lv3}, ROLE={'없음' if not role else role}")

    def _log_validation(self, result, context):
        self.log_toggle.setChecked(True)
        self._log(json.dumps(result,ensure_ascii=False,indent=2),
                  summary=f'{context}: {validation_label(result)}')
        symbols = {'passed':'✓','failed':'✕','unknown':'?','warning':'!','not_applicable':'-'}
        for check in result.get('checks',[]):
            target = f" · {check['target']}" if check.get('target') else ''
            self._log(f"{symbols.get(check.get('status'),'?')} {check.get('label','검사')}{target}: {check.get('message','')}")
        for warning in result.get('warnings',[]):
            self._log('안내: '+warning)

    @staticmethod
    def _api_log_summary(result,context):
        status = result.get('status','unknown')
        label = {'preview':'요청 검증 완료','simulated':'모의 실행 완료','succeeded':'성공',
                 'partially_succeeded':'부분 성공','failed':'실패','cancelled':'중지됨',
                 'running':'실행 중','queued':'실행 대기'}.get(status,'결과 확인 필요')
        if result.get('bus_done'):
            label += ' · BUS 완료'
        error = result.get('error_code')
        if error:
            label += ' · '+('서버 시간 초과' if error=='POWERSHELL_TIMEOUT' else str(error))
        if result.get('replayed'):
            label += ' · 이전 작업 결과 · 재실행 안 함'
        return f'{context}: {label}'

    def _log(self, message: str, seq: int = None, dry: bool = False, *, summary: str = None):
        api_key = getattr(self,'api_config',{}).get('api_key','')
        def sanitize(value):
            value = str(value or '')
            return value.replace(api_key,'[API Key]') if api_key else value
        detailed = sanitize(message)
        visible = sanitize(summary if summary is not None else message)
        ts = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        dry_str = ' Dry' if dry else ''
        tag = f'{dry_str} #{seq}' if seq is not None else ''
        norm = detailed.replace('\r','\n').replace('\t','    ').rstrip()
        block = f'[{ts}{tag}]\n{norm}\n'
        # The file keeps the full record; the GUI receives its own short message.
        with open(os.path.join(LOG_DIR,'access.log'),'a',encoding='utf-8') as handle:
            handle.write(block)
        visible = ' '.join(visible.split())
        cursor = self.log.textCursor(); cursor.movePosition(QTextCursor.End)
        self.log.setTextCursor(cursor)
        if self.log.toPlainText():
            self.log.insertPlainText('\n')
        self.log.insertPlainText(f'[{ts}{tag}] {visible}')
        self.log.ensureCursorVisible()

    def choose_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "승인 엑셀 선택", "", "Excel Files (*.xlsx *.xls)")
        if not path:
            return
        self.load_excel(path)

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            for url in event.mimeData().urls():
                lf = url.toLocalFile().lower()
                if lf.endswith(".xlsx") or lf.endswith(".xls"):
                    event.acceptProposedAction(); return
        event.ignore()

    def dropEvent(self, event):
        urls = event.mimeData().urls()
        if not urls: return
        path = urls[0].toLocalFile()
        if path.lower().endswith((".xlsx",".xls")):
            self.load_excel(path)

    def load_excel(self, file_path: str, append: bool = False, silent: bool = False, source: str = 'test-client'):
        try:
            self.table.blockSignals(True)
            #if not (append and silent):
                #self.file_label.setText(file_path)
            ext = os.path.splitext(file_path)[1].lower()

            start_offset = self.table.rowCount() if append else 0

            rows = []
            identities = []

            if ext == ".xlsx":
                wb = load_workbook(file_path, data_only=True)
                ws = wb.active
                header_raw = [str(c.value) if c.value is not None else "" for c in ws[1]]
                colmap = auto_map_columns(header_raw)

                is_end = any(("열람" in (str(h) or "")) for h in header_raw)
                kind = "종료" if is_end else "진행"

                required = {"user", "proj"} if is_end else {"user", "proj", "level2", "level3"}
                missing = [k for k in required if k not in colmap]
                if missing:
                    self._log(f"엑셀 로드 실패: 필요한 컬럼 없음 -> {missing} / 헤더: {header_raw}")
                    return False

                def gv(row, key, default=""):
                    i = colmap.get(key)
                    if i is None:
                        return default
                    v = row[i] if i < len(row) else None
                    return (str(v).strip() if v is not None else default)

                for r in ws.iter_rows(min_row=2, values_only=True):
                    if not any(r):
                        continue
                    user_id = gv(r, "user")
                    name    = gv(r, "name")
                    proj    = gv(r, "proj")
                    lv2     = "" if is_end else gv(r, "level2")
                    lv3     = "" if is_end else gv(r, "level3")
                    role    = gv(r, "role") if not is_end else ""
                    dept    = gv(r, "dept")

                    if not user_id or not proj:
                        continue
                    if not is_end and (not lv2 or not lv3):
                        continue

                    reqtype = REQ_RELEASE if _is_release_row_by_values(list(r), header_raw) else REQ_GRANT
                    path = build_path_l3(proj, lv2, lv3)
                    rows.append((kind, reqtype, user_id, name, proj, lv2, lv3, dept, role, "대기"))
                    identities.append(bus_row_identity(header_raw,r))

                wb.close()

            elif ext == ".xls":
                header_raw, data_rows = _parse_html_best_table(file_path)
                if not header_raw:
                    self._log("엑셀 로드 실패: .xls(HTML) 테이블을 찾지 못했습니다.")
                    return False

                colmap = auto_map_columns(header_raw)

                is_end = any(("열람" in (str(h) or "")) for h in header_raw)
                kind = "종료" if is_end else "진행"

                required = {"user", "proj"} if is_end else {"user", "proj", "level2", "level3"}
                missing = [k for k in required if k not in colmap]
                if missing:
                    self._log(f"엑셀 로드 실패: 필요한 컬럼 없음 -> {missing} / 헤더: {header_raw}")
                    return False

                def gv(row, key, default=""):
                    i = colmap.get(key)
                    if i is None:
                        return default
                    v = row[i] if i < len(row) else None
                    return (str(v).strip() if v is not None else default)

                for r in data_rows:
                    user_id = gv(r, "user")
                    name    = gv(r, "name")
                    proj    = gv(r, "proj")
                    lv2     = "" if is_end else gv(r, "level2")
                    lv3     = "" if is_end else gv(r, "level3")
                    role    = gv(r, "role") if not is_end else ""
                    dept    = gv(r, "dept")

                    if not user_id or not proj:
                        continue
                    if not is_end and (not lv2 or not lv3):
                        continue

                    reqtype = REQ_RELEASE if _is_release_row_by_values(r, header_raw) else REQ_GRANT
                    path = build_path_l3(proj, lv2, lv3)
                    rows.append((kind, reqtype, user_id, name, proj, lv2, lv3, dept, role, "대기"))
                    identities.append(bus_row_identity(header_raw,r))

            else:
                self._log(f"엑셀 로드 실패: 지원하지 않는 확장자 ({ext})")
                return False

            if append:
                self.table.setRowCount(start_offset + len(rows))
                base = start_offset
            else:
                self.table.setRowCount(len(rows))
                base = 0

            for i, row in enumerate(rows):
                r = base + i

                chk = QCheckBox()
                chk.setChecked(True)

                wrapper = QWidget()
                layout = QHBoxLayout(wrapper)
                layout.addWidget(chk)
                layout.setAlignment(Qt.AlignCenter)
                layout.setContentsMargins(0, 0, 0, 0)
                self.table.setCellWidget(r, self.COL_SELECT, wrapper)

                for j, val in enumerate(rows[i]):
                    col = j + 1
                    it = QTableWidgetItem(str(val))
                    it.setTextAlignment(Qt.AlignCenter)
                    if col in self.EDITABLE_COLS:
                        it.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled | Qt.ItemIsEditable)
                    else:
                        it.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)
                    self.table.setItem(r, col, it)

            for row in range(base,base+len(rows)):
                values = self._row_values(row)
                try:
                    request_id = make_bus_request_id(values,identities[row-base]) if source=='bus' else new_client_request_id()
                except ValueError:
                    request_id = new_client_request_id()
                self._set_row_metadata(row,dict(source=source,request_id=request_id,bus_done=False))
            hv = self.table.horizontalHeader()
            hv.setStretchLastSection(True)

            return True
        except Exception as e:
            self._log(f"엑셀 로드 실패: {e}")
            return False

        finally:
            self.table.blockSignals(False)

    def _on_cell_changed(self,row,col):
        self._schedule_ui_refresh()
        if col == self.COL_STATUS:
            return
        meta = self._row_metadata(row)
        if meta:
            self._set_row_metadata(row,dict(source='test-client',request_id=new_client_request_id()))
            self.table.item(row,self.COL_STATUS).setText('수정됨 (수동 요청)')


    def _open_table_menu(self, pos):
        if self.table.contextMenuPolicy() != Qt.CustomContextMenu:
            return
        row = self.table.indexAt(pos).row()
        m = QMenu(self)
        act_del_row = m.addAction("행 삭제")
        act_del_sel = m.addAction("선택 행 삭제")
        act_del_all = m.addAction("전체 삭제")

        gpos = self.table.viewport().mapToGlobal(pos)
        act = m.exec_(gpos)

        if act == act_del_row:
            self._delete_row(row)
        elif act == act_del_sel:
            self._delete_checked_rows()
        elif act == act_del_all:
            self._delete_all_rows()

    def _delete_all_rows(self):
        if self.table.contextMenuPolicy() != Qt.CustomContextMenu:
            return
        self.table.blockSignals(True)
        try:
            self.table.setRowCount(0)
        finally:
            self.table.blockSignals(False)

    def _delete_row(self, row: int):
        if self.table.contextMenuPolicy() != Qt.CustomContextMenu:
            return
        if row < 0 or row >= self.table.rowCount():
            return
        self.table.removeRow(row)

    def _get_checkbox(self, row: int) -> QCheckBox:
        w = self.table.cellWidget(row, self.COL_SELECT)
        if not w:
            return None
        return w.findChild(QCheckBox)

    def _is_row_checked(self, row: int) -> bool:
        cb = self._get_checkbox(row)
        return bool(cb and cb.isChecked())

    def _delete_checked_rows(self):
        if self.table.contextMenuPolicy() != Qt.CustomContextMenu:
            return
        for r in range(self.table.rowCount()-1, -1, -1):
            if self._is_row_checked(r):
                self.table.removeRow(r)

    def validate_row(self, row, mode='add'):
        try:
            meta = self._row_metadata(row)
            make_access_payload(self._row_values(row), meta.get('request_id') or new_client_request_id(),
                                meta.get('source','test-client'), self._requester())
            return True, ''
        except ValueError as exc:
            return False, f'{row+1}행: {exc}'


    def toggle_all_rows(self, checked: bool):
        for r in range(self.table.rowCount()):
            cb = self._get_checkbox(r)
            if cb:
                cb.setChecked(checked)

    def run_execute(self):
        if self.api_thread is not None or self.session.is_busy() or self._waiting_for_bus:
            return
        if not self.api_config.get('api_key'):
            self._open_api_settings()
            if not self.api_config.get('api_key'):
                return
        retry_rows = {}
        if not self.chk_dry.isChecked():
            for row in range(self.table.rowCount()):
                meta = self._row_metadata(row)
                result = meta.get('result') or {}
                parent = None
                if result.get('status') in {'failed','partially_succeeded'} and result.get('executor_mode')=='powershell':
                    parent = result.get('job_id')
                elif result.get('status') in {None,'preview'}:
                    parent = meta.get('_retry_parent_job_id') or meta.get('_failed_job_id')
                if self._is_row_checked(row) and parent and not meta.get('bus_done'):
                    retry_rows[row] = parent
            if retry_rows and QMessageBox.question(self,'실패 작업 재실행',
                    f'실패·부분 성공한 {len(retry_rows)}건을 새 Job으로 재실행할까요?\n'
                    '폴더 등 실패 원인을 조치했는지 확인하세요.\n'
                    '성공한 단계는 유지하고 실패한 단계와 이후 단계만 실행합니다.\n'
                    '모든 단계가 실제 성공한 경우에만 BUS 완료가 가능합니다.',
                    QMessageBox.Yes|QMessageBox.No,QMessageBox.No)!=QMessageBox.Yes:
                return
        self.stop_requested = False
        self._ignore_bus_results = False
        self._waiting_for_bus = False
        self.auto_complete_after_add = self.chk_auto_complete.isChecked()
        queue = []
        for row in range(self.table.rowCount()):
            if not self._is_row_checked(row):
                continue
            ok, reason = self.validate_row(row)
            if not ok:
                self.table.item(row,self.COL_STATUS).setText('검증실패')
                self._log(reason)
                continue
            meta = self._row_metadata(row)
            if not meta:
                meta = dict(source='test-client',request_id=new_client_request_id())
            values = self._row_values(row)
            # Keep the original operator when retrying the same idempotent payload.
            payload = meta.get('payload') or make_access_payload(values,meta['request_id'],meta['source'],self._requester())
            meta['payload'] = payload
            meta['live_result'] = {}
            if row in retry_rows:
                meta['_retry_parent_job_id'] = retry_rows[row]
            if not self.chk_dry.isChecked():
                meta['result'] = {}
            self._set_row_metadata(row,meta)
            queue.append((row,payload,self.chk_dry.isChecked(),retry_rows.get(row)))
        if not queue:
            return
        self.run_queue = queue
        self.total_jobs = len(queue)
        self.done_jobs = 0
        self._progress_started_at = time.monotonic()
        self._set_running_ui(True)
        self._start_next_job()


    def _get(self, row: int, col: int) -> str:
        it = self.table.item(row, col)
        return it.text().strip() if it else ""

    def _show_initial_api_settings(self):
        if not self.api_config.get('api_key'):
            self._open_api_settings()


    def _open_api_settings(self):
        if self.api_thread is not None or self._waiting_for_bus:
            return
        dialog = ApiSettingsDialog(self.api_config, self)
        if dialog.exec_() == QDialog.Accepted:
            self.api_config = dialog.values()
            self._executor_mode = None
            self._set_api_connection_state(getattr(dialog,'connection_verified',False))
            self._log('API 연결 설정을 저장했습니다. API Key는 로그에 기록하지 않습니다.')


    def _api_client(self):
        return FolderGrantApiClient(self.api_config.get('base_url', DEFAULT_API_URL),
                                   self.api_config.get('api_key', ''))


    def _requester(self):
        import getpass
        return self.creds.get('id') or getpass.getuser() or 'client-operator'


    def _row_values(self, row):
        return dict(kind=self._get(row, self.COL_KIND), req=self._get(row, self.COL_REQTYPE),
                    user=self._get(row, self.COL_USER), proj=self._get(row, self.COL_PROJ),
                    lv2=self._get(row, self.COL_LV2), lv3=self._get(row, self.COL_LV3),
                    role=self._get(row, self.COL_ROLE))


    def _row_metadata(self, row):
        item = self.table.item(row, self.COL_STATUS)
        return dict(item.data(Qt.UserRole) or {}) if item else {}


    def _set_row_metadata(self, row, meta):
        item = self.table.item(row, self.COL_STATUS)
        if item:
            old = self.table.blockSignals(True)
            item.setData(Qt.UserRole, dict(meta))
            editable = meta.get('source') == 'test-client'
            for col in self.EDITABLE_COLS:
                cell = self.table.item(row, col)
                if cell:
                    flags = cell.flags() & ~Qt.ItemIsEditable
                    cell.setFlags(flags | Qt.ItemIsEditable if editable else flags)
            self.table.blockSignals(old)
            self._schedule_ui_refresh()


    def _bus_target(self, row):
        values = self._row_values(row)
        result = self._row_metadata(row).get('result', {})
        # Prefer the server's executed path; never check UNC paths on the client.
        paths = [step.get('target', '') for step in result.get('steps', [])
                 if str(step.get('target', '')).startswith('\\\\')]
        path = paths[0] if paths else ('' if values['kind']=='종료' else build_path_l3(values['proj'],values['lv2'],values['lv3']))
        return {**values, 'row': row, 'path': path}


    def _on_api_progress(self, job):
        self._set_api_connection_state(True,job)
        meta = self._row_metadata(self.current_row)
        if job.get('job_id'):
            meta['job_id'] = job['job_id']
            meta['live_result'] = dict(job)
            self._set_row_metadata(self.current_row,meta)
        self._set_progress_status(f"API {job.get('status','')} / {job.get('job_id','')}",
                                 self._get(self.current_row,self.COL_PROJ),self.done_jobs,self.total_jobs,running=True)
        self._refresh_detail()


    def _on_api_result(self, result):
        self._set_api_connection_state(True,result)
        row = self.current_row
        meta = self._row_metadata(row)
        result = dict(result)
        if result.get('status') == 'validated':
            meta['validation_result'] = result
            meta['show_validation'] = True
            self._set_row_metadata(row,meta)
            self.table.item(row,self.COL_STATUS).setText(validation_label(result))
            self._log_validation(result,f'{row+1}행 · 프로젝트 {self._get(row,self.COL_PROJ)}')
            self.done_jobs += 1
            self._refresh_detail()
            return
        meta.pop('show_validation',None)
        previous = meta.get('result') or {}
        if result.get('status')=='preview':
            if previous.get('status') in {'failed','partially_succeeded'} and previous.get('executor_mode')=='powershell':
                meta['_failed_job_id'] = previous.get('job_id')
        else:
            meta.pop('_retry_parent_job_id',None)
            meta.pop('_failed_job_id',None)
        meta['result'] = result
        self._set_row_metadata(row,meta)
        status = result.get('status')
        label = {'preview':'DryRun', 'simulated':'모의완료 (BUS 미처리)',
                 'succeeded':'API 성공', 'partially_succeeded':'부분 성공 (BUS 미처리)',
                 'failed':'API 실패', 'cancelled':'취소됨'}.get(status,'알 수 없는 결과')
        self.table.item(row,self.COL_STATUS).setText(label)
        self._log(json.dumps(result,ensure_ascii=False,indent=2),
                  summary=self._api_log_summary(result,f'{row+1}행 · 프로젝트 {self._get(row,self.COL_PROJ)}'))
        if result.get('error_code') == 'POWERSHELL_TIMEOUT':
            self.table.item(row,self.COL_STATUS).setText('서버 시간 초과 (BUS 완료 보류)')
            self._log('서버 PowerShell 실행 제한을 초과하여 작업이 종료됐습니다. '
                      'API/.env의 FGT_POWERSHELL_TIMEOUT_SECONDS를 확인하세요. '
                      '실패 Job은 완료 버튼으로 성공 처리할 수 없습니다.')
        self.done_jobs += 1
        stopped = self.stop_requested or result.get('client_stopped',False)
        if (self.auto_complete_after_add and not meta.get('bus_done')
                and can_complete_bus(result,meta.get('source'),stopped)):
            # Session preparation pumps Qt events. Hold the queue before entering it.
            self._waiting_for_bus = True
            self._pending_after_add_row = row
            if self._ensure_bus_session('API 성공 후 완료 처리') and not self.stop_requested and not self._closing:
                self.table.item(row,self.COL_STATUS).setText('완료 처리중')
                self.trigger_session_process.emit([self._bus_target(row)])
            else:
                self._waiting_for_bus = False
                self._pending_after_add_row = None
                self._log('API는 성공했습니다. BUS 세션 준비 후 완료 처리만 다시 실행하세요.')
                QTimer.singleShot(0,self._start_next_job)


    def _on_api_error(self, message):
        self._set_api_connection_state(False)
        checking = bool(self.api_worker and getattr(self.api_worker,'validate_only',False))
        if checking:
            meta = self._row_metadata(self.current_row)
            meta.update(show_validation=True,validation_result=dict(status='validated',validation_status='unknown'))
            self._set_row_metadata(self.current_row,meta)
        self.table.item(self.current_row,self.COL_STATUS).setText('요청 검증 확인 불가' if checking else 'API 오류 (등록 결과 확인 필요)')
        self._log(message,summary=('요청 검증 확인 불가: '+message) if checking else 'API 요청 오류 · 상세 내용은 로그 파일을 확인하세요.')
        self.done_jobs += 1


    def _api_thread_finished(self):
        self.api_worker = None
        self.api_thread = None
        if self._closing:
            QTimer.singleShot(0,self.close)
        elif not self._waiting_for_bus:
            self._start_next_job()

    def _checkmark_path(self):
        return resource_path('assets/check.svg').replace('\\','/')


    def _preview_selected(self):
        self.chk_dry.setChecked(True)
        self.run_execute()


    def _execute_selected(self):
        self.chk_dry.setChecked(False)
        self.run_execute()


    def _toggle_log(self,expanded):
        self.log.setVisible(expanded)
        self.log_toggle.setText(('⌄' if expanded else '›')+'  작업 로그')


    def _toggle_detail_columns(self,visible):
        for column in (self.COL_NAME,self.COL_LV3,self.COL_DEPT,self.COL_ROLE):
            self.table.setColumnHidden(column,not visible)
        self.table.horizontalHeaderItem(self.COL_LV2).setText('Level2' if visible else 'Level2 / Level3')


    def _switch_workspace(self,index):
        # self._refresh_request_count_badge()
        self.workspaces.setCurrentIndex(index)
        self.actions_bar.setVisible(index==0)
        self.summary_label.setVisible(index==0)
        viewer = getattr(self,'_newdlg',None)
        if viewer is not None:
            viewer.actions_bar.setVisible(index==1)
        if index==1:
            self.open_new_viewer()


    def _schedule_ui_refresh(self,*args):
        if not self._ui_refresh_pending:
            self._ui_refresh_pending = True
            QTimer.singleShot(0,self._refresh_ui)


    def _sync_selection_label_width(self, extra=None):
        labels = [self.selection_label]
        viewer = getattr(self,'_newdlg',None)
        if viewer is not None:
            labels.append(viewer.selection_label)
        if extra is not None and extra not in labels:
            labels.append(extra)
        width = max(label.fontMetrics().horizontalAdvance(label.text()) for label in labels)+2
        for label in labels:
            label.setFixedWidth(width)

    def _refresh_ui(self):
        self._ui_refresh_pending = False
        selected = running = completed = attention = 0
        for row in range(self.table.rowCount()):
            checkbox = self._get_checkbox(row)
            if checkbox:
                selected += int(checkbox.isChecked())
                if not checkbox.property('summaryConnected'):
                    checkbox.toggled.connect(self._schedule_ui_refresh)
                    checkbox.setProperty('summaryConnected',True)
            status = self._get(row,self.COL_STATUS)
            text,tone = request_status_presentation(status)
            running += int(tone=='blue' and text!='요청 검증')
            completed += int(tone=='green')
            attention += int(tone in {'red','amber'})
        self.selection_label.setText(f'선택 {selected}건')
        self._sync_selection_label_width()
        self.table.horizontalHeader().setChecked(bool(self.table.rowCount()) and selected==self.table.rowCount())
        self.summary_label.setText(f'전체 {self.table.rowCount()}건 · 실행 중 {running} · 완료 {completed} · 확인 필요 {attention}')
        self._fit_table_columns()
        self._refresh_detail()

    def _fit_table_columns(self):
        if not self.table.isColumnHidden(self.COL_LV3):
            return
        width = max(620,self.table.viewport().width())
        for column,fraction in [(self.COL_SELECT,.05),(self.COL_KIND,.09),(self.COL_REQTYPE,.13),
                                (self.COL_USER,.115),(self.COL_PROJ,.145),(self.COL_LV2,.265)]:
            self.table.setColumnWidth(column,int(width*fraction))

    def resizeEvent(self,event):
        super().resizeEvent(event)
        if hasattr(self,'table'):
            QTimer.singleShot(0,self._fit_table_columns)


    def _refresh_detail(self):
        row = self.table.currentRow()
        valid = 0 <= row < self.table.rowCount()
        values = self._row_values(row) if valid else {}
        meta = self._row_metadata(row) if valid else {}
        result = meta.get('result') or {}
        active = self._get(row,self.COL_STATUS) in {'API 실행중','요청 검증 중'}
        if active:
            result = meta.get('live_result') or (result if result.get('status')=='running' else {})
        self.detail_project.setText(values.get('proj') or '-')
        self.detail_employee.setText(values.get('user') or '-')
        access = values.get('req','')
        folder = ' / '.join(x for x in (values.get('lv2'),values.get('lv3')) if x)
        self.detail_access.setText(f'{access} · {folder}' if folder else access or '요청을 선택하세요')
        self.detail_role.setText(values.get('role') or '-')
        self.detail_name.setText(self._get(row,self.COL_NAME) if valid else '')
        self.detail_department.setText(self._get(row,self.COL_DEPT) if valid else '')
        extra = ' · '.join(x for x in (self.detail_name.text(),self.detail_department.text()) if x)
        self.detail_extra.setText(extra); self.detail_extra.setVisible(bool(extra))
        validation = meta.get('validation_result') if meta.get('show_validation') else None
        raw_steps = result.get('steps') or []
        statuses = ['queued']*4
        if result.get('job_id') or result.get('valid'):
            statuses[0] = 'succeeded'
        if valid and '검증실패' in self._get(row,self.COL_STATUS):
            statuses[0] = 'failed'
        if active and self._get(row,self.COL_STATUS)=='요청 검증 중':
            statuses[0] = 'running'
        elif active and result.get('status')=='submitting':
            statuses[0] = 'running'
        if validation and not active:
            statuses[0] = 'succeeded' if validation.get('validation_status')=='passed' else 'failed' if validation.get('validation_status')=='failed' else 'queued'
        groups = [s for s in raw_steps if any(k in s.get('type','') for k in ('group','member'))]
        acls = [s for s in raw_steps if s not in groups]
        def phase(steps):
            if any(s.get('status')=='failed' for s in steps): return 'failed'
            if any(s.get('status')=='running' for s in steps): return 'running'
            if steps and all(s.get('status')=='succeeded' for s in steps): return 'succeeded'
            return 'queued'
        statuses[1] = phase(groups); statuses[2] = phase(acls)
        if meta.get('bus_done'):
            statuses[3] = 'succeeded'
        elif valid and self._get(row,self.COL_STATUS)=='완료 처리중':
            statuses[3] = 'running'
        labels = ['요청 검증','그룹 권한 적용','폴더 권한 적용','BUS 완료 처리']
        if valid and meta.get('source')!='bus': labels[3] += ' · 수동 요청 제외'
        elif values.get('kind')=='종료': labels[1] += ' · 해당 없음'
        for index,(icon,label,state) in enumerate(zip(self.step_icons,self.step_labels,statuses)):
            icon.setText('✓' if state=='succeeded' else str(index+1))
            label.setText(labels[index])
            for widget in (icon,label):
                widget.setProperty('state',state)
                widget.style().unpolish(widget); widget.style().polish(widget)
        if valid:
            raw = self._get(row,self.COL_STATUS)
            shown,tone = request_status_presentation(raw)
            self.detail_banner_title.setText('요청 검증 중' if active and raw=='요청 검증 중' else '서버 작업 진행 중' if tone=='blue' and 'BUS' not in shown else shown)
            message = result.get('error_message') or result.get('error_code') or ''
            if result.get('status')=='running':
                elapsed = self._format_elapsed()
                message = f'이번 요청 경과 {elapsed} · 최종 결과 대기'
            if not message:
                message = '서버 실제 성공 후에만 BUS 완료 가능' if meta.get('source')=='bus' else '선택한 요청의 서버 결과를 표시합니다.'
            if validation:
                message = '아래 작업 로그에서 각 항목의 실제 조회 결과를 확인하세요. 검증은 권한을 변경하지 않습니다.'
            self.status_label.setText(message)
        else:
            self.detail_banner_title.setText('선택 요청 대기')
            if not self._progress_running:
                self.status_label.setText('목록에서 요청을 선택하세요.')
        self.progress_panel.show()


    def _refresh_connection_strip(self):
        t = THEMES.get(self.current_theme,THEMES['light'])
        api_ok = getattr(self,'_api_verified',False)
        self.api_status_label.setText('●  API 연결됨' if api_ok else '○  API 미확인')
        self.api_status_label.setStyleSheet(f"color:{t['green'] if api_ok else t['muted']};")
        # session = getattr(self,'session',None); watcher = getattr(self,'watch_session',None)
        # bus_ok = bool((session and session.is_ready()) or (watcher and watcher.is_ready()))
        session = getattr(self,'session',None)
        bus_ok = bool(session and session.is_ready())
        self.bus_status_label.setText('●  BUS 로그인됨' if bus_ok else '○  BUS 미연결')
        self.bus_status_label.setStyleSheet(f"color:{t['green'] if bus_ok else t['muted']};")
        parsed = urlsplit(getattr(self,'api_config',{}).get('base_url',DEFAULT_API_URL))
        self.server_label.setText(parsed.netloc or '-')



    def _set_api_connection_state(self,connected,result=None):
        self._api_verified = bool(connected)
        if result and result.get('executor_mode') in {'mock','powershell'}:
            self._executor_mode = result['executor_mode']
        self._refresh_connection_strip()



def _excepthook(etype, value, tb):
    try:
        import traceback
        msg = "".join(traceback.format_exception(etype, value, tb))
        print(msg)
    except:
        pass

    QMessageBox.critical(None, "치명적 오류", str(value))

import sys as _sys
if __name__ == "__main__":
    _sys.excepthook = _excepthook

if __name__ == "__main__":
    QApplication.setAttribute(Qt.AA_EnableHighDpiScaling, True)
    QApplication.setAttribute(Qt.AA_UseHighDpiPixmaps, True)

    app = QApplication(sys.argv)
    app.setFont(QFont("Segoe UI", 9))
    window = AccessManager(startup=True)
    startup_dialog = StartupLoadingDialog(window)
    def _on_quit():
        try:
            window.close()
        except Exception:
            pass
    startup_dialog.show()
    QTimer.singleShot(0,startup_dialog.start)
    sys.exit(app.exec_())
