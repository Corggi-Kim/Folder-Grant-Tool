"""Generate the three Office beta documents. Run from any directory.

Dependencies: python-docx, python-pptx, Pillow.
Input: beta_content.py and docs/beta-assets/{access,projects,startup}.png.
No BUS/API connection or external operation is performed.
"""
from pathlib import Path
import sys
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PC
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from beta_content import DATE, BASE, API_PAGES, CLIENT_PAGES

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT/'docs'
ASSETS = OUT/'beta-assets'
NAVY='14233D'; BLUE='2962FF'; INK='23324A'; MUTED='60718A'; PALE='EDF3FF'; BG='F4F6FA'; BORDER='DBE3F0'; TEAL='00856B'
FONT='Noto Sans CJK KR'


def architecture():
    img=Image.new('RGB',(1800,610),'white');d=ImageDraw.Draw(img)
    font='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
    bold='/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc'
    f=lambda n,b=False:ImageFont.truetype(bold if b else font,n)
    def box(x,y,w,h,title,body,color):
        d.rounded_rectangle((x,y,x+w,y+h),18,fill=color,outline='#DBE3F0',width=2)
        d.text((x+26,y+22),title,font=f(36,True),fill='#23324A')
        for i,line in enumerate(body):d.text((x+26,y+80+i*42),line,font=f(26),fill='#60718A')
    def arrow(a,b,label='',above=True):
        d.line((a,b),fill='#2962FF',width=5)
        x,y=b;sign=1 if b[0]>a[0] else -1
        d.polygon([(x,y),(x-sign*17,y-10),(x-sign*17,y+10)],fill='#2962FF')
        if label:d.text(((a[0]+b[0])/2-20,(a[1]+b[1])/2-38 if above else (a[1]+b[1])/2+12),label,font=f(24),fill='#2962FF',anchor='mm')
    box(25,110,350,235,'BUS',['요청 목록','완료 처리'],'#F4F6FA')
    box(495,110,410,235,'Client',['조회 · 선택 · 결과 확인','BUS 연결 / API 요청'],'#EDF3FF')
    box(1040,110,370,235,'API Server',['정책 · 계획 · Job 실행','공유 API Key 인증'],'#EDF3FF')
    box(1480,110,295,235,'AD / 폴더',['그룹 · ACL 변경','템플릿 복사'],'#F4F6FA')
    arrow((375,200),(495,200),'조회')
    arrow((495,285),(375,285),'완료',False)
    arrow((905,200),(1040,200),'요청')
    arrow((1040,285),(905,285),'결과',False)
    arrow((1410,225),(1480,225))
    box(1040,435,370,150,'SQLite',['Job · 단계 · 감사 기록'],'#F4F6FA')
    d.line((1225,345,1225,435),fill='#2962FF',width=5)
    img.save(ASSETS/'architecture.png')


def shade(cell,fill):
    el=OxmlElement('w:shd');el.set(qn('w:fill'),fill);cell._tc.get_or_add_tcPr().append(el)


def font_run(r,size=10.5,bold=False,color=INK):
    r.font.name=FONT;r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)


def paragraph(doc,text='',style=None,size=10.5,bold=False,color=INK):
    p=doc.add_paragraph(style=style)
    p.paragraph_format.space_after=Pt(8)
    p.paragraph_format.line_spacing=1.25
    font_run(p.add_run(text),size,bold,color)
    return p


def table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    width=6.66
    widths=([2.18,4.48] if len(headers)==2 else [width/len(headers)]*len(headers))
    for i,h in enumerate(headers):
        cell=t.rows[0].cells[i];cell.width=Inches(widths[i]);shade(cell,NAVY)
        p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
        font_run(p.add_run(h),9.5,True,'FFFFFF')
    # Repeat the heading if an unusually long table crosses a page.
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for n,row in enumerate(rows):
        cells=t.add_row().cells
        for i,text in enumerate(row):
            cell=cells[i];cell.width=Inches(widths[i]);shade(cell,BG if n%2==0 else 'FFFFFF')
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5);p.paragraph_format.line_spacing=1.12
            font_run(p.add_run(text),9.4,i==0)
        no_split=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(no_split)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


def callout(doc,title,text):
    t=doc.add_table(rows=1,cols=1);cell=t.cell(0,0);shade(cell,PALE)
    p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(4)
    font_run(p.add_run(title),10.3,True,BLUE)
    p=cell.add_paragraph();p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1.18
    font_run(p.add_run(text),9.8)


def office_doc(name,title,subtitle,pages):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=Mm(19);sec.bottom_margin=Mm(18);sec.left_margin=Mm(20);sec.right_margin=Mm(20)
    sec.header_distance=Mm(8);sec.footer_distance=Mm(8)
    sec.different_first_page_header_footer=True
    normal=doc.styles['Normal'];normal.font.name=FONT;normal.font.size=Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    for sty in ('Heading 1','Heading 2'):
        doc.styles[sty].font.name=FONT;doc.styles[sty].font.color.rgb=RGBColor.from_string(NAVY)
    hp=sec.header.paragraphs[0];font_run(hp.add_run(f'FOLDER GRANT  /  {title}'),8.5,False,MUTED)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font_run(fp.add_run(f'내부 베타 배포 · {DATE}    |    '),8,False,MUTED)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
    paragraph(doc,'BETA DOCUMENTATION',size=11,bold=True,color=BLUE)
    paragraph(doc,'',size=26)
    paragraph(doc,'Folder Grant',size=34,bold=True,color=NAVY)
    paragraph(doc,title,size=29,bold=True,color=NAVY)
    paragraph(doc,subtitle,size=14,color=MUTED)
    paragraph(doc,'',size=22)
    callout(doc,'설계에서 장애 대응까지','현재 구현을 기준으로 구조, 설정, 작업 흐름과 유지보수 시 확인할 사항을 정리했습니다. 베타 배포 전 점검과 인수인계에 함께 사용하세요.')
    paragraph(doc,'',size=15)
    table(doc,['문서 정보','내용'],[
        ['문서 버전','1.0 · Beta'],['작성일',DATE],['프로그램 기준','API 0.4.0 / Client V1.1.0 Build:20261022'],
        ['소스 기준',BASE[:12]],['자료 성격','현재 구현 + 명시된 운영 권장 사항'],['데이터 원칙','실제 계정·API Key 미포함 / 화면은 오프라인 샘플']])
    paragraph(doc,'자동 테스트는 외부 Windows·BUS·AD 환경의 검증을 대신하지 않습니다.',size=9.5,color=MUTED)
    for num,(heading,blocks) in enumerate(pages,1):
        doc.add_page_break()
        paragraph(doc,f'{num:02d}  {heading}',style='Heading 1',size=20,bold=True,color=NAVY)
        for kind,data in blocks:
            if kind=='p':paragraph(doc,data)
            elif kind=='h':
                p=paragraph(doc,data,style='Heading 2',size=12,bold=True,color=BLUE)
                p.paragraph_format.space_before=Pt(7);p.paragraph_format.keep_with_next=True
            elif kind=='table':table(doc,*data)
            elif kind=='callout':callout(doc,*data)
            elif kind=='steps':
                for i,text in enumerate(data,1):
                    p=paragraph(doc,f'{i:02d}  {text}',size=10.3)
                    p.paragraph_format.left_indent=Mm(2)
            elif kind=='code':
                p=paragraph(doc,data,size=8.4,color=INK)
                p.paragraph_format.line_spacing=1.1
                el=OxmlElement('w:shd');el.set(qn('w:fill'),BG);p._p.get_or_add_pPr().append(el)
            elif kind=='image':
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(ASSETS/data[0]),width=Inches(data[1]))
    doc.core_properties.title=f'Folder Grant {title}'
    doc.core_properties.subject='Beta 설계·운영·유지보수 기준서'
    doc.core_properties.author='Folder Grant Project'
    doc.core_properties.keywords='Beta, API, Client, 설계, 운영, 유지보수'
    doc.save(OUT/name)


def pcolor(value):return PC.from_string(value)


def rect(sl,x,y,w,h,fill=BG,line=None,rounded=False):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,PI(x),PI(y),PI(w),PI(h))
    sh.fill.solid();sh.fill.fore_color.rgb=pcolor(fill)
    if line:sh.line.color.rgb=pcolor(line)
    else:sh.line.fill.background()
    if rounded:
        try:sh.adjustments[0]=.08
        except (AttributeError,IndexError):pass
    return sh


def text(sl,x,y,w,h,value,size=18,color=INK,bold=False):
    box=sl.shapes.add_textbox(PI(x),PI(y),PI(w),PI(h));tf=box.text_frame
    tf.word_wrap=True;tf.margin_left=0;tf.margin_top=0;tf.margin_bottom=0;tf.margin_right=0
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.space_after=PP(8);p.font.name=FONT;p.font.size=PP(size);p.font.bold=bold;p.font.color.rgb=pcolor(color)
        p.text=line
    return box


def card(sl,x,y,w,h,kicker,title,body,dark=False):
    rect(sl,x,y,w,h,'20334F' if dark else 'FFFFFF',rounded=True)
    text(sl,x+.23,y+.22,w-.46,.35,kicker,11,'8DB1FF' if dark else BLUE,True)
    text(sl,x+.23,y+.7,w-.46,.65,title,22,'FFFFFF' if dark else NAVY,True)
    text(sl,x+.23,y+1.48,w-.46,h-1.65,body,16,'D1DDEF' if dark else MUTED)


def presentation():
    prs=Presentation();prs.slide_width=PI(13.333333);prs.slide_height=PI(7.5)
    def slide(title,subtitle='',dark=False):
        sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=pcolor(NAVY if dark else BG)
        text(sl,.5,.25,9,.3,'FOLDER GRANT  /  BETA',10,'8DB1FF' if dark else BLUE,True)
        text(sl,.5,.83,12.2,.65,title,30,'FFFFFF' if dark else NAVY,True)
        text(sl,.5,1.57,12.2,.55,subtitle,15,'B7C7DF' if dark else MUTED)
        text(sl,.5,7.12,11,.22,f'내부 베타 소개  ·  {DATE}  ·  API 0.4.0 / Client 1.1.0',8.5,'91A6C4' if dark else MUTED)
        text(sl,12.3,7.08,.55,.3,f'{len(prs.slides):02d}',10,'91A6C4' if dark else MUTED)
        return sl
    sl=slide('업무 요청부터 권한 적용까지','Folder Grant API + Client 프로그램 소개',True)
    text(sl,.65,2.6,8.2,1.5,'익숙한 BUS 업무 흐름에\n서버 중심의 권한 처리를 더합니다.',34,'FFFFFF',True)
    text(sl,.65,4.55,7.7,1.25,'권한 요청 조회 · 신규 프로젝트 생성 · 결과 확인 · BUS 완료\n제한된 시험 대상부터 검증하는 Beta 릴리스',19,'B7C7DF')
    rect(sl,9.5,2.65,2.85,2.85,'2962FF',rounded=True)
    text(sl,9.82,3.12,2.3,.6,'BETA',34,'FFFFFF',True)
    text(sl,9.82,4.05,2.2,1.2,'API 0.4.0\nClient 1.1.0',17,'FFFFFF')

    sl=slide('업무 화면과 권한 실행의 역할을 나눴습니다','사용자는 Client에서 요청을 확인하고, 서버는 정해진 정책으로 작업을 처리합니다.')
    card(sl,.5,2.4,3.95,3.75,'01  사용자 경험','조회부터 완료까지','BUS 조회·완료 흐름을 유지하고\n수동·Excel 입력도 지원합니다.')
    card(sl,4.69,2.4,3.95,3.75,'02  실행 기준','권한 정책을 서버에','AD 그룹과 폴더 ACL 처리를\nAPI 서버에서 수행합니다.')
    card(sl,8.88,2.4,3.95,3.75,'03  결과 확인','단계별 이력 확인','일부 단계가 실패해도 원인과\n이전 성공 결과를 확인합니다.')

    sl=slide('두 프로그램이 하나의 업무 흐름을 만듭니다','Client는 BUS와 통신하고, API 서버는 AD·공유 폴더 작업과 이력을 담당합니다.')
    sl.shapes.add_picture(str(ASSETS/'architecture.png'),PI(.5),PI(2.25),width=PI(12.3))
    text(sl,.7,6.65,11.9,.4,'BUS 완료는 API 작업 성공과 별도 단계이며, 실제 성공을 확인한 요청만 처리합니다.',15,BLUE,True)

    sl=slide('API와 Client, 맡은 일이 명확합니다','배포·운영 담당자는 API를, 베타 사용자는 Client를 중심으로 확인합니다.')
    card(sl,.5,2.35,6.0,3.95,'SERVER  /  API','실행과 이력의 중심','입력 검증 · 실행 계획 생성\n권한 부여·해제 · 신규 프로젝트 생성\nJob 상태 · 단계 결과 · 감사 기록')
    card(sl,6.8,2.35,6.0,3.95,'DESKTOP  /  CLIENT','업무 요청의 시작과 끝','BUS 조회 · 수동/Excel 입력\n미리보기 · 선택 실행 · 중지\n결과 확인 · 실제 성공 후 BUS 완료')

    sl=slide('처음 사용해도 작업 순서가 보입니다','시작 시 연결과 요청 목록을 준비하고, 필요한 항목을 선택해 처리합니다.')
    steps=[('01','연결·로그인','API 확인과 BUS 로그인'),('02','요청 확인','두 탭의 목록·상세 확인'),('03','미리보기','대상과 계획을 먼저 확인'),('04','선택 실행','체크한 항목을 서버 처리'),('05','완료 확인','결과 검증 후 BUS 완료')]
    for i,(num,title,body) in enumerate(steps):
        x=.5+i*2.5;rect(sl,x,2.6,2.3,2.55,'FFFFFF',rounded=True)
        text(sl,x+.2,2.85,1.8,.5,num,26,BLUE,True)
        text(sl,x+.2,3.6,1.9,.45,title,19,NAVY,True)
        text(sl,x+.2,4.22,1.9,.7,body,13,MUTED)
    text(sl,.7,5.7,11.8,.7,'베타 첫 실제 시험은 자동 BUS 완료를 끄고, 대상과 실제 권한을 확인한 후 완료 처리합니다.',19,INK,True)

    sl=slide('권한 요청: 목록과 상세를 한 화면에','진행·종료, 권한 부여·해제 요청을 선택하고 결과를 확인합니다.')
    sl.shapes.add_picture(str(ASSETS/'access.png'),PI(.5),PI(2.15),height=PI(4.76))
    card(sl,6.45,2.4,6.35,3.9,'CLIENT  /  권한 요청','선택하고, 확인하고, 실행','수동 입력 → 엑셀 불러오기 → BUS 요청 조회\n행 선택: 오른쪽에서 대상·진행·오류 확인\n미리보기 / 선택 실행 / BUS 완료 재확인 / 중지')
    text(sl,6.65,6.5,5.9,.35,'화면은 실제 Qt 앱의 오프라인 샘플 데이터입니다.',11,MUTED)

    sl=slide('신규 프로젝트: 같은 구조로 쉽게 처리','프로젝트를 선택하면 오른쪽 상세에서 원본 정보를 확인할 수 있습니다.')
    sl.shapes.add_picture(str(ASSETS/'projects.png'),PI(.5),PI(2.15),height=PI(4.76))
    card(sl,6.45,2.4,6.35,3.9,'CLIENT  /  신규 프로젝트','그룹·폴더·기본 권한 준비','수동 추가 또는 BUS 신규조회\n템플릿 폴더 복사와 기본 ACL 설정\nDry Run 확인 후 실제 실행 · 결과 검증')
    text(sl,6.65,6.5,5.9,.35,'화면은 샘플 데이터이며, 생성 작업은 서버에서 수행합니다.',11,MUTED)

    sl=slide('미리보기·Mock·실제 실행은 다릅니다','검증 목적에 맞는 모드를 사용하고, Mock과 실제 실행용 DB를 분리합니다.')
    card(sl,.5,2.4,3.95,3.9,'PLAN','미리보기','실행 계획만 확인\n실제 자원 변경 없음\nJob 등록·BUS 완료 없음')
    card(sl,4.69,2.4,3.95,3.9,'SIMULATION','Mock 서버','Job·DB·상태 흐름 시험\nAD·ACL 변경 없음\nBUS 완료 없음')
    card(sl,8.88,2.4,3.95,3.9,'EXECUTION','PowerShell 서버','실제 그룹·권한·폴더 변경\n테스트 자원부터 검증\n실제 성공 후 BUS 완료 가능')

    sl=slide('성공 여부와 BUS 완료 여부를 구분합니다','부분 성공과 조회 오류를 완료로 처리하지 않습니다.',True)
    for i,(title,body,col) in enumerate([
        ('실제 성공','BUS 원본 요청이면\n완료 처리가 가능합니다.','78DFBE'),
        ('부분 성공·실패','오류와 실제 권한 확인 후\n원인을 조치합니다.','FFC76C'),
        ('결과 확인 필요','기존 Job 상태를 확인하고\n중복 실행을 피합니다.','9AB8FF')]):
        x=.5+i*4.2;rect(sl,x,2.55,3.95,2.7,'20334F',rounded=True)
        text(sl,x+.24,2.85,3.45,.65,title,23,col,True)
        text(sl,x+.24,3.75,3.45,1.1,body,18,'FFFFFF')
    text(sl,.75,5.85,11.8,.75,'BUS 완료 재확인은 서버 결과를 확인한 뒤 BUS 완료만 시도하며, 권한 작업을 다시 실행하지 않습니다.',19,'D1DDEF')

    sl=slide('실패한 권한 작업은 원인 조치 후 재시도','같은 요청을 다시 보내는 것과 실패 단계를 재실행하는 것은 다릅니다.')
    card(sl,.5,2.4,3.95,3.7,'01  조사','실패 단계 확인','오른쪽 상세와 파일 로그에서\n경로·오류·Job을 확인합니다.')
    card(sl,4.69,2.4,3.95,3.7,'02  조치','원인을 먼저 해결','누락 폴더·접근 권한 등을\n서버 담당자와 점검합니다.')
    card(sl,8.88,2.4,3.95,3.7,'03  재시도','선택 실행 → 확인','새 Job으로 실패·미실행 단계 처리\n이전 성공 단계와 이력은 보존합니다.')
    text(sl,.7,6.45,11.9,.35,'대상: 실제 PowerShell의 실패·부분 성공 권한 Job. 프로젝트·Mock·진행 중 Job은 제외됩니다.',12,MUTED)

    sl=slide('설치와 운영에 필요한 준비','Client PC와 API 서버의 준비 항목을 각각 확인합니다.')
    card(sl,.5,2.35,6.0,3.95,'사용 PC','Client 환경','Windows · Python 3.11+ · Chrome\nBUS와 API 주소로 연결 가능한 네트워크\nAPI Key · BUS 계정 · 로컬 기록 폴더 쓰기 권한')
    card(sl,6.8,2.35,6.0,3.95,'실행 서버','API 환경','Windows · AD PowerShell 모듈\n도메인 서비스 계정 · AD/공유/NTFS 권한\nAPI 포트 · 템플릿 · DB/로그 백업')
    text(sl,.7,6.5,11.9,.35,'저장 계정·API Key는 평문 파일로 보관됩니다. 접근 범위를 제한하고 공유 배포본에 포함하지 마세요.',12,MUTED)

    sl=slide('베타는 작은 범위에서 확인하며 엽니다','참여 대상과 시험 자원을 정하고, 단계별로 실제 동작을 검증합니다.')
    for i,(n,title,body) in enumerate([
        ('1','연결·계획 검증','Mock + Preview\n입력·경로·정책 확인'),
        ('2','실제 1건 시험','자동 완료 OFF\n실제 AD·ACL 직접 확인'),
        ('3','실패·복구 검증','부분 성공·시간 초과\n재시도·BUS 완료 확인'),
        ('4','제한된 베타 운영','대상·문의 경로 지정\n로그와 사용자 의견 수집')]):
        x=.5+i*3.15;rect(sl,x,2.45,2.93,3.75,'FFFFFF',rounded=True)
        text(sl,x+.22,2.75,2.4,.6,n,30,BLUE,True)
        text(sl,x+.22,3.6,2.48,.6,title,21,NAVY,True)
        text(sl,x+.22,4.5,2.48,1.1,body,16,MUTED)
    text(sl,.7,6.48,11.8,.35,'자동 검증: API 60개 / Client 93개 통과. 실제 Windows·BUS·AD 연동은 별도 시험 대상입니다.',12,MUTED)

    sl=slide('현재 제한을 알고, 다음 개선을 준비합니다','자동 복구와 권한 분리는 후속 개선 과제입니다. 아래 항목은 구현 완료 기능이 아닙니다.',True)
    card(sl,.5,2.4,6.0,3.95,'현재 BETA','운영 시 알아둘 점','API 재시작 후 진행 Job 자동 복구 없음\n취소는 단계 사이에서 확인 · 자동 롤백 없음\n공유 API Key · 로그 보관/정리 수동 관리',True)
    card(sl,6.8,2.4,6.0,3.95,'개선 제안','검증 결과에 따라 우선순위 결정','별도 Worker · 재시작 복구\n사용자 인증·권한 분리 · 실행 동시성 관리\n모니터링 · 로그 보관 · 배포 체계',True)

    sl=slide('도입·운영·유지보수 자료를 함께 제공합니다','docs/ 폴더의 문서를 기준으로 설치하고, 베타 결과를 기록하세요.')
    card(sl,.5,2.4,6.0,3.95,'WORD  /  SERVER','API 설계·운영·유지보수','구조와 정책 · API 계약\n설정·서비스 배포 · 백업·장애 대응\n변경 검증 · 베타 오픈 확인 목록')
    card(sl,6.8,2.4,6.0,3.95,'WORD  /  DESKTOP','Client 설계·사용·유지보수','설치·로그인 · 화면·작업 절차\n결과·재시도 · 로그·문제 해결\n내부 구조 · 변경 검증 · 참여 안내')
    text(sl,.7,6.5,11.9,.35,'피드백에는 발생 시각·버전·기대/실제 결과·request_id/job_id를 포함하고 키·암호는 제거하세요.',12,MUTED)
    prs.core_properties.title='Folder Grant API + Client Beta 프로그램 소개'
    prs.core_properties.author='Folder Grant Project'
    prs.core_properties.subject='프로그램 소개와 베타 도입 안내'
    prs.save(OUT/'Folder_Grant_프로그램_소개서_Beta.pptx')


def main():
    ASSETS.mkdir(exist_ok=True,parents=True)
    for name in ('access.png','projects.png','startup.png'):
        if not (ASSETS/name).is_file():raise SystemExit(f'Missing input: {ASSETS/name}')
    architecture()
    office_doc('API_설계_운영_유지보수_Beta.docx','API 설계·운영·유지보수','서버 설치 담당자 · 운영 담당자 · 유지보수 개발자를 위한 기준서',API_PAGES)
    office_doc('Client_설계_사용_유지보수_Beta.docx','Client 설계·사용·유지보수','베타 사용자 · 배포 담당자 · 유지보수 개발자를 위한 안내서',CLIENT_PAGES)
    presentation()
    print('Created 2 Word documents and 1 PowerPoint in',OUT)

if __name__=='__main__':main()
