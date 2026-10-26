"""Build release copies from manually edited Beta Office inputs; never overwrite inputs.

Default refuses to overwrite existing outputs. --refresh allows regeneration only
when outputs still match release-manifest.json. For intentional manual edits,
edit Office files directly or generate into a different --output directory.
"""
from pathlib import Path
import argparse
import hashlib
import json
import shutil
import sys
from tempfile import TemporaryDirectory
from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
import build_beta_documents as style
from release_content import DATE, BASE, API_ADDITIONS, CLIENT_ADDITIONS, HANDOVER_PAGES

ROOT=Path(__file__).resolve().parents[2]
DOCS=ROOT/'docs'
FONT='Noto Sans CJK KR'
OUTPUT_NAMES=['API_설계_운영_유지보수.docx','Client_설계_사용_유지보수.docx','오픈_운영_인수인계.docx',
              'Folder_Grant_프로그램_소개서.pptx','Folder_Grant_개발완료_보고서.pptx']
COMMON={
'2026-10-23':DATE,'26f3f761d4f1':BASE[:12],'26f3f76':BASE[:7],
'1.0 · Beta':'2.0 · 내부 오픈 기준','BETA DOCUMENTATION':'RELEASE DOCUMENTATION',
'베타 배포 전 점검과 인수인계에 활용하세요.':'내부 오픈 점검과 운영 인수인계에 활용하세요.',
'베타 사용자':'업무 사용자','베타 안내서':'운영 안내서','베타 버전의 핵심 가치':'현재 구현의 핵심 가치',
'유지보수·베타 점검':'유지보수·오픈 점검','베타 오픈':'내부 오픈','베타 점검':'오픈 점검',
'베타 참여 안내':'업무 오픈 안내','현재 베타의 제한 사항':'현재 운영 시의 제한 사항',
'베타 배포':'내부 배포','제한된 베타':'제한된 내부 운영','베타 첫':'오픈 전 첫',
'API 자동 테스트 60개가 모두 통과했습니다.':'API 자동 테스트 76개가 통과했으며 선택 PowerShell 시험 5개는 건너뛰었습니다.',
'Client 자동 테스트 93개가 모두 통과했습니다.':'Client 자동 테스트 125개가 모두 통과했습니다.',
'사용 안내':'도움말',
}
CLIENT={
'미리보기를 눌러 서버가 계산한 경로·그룹·실행 계획을 확인합니다. 실제 권한은 바뀌지 않습니다.':'요청 검증을 눌러 서버에서 실제 사용자·그룹·요청 폴더·현재 ACL을 조회합니다. 결과는 항목별로 화면 로그에 표시하며 실제 권한은 바뀌지 않습니다.',
'수동 요청으로 전환되며 이전 결과는 무효화됨. BUS 완료 연결도 해제됨':'BUS 조회 행은 읽기 전용으로 직접 편집할 수 없음. 수동·직접 Excel 입력 행만 편집 가능',
'톱니바퀴 메뉴에서 이름·부서·Role 등 편집 열을 표시':'공통 토글로 권한 목록의 상세 열과 신규 프로젝트의 BUS 원본 상세 열을 함께 표시',
'미리보기(Dry Run)를 켜고 선택 실행하여 생성 계획을 확인합니다.':'요청 검증 버튼으로 그룹·생성 위치·템플릿 구조 등 실제 준비 상태를 확인합니다.',
'실제 시험은 Dry Run을 해제하고, 실제 성공 후 BUS 완료 옵션을 끈 상태에서 실행합니다.':'첫 실제 시험은 성공 후 BUS 자동 완료를 끄고 선택 실행합니다.',
'실제 자원을 검증한 뒤 BUS 완료 옵션을 켜고 같은 항목을 다시 실행하면, 보관된 성공 결과를 이용해 BUS 완료만 재시도합니다.':'실제 자원을 확인한 뒤 BUS 완료 재확인을 누릅니다. 서버의 최신 성공을 새로 조회하고 BUS 완료만 시도하며 프로젝트를 재생성하지 않습니다.',
'미리보기 / DryRun':'요청 검증 / validated',
'계획만 확인. 실제 실행·BUS 완료 없음':'실제 자원을 읽기 전용 조회. Job 등록·실제 변경·BUS 완료 없음',
'권한 탭은 BUS 완료 재확인을 실행. 프로젝트는 완료 옵션을 켜고 같은 성공 항목을 다시 처리':'양쪽 탭 모두 BUS 완료 재확인을 실행. 최신 서버 성공을 조회한 뒤 BUS 완료만 시도',
'수동·Excel·BUS 요청의 사번·코드·경로·Role을 Preview로 비교':'수동·Excel·BUS 입력과 실제 자원을 요청 검증으로 확인. Preview는 개발용 계획 API로 구분',
'수동·Preview·Mock 요청':'수동·요청 검증·Preview·Mock 요청',
'베타의 의미':'개발 완료와 업무 인수의 구분',
'베타를 배포할 때마다':'내부 배포할 때마다',
'BUS 원본 번호가 있으면 그대로 사용하고, 없으면 원본 필드·대상·작업·경로·Role 등을 해시하여 request_id를 만듭니다.':'BUS 원본 번호가 있으면 우선 식별 재료로 사용하고 없으면 요청일 등 원본 필드를 사용합니다. 대상·작업·경로·Role 등과 함께 SHA-256 해시를 만들어 BUS-<해시>를 API request_id로 사용합니다.',
'현재 기능을 제한된 대상에서 검증하는 단계입니다.':'현재 합의된 기능을 구현했으며 오픈 전 제한된 대상의 현장 인수 시험이 필요합니다.',
}
API={
'운영·장애 → 유지보수·베타 점검':'운영·장애 → 유지보수·오픈 점검',
'Preview로 검증':'Preview로 입력·계획을 확인하고 /validate로 실제 자원을 조회',
'cancel_requested를 거쳐 cancelled':'cancel_requested 플래그가 설정된 뒤 cancelled',
}


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def replace_paragraph(p,mapping):
    before=p.text
    after=before
    for old,new in mapping.items():after=after.replace(old,new)
    if after==before:return
    # Preserve run formatting where possible, including edits made in Office.
    individual=''.join(replaced_run(r.text,mapping) for r in p.runs)
    if individual==after:
        for r in p.runs:r.text=replaced_run(r.text,mapping)
    elif p.runs:
        p.runs[0].text=after
        for r in p.runs[1:]:r.text=''
    else:p.add_run(after)


def replaced_run(s,mapping):
    for old,new in mapping.items():s=s.replace(old,new)
    return s


def paragraphs(container):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:yield from paragraphs(cell)


def add_blocks(doc,title,blocks):
    doc.add_page_break()
    style.paragraph(doc,title,style='Heading 1',size=20,bold=True,color=style.NAVY)
    for kind,data in blocks:
        if kind=='p':style.paragraph(doc,data)
        elif kind=='h':style.paragraph(doc,data,style='Heading 2',size=12,bold=True,color=style.BLUE)
        elif kind=='table':style.table(doc,*data)
        elif kind=='callout':style.callout(doc,*data)
        elif kind=='steps':
            for i,text in enumerate(data,1):style.paragraph(doc,f'{i:02d}  {text}',size=10.3)


def swap_images(package,assets):
    matches={}
    for filename in ('access.png','access-dark.png','projects.png'):
        old=DOCS/'beta-assets'/filename
        new=assets/filename
        if old.exists() and new.exists():matches[sha(old)]=new.read_bytes()
    for part in (package.parts if hasattr(package,'parts') else package.iter_parts()):
        if part.content_type.startswith('image/'):
            replacement=matches.get(hashlib.sha256(part.blob).hexdigest())
            if replacement:part._blob=replacement


def update_doc(kind,out,assets):
    source=DOCS/f'{kind}_설계_{"운영" if kind=="API" else "사용"}_유지보수_Beta.docx'
    d=Document(source)
    mapping={**COMMON,**(API if kind=='API' else CLIENT)}
    for p in paragraphs(d):replace_paragraph(p,mapping)
    for section in d.sections:
        for block in (section.header,section.footer,section.first_page_header,section.first_page_footer):
            for p in paragraphs(block):replace_paragraph(p,{**mapping,'내부 베타':'내부 오픈'})
    # Insert new contracts alongside the preserved original endpoint/config tables.
    if kind=='API':
        for t in d.tables:
            if t.rows and t.cell(0,0).text=='Method · 경로':
                for endpoint in ('access-jobs','project-jobs'):
                    cells=t.add_row().cells
                    cells[0].text=f'POST /api/v1/{endpoint}/validate'
                    cells[1].text='실제 자원의 읽기 전용 요청 검증; Job 등록·권한 변경 없음'
            if t.rows and t.cell(0,0).text=='환경 변수':
                cells=t.add_row().cells;cells[0].text='FGT_VALIDATION_TIMEOUT_SECONDS';cells[1].text='20 / 검증 전체 조회 제한; 1~25초'
    # Do not rely on brittle old page-number references after new content is inserted.
    for p in d.paragraphs:
        if ('쪽 /' in p.text or '3–4쪽' in p.text):
            p.text='본문의 번호 붙은 장 제목을 따라 읽고, 마지막 최신 기능·유지보수 장과 별도 오픈_운영_인수인계.docx를 함께 확인하세요.'
    for title,blocks in API_ADDITIONS if kind=='API' else CLIENT_ADDITIONS:add_blocks(d,title,blocks)
    d.core_properties.title=f'Folder Grant {kind} 설계·운영·유지보수 / 내부 오픈 기준'
    d.core_properties.subject=f'{DATE} / main {BASE[:12]} / 사용자 수정 베타본 기반 운영본'
    d.core_properties.comments='기존 Beta 파일을 수정하지 않고 복사하여 최신 내용을 반영했습니다.'
    swap_images(d.part.package,assets)
    d.save(out/OUTPUT_NAMES[0 if kind=='API' else 1])


def handover(out):
    # Reuse the established document typography without invoking the Beta builder.
    style.office_doc(OUTPUT_NAMES[2],'오픈·운영 인수인계','업무 책임자 · 서버 운영자 · Client 배포 담당자를 위한 체크리스트',HANDOVER_PAGES)
    p=out/OUTPUT_NAMES[2];d=Document(p)
    for para in paragraphs(d):
        replace_paragraph(para,{**COMMON,'내부 베타':'내부 오픈','베타 배포 전 점검과 인수인계에 함께 사용하세요.':'내부 오픈 점검과 인수인계에 사용하세요.'})
    for sec in d.sections:
        for para in sec.footer.paragraphs:replace_paragraph(para,{'내부 베타 배포':'내부 오픈 기준'})
    d.core_properties.subject=f'내부 오픈·운영 인수인계 / {DATE} / {BASE[:12]}'
    d.save(p)


def update_intro(out,assets):
    p=Presentation(DOCS/'Folder_Grant_프로그램_소개서_Beta.pptx')
    mapping={**COMMON,'FOLDER GRANT  /  BETA':'FOLDER GRANT  /  RELEASE','내부 베타 소개':'내부 오픈 안내',
      '베타 사용자는':'업무 사용자는','BETA':'RELEASE',
      '제한된 시험 대상에서 먼저 검증하는 Beta 릴리스':'개발 범위 구현 완료 · 현장 인수 후 내부 오픈',
      '미리보기 · 선택 실행 · 중지':'요청 검증 · 선택 실행 · 중지',
      '미리보기 / 선택 실행':'요청 검증 / 선택 실행','미리보기':'요청 검증',
      '대상과 계획을 먼저 확인':'실제 자원을 읽기 전용 조회',
      'Dry Run 확인 후 실제 실행 · 결과 검증':'요청 검증 → 선택 실행 → BUS 완료 재확인',
      'PLAN':'VALIDATION','실행 계획만 확인':'실제 그룹·폴더·ACL 조회',
      '실제 자원 변경 없음':'읽기 전용 · 권한 변경 없음',
      '작은 범위부터 검증하며 베타를 시작합니다':'현장 인수 후 내부 오픈을 진행합니다',
      '제한된 베타 운영':'제한된 내부 운영','Mock + Preview':'Mock + 요청 검증',
      'API 60개 / Client 93개 통과':'API 76개 통과·5개 선택 시험 제외 / Client 125개 통과',
      '베타 결과':'현장 인수 결과','베타 오픈 확인 목록':'오픈 확인 목록',
      '현재 BETA':'현재 범위','현재 RELEASE':'현재 범위'}
    for slide in p.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:replace_paragraph(para,mapping)
    # Clarify that old plan Preview APIs remain separate from the UI validation button.
    p.slides[7].notes_slide.notes_text_frame.text='요청 검증은 /validate로 실제 자원을 읽기 전용 조회합니다. 기존 /preview는 실행 계획 API로 별도 유지합니다. Mock에서는 실제 검증하지 않습니다.'
    p.core_properties.title='Folder Grant API + Client 프로그램 소개 / 내부 오픈 기준'
    p.core_properties.subject=f'{DATE} / {BASE[:12]} / 기존 편집본 기반 업데이트'
    swap_images(p.part.package,assets)
    p.save(out/OUTPUT_NAMES[3])


def completion_deck(out,assets):
    from pptx.util import Inches as I
    p=Presentation();p.slide_width=I(13.333333);p.slide_height=I(7.5)
    def slide(title,subtitle='',dark=False):
        s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=style.pcolor(style.NAVY if dark else style.BG)
        fg='FFFFFF' if dark else style.NAVY;muted='B7C7DF' if dark else style.MUTED
        style.text(s,.55,.25,11,.3,'FOLDER GRANT  /  DEVELOPMENT COMPLETION',10,'8DB1FF' if dark else style.BLUE,True)
        style.text(s,.55,.88,12.2,.7,title,29,fg,True)
        style.text(s,.55,1.63,12,.55,subtitle,14,muted)
        style.text(s,.55,7.13,11.8,.22,f'{DATE}  ·  API 0.4.0 / Client V1.1.0 Build:20261022  ·  main {BASE[:12]}',8,muted)
        style.text(s,12.25,7.1,.55,.25,f'{len(p.slides):02d}',10,muted)
        return s
    def cards(title,subtitle,items,dark=False):
        s=slide(title,subtitle,dark)
        n=len(items);gap=.24;width=(12.23-(n-1)*gap)/n
        for i,(tag,heading,body) in enumerate(items):style.card(s,.55+i*(width+gap),2.45,width,3.95,tag,heading,body,dark)
        return s
    def screenshot(title,subtitle,filename,heading,body):
        s=slide(title,subtitle)
        s.shapes.add_picture(str(assets/filename),I(.55),I(2.25),height=I(4.45))
        style.card(s,6.1,2.4,6.65,3.95,'실제 Qt 화면 / 오프라인 샘플',heading,body)
        style.text(s,.65,6.84,11.8,.2,'샘플 화면이며 실제 BUS·AD에 연결하거나 권한 작업을 수행한 증거는 아닙니다.',10,style.MUTED)
    s=slide('Folder Grant 개발 완료 보고','API 서버 + Windows Client  |  내부 오픈·운영 인수인계',True)
    style.text(s,.75,2.65,10.9,1.35,'요청 확인부터 권한 적용,\nBUS 완료까지 연결했습니다.',33,'FFFFFF',True)
    style.text(s,.75,4.55,10.7,1.1,'합의된 기능 구현과 자동 검증 결과를 정리하고\n현장 인수·운영 전환에 필요한 기준을 전달합니다.',20,'B7C7DF')
    style.text(s,.75,6.05,11,.5,'개발 범위 구현 완료  /  실제 환경 인수·오픈 승인은 별도 기록',17,'8DB1FF',True)
    cards('개발 목적과 완료 범위','기존 BUS 업무 흐름을 유지하면서 권한 실행을 서버로 분리했습니다.',[
      ('01  업무','조회에서 완료까지','권한·신규 프로젝트 요청 조회\n수동·Excel 입력과 선택 실행\n실제 성공 후 BUS 완료'),
      ('02  서버','공통 실행 기준','입력·정책 검증과 Job 이력\nAD 그룹·폴더 ACL·프로젝트 생성\n실패·부분 성공·재실행 구분'),
      ('03  인수','운영 가능한 기록','화면 요약 + 파일 상세 로그\n설계·사용·운영·유지보수 문서\n오픈·장애·백업 체크리스트')])
    s=slide('구성: Client와 API의 책임을 분리','BUS 연결은 Client가, 실제 AD·파일 권한 변경은 API 서버가 수행합니다.')
    s.shapes.add_picture(str(DOCS/'beta-assets/architecture.png'),I(.55),I(2.35),width=I(12.2))
    style.text(s,.8,6.6,11.8,.4,'SQLite는 Job·단계·감사를 보존합니다. BUS 완료 여부는 API 성공과 별도로 관리합니다.',15,style.BLUE,True)
    cards('구현 완료 기능','현재 코드에 있는 기능과 후속 개선 과제를 구분합니다.',[
      ('API 0.4.0','실행·검증·이력','실제 자원 읽기 전용 검증\n권한 부여·해제 / 신규 생성\n조회·취소·권한 실패 단계 재실행'),
      ('CLIENT 1.1.0','사용자 작업 흐름','시작 연결·로그인·두 목록 준비\n상세·실시간 단계·회전 표시\n양쪽 탭 BUS 완료 재확인')])
    screenshot('권한 요청: 대상을 확인하고 처리','BUS 요청은 읽기 전용, 수동·Excel 요청만 목록 편집이 가능합니다.','access.png','선택·검증·실행','체크한 항목을 순차 처리\n시작 시 처리 행을 자동 선택\n그룹·폴더·BUS 단계 상태 확인\n우클릭 삭제는 화면 목록만 제거')
    screenshot('신규 프로젝트: 생성 단계와 원본 상세','생성 준비부터 결과·BUS 완료까지 한 화면에서 확인합니다.','projects.png','공통 구조와 상세 정보','코드·이름·상태 중심 요약\n공통 상세 열 토글로 원본 표시\n서버 6단계 + 요청·BUS 진행\n긴 오류·요청번호는 스크롤·복사')
    cards('요청 검증·Preview·실제 실행의 구분','현재 화면의 요청 검증 버튼은 /validate를 호출합니다.',[
      ('READ ONLY','요청 검증','실제 사용자·그룹·요청 폴더\n현재 ACL·생성 준비 상태 조회\nJob 등록·권한·BUS 변경 없음'),
      ('DEVELOPMENT','Preview / Mock','Preview: 실행 계획만 계산\nMock: Job·DB 흐름 모의 실행\n실제 검증·변경·BUS 완료 없음'),
      ('POWERSHELL','선택 실행','서버 서비스 계정으로 변경\n그룹·ACL 또는 프로젝트 생성\n최종 실제 성공을 별도로 확인')])
    cards('진행 상태를 결과와 혼동하지 않도록','서버 조회 응답을 기준으로 표시하며 애니메이션은 처리 중임을 뜻합니다.',[
      ('PROGRESS','실시간 상세','각 요청 시작 시 해당 행 표시\n서버 단계에 맞춰 회전·체크\n다른 행을 선택해 상세 확인'),
      ('STATE','결과 보존','검증·라이브 진행·실제 Job 분리\n모의·미실행·실패·중지 구분\n타이머는 새 작업마다 초기화')])
    cards('BUS 완료는 실제 성공 뒤에만 수행','두 탭 모두 BUS 완료 재확인으로 최신 서버 결과를 다시 조회합니다.',[
      ('허용 조건','원본·실제 성공','BUS 원본 요청\npowershell + succeeded\n응답의 대상 일치 확인'),
      ('보류 조건','불확실하면 멈춤','부분 성공·실패·중지·Mock\n조회 오류·요청 불일치\n과거 성공 캐시로 우회 금지'),
      ('수행 범위','완료만 재확인','권한 Job 재등록 안 함\n프로젝트 재생성 안 함\n자동 완료 OFF여도 사용 가능')])
    cards('실패한 권한 작업의 복구 흐름','원인을 먼저 조치하고, 명시적 재실행으로 이력을 보존합니다.',[
      ('01  조사','실패 단계 확인','Job ID·요청번호·UNC·오류\n실제 그룹·ACL과 단계 비교\nAPI 성공·BUS 실패 구분'),
      ('02  조치','원인 해결','누락 폴더·서비스 계정 점검\n시간 초과 설정·부분 적용 확인\n불확실 Job 중복 실행 금지'),
      ('03  재실행','선택 실행 → 확인','새 Job으로 실패 단계부터 처리\n기존 성공 단계·부모 이력 보존\n프로젝트 전용 /retry는 없음')])
    cards('오픈 전 필요한 환경','설치·권한·연결과 비밀정보 취급을 인수인계합니다.',[
      ('SERVER','API 운영 환경','Windows·Python·AD 모듈\n전용 서비스 계정·공유/NTFS\n서비스·DB·템플릿·키·포트'),
      ('DESKTOP','Client 배포 환경','Windows·Python·Chrome\nBUS/API 연결·드라이버 준비\nassets·설정·로그 경로 쓰기')])
    cards('운영·보안·장시간 작업 기준','기본 설정과 운영자가 별도로 마련할 절차입니다.',[
      ('WAIT','서로 다른 제한','HTTP 30초 / 조회 최대 4시간\nPowerShell 단계 10800초\n읽기 전용 검증 20초'),
      ('RECORD','로그와 추적','화면은 요약, 파일은 상세\naccess.log 한 파일 누적\n서버 UTC·PC 로컬 시각 비교'),
      ('ACCESS','현재 보안 경계','공유 X-API-Key 인증\n키·저장 암호는 평문 파일\n접근 범위·HTTPS 별도 구성')])
    s=cards('검증 결과와 증거 범위','2026-10-26 / main ebbdfb9d4b8f 기준 자동 시험을 다시 실행했습니다.',[
      ('API','76개 통과','계약·계획·검증·취소·재실행\n선택 PowerShell 시험 5개 제외\n실제 AD·SMB 변경 시험 아님'),
      ('CLIENT','125개 통과','완료 차단·UI·스레드·진행\n상세 열·프로젝트 완료 재확인\n외부 API/BUS는 테스트 대역')])
    style.text(s,.8,6.65,11.8,.3,'Word·PPT는 구조 검사와 PDF 변환 후 배치를 확인합니다. 실제 업무 인수 시험은 운영자가 기록합니다.',12,style.MUTED)
    cards('현장 인수와 오픈 전환 순서','완료 보고는 오픈 승인 자체를 대신하지 않습니다.',[
      ('01','환경·연결 확인','서비스 계정·방화벽·DB 백업\nMock과 실제 시험 환경 분리'),
      ('02','실제 소량 시험','자동 완료 OFF → 실제 자원 확인\n실패·재실행·완료 재확인 시험'),
      ('03','승인 후 범위 확대','담당자·문의 경로·중단 기준\n인수 결과·미결·승인 기록')])
    cards('업데이트·장애·복구 시 운영 원칙','서비스 재시작이 작업 복구나 권한 롤백을 뜻하지 않습니다.',[
      ('UPDATE','작업 종료 후 배포','접수 중단·활성 Job 확인\n서비스 중지·코드/DB/설정 백업\n호환 버전 갱신 후 단계별 확인'),
      ('INCIDENT','실제 자원과 대조','기존 Job·잔존 PowerShell 조사\n부분 변경·BUS 상태 분리\n코드·DB 복원만으로 ACL 복구 안 됨')])
    cards('현재 제한과 후속 개선 과제','아래 항목은 이번 릴리스의 구현 완료 기능이 아닙니다.',[
      ('CURRENT','운영 시 알고 쓸 점','프로세스 내 BackgroundTasks\n재시작 후 진행 Job 복구 없음\n자동 롤백·로그 순환 없음'),
      ('NEXT','별도 개선 검토','외부 Worker·내구성 큐·복구\n개인 인증·권한·실행 잠금\n새 요청 자동 감지는 현재 보류')],True)
    cards('개발 산출물과 인수인계','최신 운영본과 사용자가 수정한 기존 베타본을 함께 보존합니다.',[
      ('WORD','설계·운영 기준','API 설계·운영·유지보수\nClient 설계·사용·유지보수\n오픈·운영 인수인계'),
      ('PPT / README','설명·완료 보고','프로그램 소개서 갱신본\n개발 완료 보고서\n설치·기능·제한·체크리스트')])
    s=slide('개발 범위 구현 완료 · 운영 인수 준비','실제 환경 시험 결과와 승인 기록을 남긴 뒤 내부 오픈을 진행합니다.',True)
    style.text(s,.8,2.7,11.8,1.3,'다음 단계는 현장 인수입니다.\n담당자·대상·오픈 범위를 확정합니다.',31,'FFFFFF',True)
    style.text(s,.8,4.75,11.6,1.1,'확인 항목: 실제 AD/ACL · BUS 완료 · 실패 복구 · 백업\n기록 위치: docs/오픈_운영_인수인계.docx',20,'B7C7DF')
    for s in p.slides:
        s.notes_slide.notes_text_frame.text='설명 기준: main '+BASE+' / '+DATE+'\n자동 시험과 실제 환경 인수 시험을 구분합니다. 구현되지 않은 과제를 완료로 설명하지 않습니다.'
    p.core_properties.title='Folder Grant 개발 완료 보고서'
    p.core_properties.subject=f'구현 범위·검증 결과·내부 오픈·운영 인수 / {DATE}'
    p.core_properties.author='Folder Grant Project'
    p.save(out/OUTPUT_NAMES[4])


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=DOCS)
    parser.add_argument('--refresh',action='store_true')
    args=parser.parse_args();out=args.output.resolve();out.mkdir(exist_ok=True,parents=True)
    manifest=out/'release-manifest.json'
    previous=json.loads(manifest.read_text()) if manifest.exists() else {}
    for name in OUTPUT_NAMES:
        path=out/name
        if path.exists() and (not args.refresh or previous.get('outputs',{}).get(name)!=sha(path)):
            raise SystemExit(f'Refusing to overwrite existing/manual-edited output: {path}. Use another --output directory.')
    assets=DOCS/'release-assets'
    for name in ('access.png','access-dark.png','projects.png'):
        if not (assets/name).exists():raise SystemExit(f'Missing current UI capture: {name}')
    style.DATE=DATE;style.BASE=BASE
    with TemporaryDirectory(prefix='fgt-release-') as temporary:
        stage=Path(temporary);style.OUT=stage
        update_doc('API',stage,assets);update_doc('Client',stage,assets);handover(stage)
        update_intro(stage,assets);completion_deck(stage,assets)
        for name in OUTPUT_NAMES:shutil.copy2(stage/name,out/name)
    inputs=list(DOCS.glob('*Beta.docx'))+list(DOCS.glob('*Beta.pptx'))
    manifest.write_text(json.dumps(dict(date=DATE,source_commit=BASE,
      inputs={str(x.relative_to(ROOT)):sha(x) for x in inputs},
      outputs={name:sha(out/name) for name in OUTPUT_NAMES}),ensure_ascii=False,indent=2)+'\n')
    print('Created release Office copies:',*OUTPUT_NAMES,sep='\n')

if __name__=='__main__':main()
