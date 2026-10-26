"""Check release Office packages, links, and preservation of edited Beta inputs."""
from pathlib import Path
import hashlib
import json
import re
import subprocess
import sys
import zipfile
from xml.etree import ElementTree as ET
from docx import Document
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[2]
DOCS=ROOT/'docs'
sys.path.insert(0,str(Path(__file__).parent))
from release_content import BASE


def digest(data):return hashlib.sha256(data).hexdigest()


def word_text(document):
    yield from (p.text for p in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:yield from word_text(cell)


def main():
    manifest=json.loads((DOCS/'release-manifest.json').read_text())
    assert manifest['source_commit']==BASE
    for name,expected in manifest['inputs'].items():
        data=(ROOT/name).read_bytes()
        baseline=subprocess.check_output(['git','show',f'{BASE}:{name}'],cwd=ROOT)
        assert digest(data)==expected==digest(baseline),f'Beta input changed: {name}'
    for name,expected in manifest['outputs'].items():
        path=DOCS/name
        assert digest(path.read_bytes())==expected,f'Office file changed since generation: {name}'
        with zipfile.ZipFile(path) as z:
            assert z.testzip() is None
            for item in z.namelist():
                if item.endswith('.xml'):ET.fromstring(z.read(item))
        if path.suffix=='.docx':
            text='\n'.join(word_text(Document(path)))
            assert BASE[:12] in text
            for old in ('수동 요청으로 전환되며','미리보기(Dry Run)를 켜고','보관된 성공 결과를 이용해','자동 테스트 60개','자동 테스트 93개'):
                assert old not in text,(name,old)
        else:
            p=Presentation(path)
            expected_slides=18 if '개발완료' in name else 14
            assert len(p.slides)==expected_slides
            text='\n'.join(s.text for slide in p.slides for s in slide.shapes if s.has_text_frame)
            assert '요청 검증' in text and 'BUS 완료 재확인' in text
            assert 'API 60개' not in text and 'Client 93개' not in text
            for slide in p.slides:
                for shape in slide.shapes:
                    assert shape.left>=0 and shape.top>=0,(name,shape.name)
                    assert shape.left+shape.width<=p.slide_width+10000,(name,shape.name)
                    assert shape.top+shape.height<=p.slide_height+10000,(name,shape.name)
        print('OK',name)
    for file in [ROOT/'README.md',ROOT/'API/README.md',ROOT/'Client/README.md',ROOT/'API/test_client/README.md',DOCS/'README.md',DOCS/'OPEN_CHECKLIST.md',DOCS/'RELEASE_NOTES.md']:
        for target in re.findall(r'\]\(([^)]+)\)',file.read_text()):
            if '://' in target or target.startswith('#'):continue
            path=target.split('#',1)[0]
            assert (file.parent/path).exists(),f'Broken link: {file}: {target}'
    print('OK local README links and unchanged manually edited Beta inputs')

if __name__=='__main__':main()
